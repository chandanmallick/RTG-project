import asyncio
import io
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock,patch
from openpyxl import load_workbook
from docx import Document
from test_frequency_threshold_analysis import fixture,DB,USER
from routes import frequency_routes
from routes.frequency_analysis_routes import ResultPayload,FetchSourcesPayload,export_result
from services import frequency_analysis_sessions as sessions
from services.frequency_threshold_analysis import calculate
from services.frequency_auto_sources import fetch_sources
from services.frequency_compact_reports import performance_table


async def response_bytes(response):
    output=[]
    async for chunk in response.body_iterator:output.append(chunk)
    return b''.join(output)


class CompactReportsTests(unittest.TestCase):
    def setUp(self):sessions._sessions.clear()

    def test_compact_daily_and_monthly_rows_format_frequency_and_od(self):
        row={'date':'2026-10-03','event':1,'entity':'Bihar','period_start':'2026-10-03T17:00:00',
             'period_end':'2026-10-03T17:01:00','selected_minutes':1,'lowest_frequency':49.876,
             'thresholds':{'49.90':{'frequency_minutes':1,'adverse_minutes':.5,'adverse_pct':50,
                                    'average_od_ui_mw':5.49,'maximum_od_ui_mw':12.51}}}
        _,records=performance_table([row],['49.90'])
        self.assertEqual(records[0]['lowest_frequency'],'49.876')
        self.assertEqual(records[0]['49.90_mw'],'5 / 13')

    def test_high_uses_opposite_adverse_signs_and_separate_threshold(self):
        data=fixture();data['frequency']=data['frequency'].copy();data['frequency'][:]=50.10
        result=calculate(data,[('2026-10-03T17:00:00','2026-10-03T17:03:00')],[],event_type='high')
        self.assertEqual(list(result['summary']['thresholds']),['50.05'])
        state=result['overall_performance']['State'][0]['thresholds']['50.05']
        generator=result['overall_performance']['ISGS'][0]['thresholds']['50.05']
        self.assertEqual(state['adverse_minutes'],.5)
        self.assertEqual(generator['adverse_minutes'],.5)
        self.assertEqual(state['maximum_od_ui_mw'],-5)
        self.assertEqual(generator['maximum_od_ui_mw'],5)

    @patch.object(frequency_routes,'get_crms_frequency_transmission_lines',new_callable=AsyncMock,return_value={'success':True,'events':[]})
    def test_all_report_formats_compact_one_entity_event_row_and_monthly(self,physical):
        token=sessions._put(USER,fixture(),{'filename':'Compact source'})['session_token']
        with patch.object(frequency_routes,'fetch_crms_frequency_messages',new=AsyncMock(return_value=([],0))):
            result=sessions.analyse(token,USER,[('2026-10-03T17:00:00','2026-10-03T17:01:00'),('2026-10-03T17:02:00','2026-10-03T17:03:00')],db=DB)
        payload=ResultPayload(session_token=token,result_token=result['result_token'],layout='compact',format='html')
        response=export_result(payload,USER)
        html=response['html_document']
        self.assertNotIn('echarts',html);self.assertNotIn('series_timestamps',html)
        self.assertIn('Daily Event Chronology',html);self.assertIn('State Monthly Performance',html)
        self.assertIn('Download HTML',html);self.assertIn('Daily events side by side',html)
        for fmt in ('xlsx','docx','pdf'):
            payload.format=fmt;contents=asyncio.run(response_bytes(export_result(payload,USER)))
            if fmt=='xlsx':
                book=load_workbook(io.BytesIO(contents))
                self.assertEqual(book.sheetnames[0],'Bihar')
                self.assertEqual(book['State Event Performance'].max_row,3)
                headers=[cell.value for cell in book['State Event Performance'][1]]
                self.assertTrue(any('<49.90' in str(h) for h in headers))
                self.assertTrue(any('<49.70' in str(h) for h in headers))
                self.assertTrue(any('<49.50' in str(h) for h in headers))
                self.assertIn('<49.90 Minutes',[cell.value for cell in book['Event Frequency Statistics'][1]])
            elif fmt=='docx':self.assertTrue(Document(io.BytesIO(contents)).tables)
            else:self.assertTrue(contents.startswith(b'%PDF'))
        physical.assert_awaited_once()

    def test_automatic_fetch_reuses_readers_and_routes_state_generators_to_rtg(self):
        mappings=[{'plant_id':'STATE_BIHAR','plant_name':'Bihar','type':'State','is_state':True,'wbes_acronym':'BSP','mis_name':'Bihar'},
            {'plant_id':'unit','plant_name':'Unit','type':'ISGS','wbes_acronym':'UNIT','mis_name':'Unit'},
            {'plant_id':'local','plant_name':'Local','type':'IPP','utility_type':'State_IPP','rtg_plant_id':'rtg-local','mis_name':'Local'}]
        db=SimpleNamespace(map_collection=SimpleNamespace(find=lambda *args:mappings))
        curve={'success':True,'points':[{'timestamp':'2026-10-03T17:00:00','frequency':49.8},{'timestamp':'2026-10-03T17:00:30','frequency':50.1},{'timestamp':'2026-10-03T17:01:00','frequency':50.0}]}
        async def actual(start,end,names,frequency,kind):
            return {'rows':[{'timestamp':'2026-10-03T17:00:00',**{name:90 for name in names.split(',')}},{'timestamp':'2026-10-03T17:01:00',**{name:0 for name in names.split(',')}}]}
        wbes=[{'Acronym':name,'NetScheduleSummary':{'TotalNetSchdAmount':[100]*96}} for name in ('BSP','UNIT')]
        with patch('services.curve_frequency_service.load_curve_frequency_range',return_value=curve) as reader,patch.object(frequency_routes,'fetch_wbes_schedule_raw',return_value=wbes) as wb,patch.object(frequency_routes,'fetch_rtg_schedule_raw',return_value={'schedule':[30]*96}) as rtg,patch.object(frequency_routes,'get_schedule_data_actual',new=actual),patch.object(frequency_routes,'fetch_crms_frequency_messages',new=AsyncMock(return_value=([],0))):
            payload=FetchSourcesPayload(start_date='2026-10-03',end_date='2026-10-03',sources=['wbes','rtg','mis','crms'])
            response=fetch_sources(payload,USER,db)
            dataset=sessions.get_session(response['session_token'],USER)['dataset']
            self.assertEqual(dataset['entities'][0]['deviation'][0],-10)
            self.assertEqual(dataset['entities'][0]['actual'][-1],0)
            self.assertIsNone(dataset['entities'][2]['group'])
            self.assertEqual(dataset['entities'][2]['schedule'][0],30)
            rtg.assert_called_once_with('2026-10-03','rtg-local',force_refresh=True)
            self.assertEqual(wb.call_args.args[1],['BSP','UNIT'])
            payload.session_token=response['session_token'];payload.sources=['mis']
            refreshed=fetch_sources(payload,USER,db);reader.assert_called_once()
            wb.return_value=[];payload.session_token=refreshed['session_token'];payload.sources=['wbes']
            failed=fetch_sources(payload,USER,db)
            self.assertEqual(sessions.get_session(failed['session_token'],USER)['dataset']['entities'][0]['schedule'][0],100)
            self.assertFalse(failed['source_status'][0]['available'])
            self.assertTrue(failed['warnings'])

    def test_low_high_cache_segregation_and_all_high_exports(self):
        dataset=fixture();dataset['frequency']=dataset['frequency'].copy();dataset['frequency'][1]=50.1
        token=sessions._put(USER,dataset,{'filename':'Mixed frequency','crms_enabled':False})['session_token']
        ranges=[('2026-10-03T17:00:00','2026-10-03T17:03:00')]
        low=sessions.analyse(token,USER,ranges,db=DB,event_type='low')
        high=sessions.analyse(token,USER,ranges,db=DB,event_type='high')
        self.assertNotEqual(low['result_token'],high['result_token'])
        self.assertEqual(high['summary']['thresholds']['50.05']['frequency_minutes'],.5)
        for fmt in ('html','xlsx','pdf','docx'):
            payload=ResultPayload(session_token=token,result_token=high['result_token'],format=fmt)
            response=export_result(payload,USER)
            if fmt=='html':self.assertIn('High Frequency Analysis',response['html_document'])
            else:self.assertTrue(asyncio.run(response_bytes(response)))
