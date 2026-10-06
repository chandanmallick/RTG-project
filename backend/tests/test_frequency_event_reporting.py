import asyncio
import io
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from openpyxl import load_workbook
from docx import Document
from routes import frequency_routes as routes
from services import frequency_event_reporting as reporting

EVENT = {'event_id':'one','name':'Name does not define timestamps','start_time':'2026-10-05T17:02:00','end_time':'2026-10-05T17:15:00','event_type':'low'}

def point(kind='State', deviations=None):
    return {'plant_id':kind,'plant_name':kind,'type':kind,'series':{'timestamps':['2026-10-05T17:02:00','2026-10-05T17:02:30','2026-10-05T17:03:00','2026-10-05T17:15:00'], 'frequency':[49.8,49.9,49.7,49.85],'deviation':deviations or [10,-4,20,6]}}

def entity(kind='State', deviations=None):
    return {'entity_id':kind,'display_name':kind,'group':kind,'point':point(kind,deviations),'mapping':{}}

class SavedEventTests(unittest.TestCase):
    def test_signed_average_threshold_denominator_and_partial_block(self):
        messages = [{'entity_id':'State','timestamp':'2026-10-05T17:02:30','message_no':'M1'}]*2
        messages += [{'entity_id':'State','timestamp':EVENT['end_time'],'message_no':'M2'}]
        performance, lowest = reporting.build_performance(EVENT,[entity(),entity('ISGS',[-10,4,-20,-6]),entity('IPP',[None]*4)],messages)
        state=performance['State'][0]
        self.assertEqual(state['average_od_ui_mw'],8)
        self.assertEqual(state['maximum_od_ui_mw'],20)
        self.assertEqual(state['od_ui_time_pct'],100)
        self.assertEqual(state['message_count'],2)
        self.assertEqual(state['period_start'],EVENT['start_time'])
        self.assertEqual(state['period_end'],EVENT['end_time'])
        self.assertEqual(len(performance['State']),1)
        self.assertEqual(performance['ISGS'][0]['average_od_ui_mw'],-8)
        self.assertEqual(performance['ISGS'][0]['maximum_od_ui_mw'],-20)
        self.assertIsNone(performance['IPP'][0]['average_od_ui_mw'])
        self.assertEqual(lowest,49.7)
        unknown,_=reporting.build_performance(EVENT,[entity()],messages,False)
        self.assertIsNone(unknown['State'][0]['message_count'])

    def test_classification_and_invalid_structured_metadata(self):
        self.assertEqual(reporting.point_group({'type':'Generator'},{'type':'ISGS'}),'ISGS')
        self.assertIsNone(reporting.point_group({'type':'Generator'},{'type':'State_IPP'}))
        self.assertEqual(reporting.point_group({'type':'State'},{}),'State')
        with self.assertRaises(ValueError): reporting.event_period({'name':'Low Freq 3-Oct-26 (17:00-22:59)'})
        self.assertEqual(reporting.event_period({'start_time':'2026-10-05T11:32:00Z','end_time':'2026-10-05T11:45:00Z'})[0].hour,17)

    def test_exports_identify_events_and_respect_sections(self):
        performance,lowest=reporting.build_performance(EVENT,[entity(),entity('ISGS',[-10,4,-20,-6])],[])
        event={**EVENT,'event_name':'=Unsafe <name>','lowest_frequency':lowest,'performance':performance,'chronology':[{'timestamp':EVENT['start_time'],'message_details':'=SUM(A1) <script>','state':'State'}],'warnings':['Counts unavailable'],'calculation_note':'Signed convention'}
        payload={'include_chronology':True,'include_entity_performance':True,'performance_groups':['State','ISGS','IPP'],'supplemental_events':[event,{**event,'event_name':'Second event'}]}
        wb=load_workbook(reporting.supplements_excel(payload['supplemental_events'],payload))
        self.assertEqual(wb.sheetnames,['Chronology','State Performance','ISGS Performance','IPP Performance'])
        self.assertEqual(wb['Chronology']['A2'].data_type,'s')
        self.assertEqual(wb['Chronology']['J2'].data_type,'s')
        self.assertEqual(wb['Chronology']['K2'].value,'Counts unavailable')
        self.assertIn('Second event',[row[0] for row in wb['Chronology'].iter_rows(values_only=True)])
        html=reporting.supplements_html([event],payload)
        self.assertIn('&lt;script&gt;',html)
        doc=Document(); reporting.append_docx_supplements(doc,payload)
        self.assertEqual(len(doc.tables),6)
        off={**payload,'include_chronology':False,'include_entity_performance':False}
        doc=Document();reporting.append_docx_supplements(doc,off)
        self.assertEqual(len(doc.sections),1)
        self.assertEqual(reporting.supplements_html([event],off),'')

    def test_chronology_reuses_service_and_selected_event_only(self):
        event={**EVENT,'data_points':[point()]}
        collection=SimpleNamespace(find_one=lambda *args,**kwargs:event)
        db=SimpleNamespace(db={routes.EVENT_COLLECTION:collection})
        messages=[{'timestamp':'2026-10-05 17:02:00','message_no':'M1','issued_to':['State'],'remarks':'Text','category':['Frequency']}, {'timestamp':'2026-10-05 18:00:00','message_no':'OUT','issued_to':['State']}]
        aliases={'STATE':[entity()]}
        with patch.object(routes,'MongoService',return_value=db),patch.object(reporting,'timeline_aliases',return_value=aliases),patch.object(routes,'fetch_crms_frequency_messages',new=AsyncMock(return_value=(messages,0))),patch.object(routes,'_nearest_saved_event_values',side_effect=AssertionError('Must not load another event')):
            response=asyncio.run(routes.build_frequency_message_timeline(routes.FrequencyMessageTimelinePayload(event_id='one',ranges=[routes.FrequencyMessageRange(start_time=EVENT['start_time'],end_time=EVENT['end_time'])])))
        self.assertEqual(len(response['rows']),1)
        self.assertEqual(response['rows'][0]['deviation_mw'],10)
        self.assertEqual(response['rows'][0]['message_details'],'Text')

    def test_crms_failure_preserves_saved_messages_without_zero_counts(self):
        data=point();data['crms_messages']=[{'timestamp':'2026-10-05 17:02:00','message_no':'saved','remarks':'Saved text'}]
        event={**EVENT,'data_points':[data]}
        db=SimpleNamespace(db={routes.EVENT_COLLECTION:SimpleNamespace(find_one=lambda *args,**kwargs:event)})
        with patch.object(routes,'MongoService',return_value=db),patch.object(reporting,'timeline_aliases',return_value={'STATE':[entity()]}),patch.object(routes,'fetch_crms_frequency_messages',new=AsyncMock(side_effect=RuntimeError('Offline'))):
            response=asyncio.run(routes.build_frequency_message_timeline(routes.FrequencyMessageTimelinePayload(event_id='one',ranges=[routes.FrequencyMessageRange(start_time=EVENT['start_time'],end_time=EVENT['end_time'])])))
        self.assertFalse(response['messages_complete'])
        self.assertEqual(response['rows'][0]['message_no'],'saved')
        self.assertTrue(response['warnings'])

    def test_report_export_reuses_existing_generators_and_rejects_invalid_groups(self):
        performance,lowest=reporting.build_performance(EVENT,[entity()],[])
        event={**EVENT,'event_name':'Event 1','lowest_frequency':lowest,'chronology':[],'performance':performance,'warnings':[],'calculation_note':'Signed convention'}
        doc={**EVENT,'data_points':[point()]}
        async def run():
            with patch.object(routes,'_saved_report_events',new=AsyncMock(return_value=(None,[event],[doc]))):
                payload=routes.FrequencySavedReportPayload(event_ids=['one'],format='docx',include_existing_sections=False)
                response=await routes.export_saved_frequency_events(payload)
                parts=[]
                async for part in response.body_iterator:parts.append(part)
                document=Document(io.BytesIO(b''.join(parts)))
                self.assertEqual(len(document.tables),1)
                self.assertNotIn('Executive Summary', '\n'.join(p.text for p in document.paragraphs))
                payload.format='pdf'
                response=await routes.export_saved_frequency_events(payload)
                parts=[]
                async for part in response.body_iterator:parts.append(part)
                self.assertTrue(b''.join(parts).startswith(b'%PDF'))
                payload.format='html'
                response=await routes.export_saved_frequency_events(payload)
                self.assertEqual(response['rows'],[])
                self.assertIn('State Performance',response['supplemental_html'])
                payload.performance_groups=['invalid']
                with self.assertRaises(routes.HTTPException):await routes.export_saved_frequency_events(payload)
        asyncio.run(run())

    def test_existing_saved_report_preserves_separate_stages(self):
        first={**point('ISGS'), 'plant_id':'plant','stage_id':'1','stage_name':'Stage 1'}
        second={**point('ISGS',[-1,-2,-3,-4]), 'plant_id':'plant','stage_id':'2','stage_name':'Stage 2'}
        event={**EVENT,'data_points':[first,second]}
        db=SimpleNamespace(db={routes.EVENT_COLLECTION:SimpleNamespace(find_one=lambda *args,**kwargs:event)})
        with patch.object(routes,'lookup_rtg_capacity_on_bar',return_value={}),patch.object(routes,'lookup_unit_fuel_name',return_value='Coal'),patch.object(routes,'build_source_status',return_value={}):
            result,error=routes.build_saved_event_response(db,'one',[],*reporting.event_period(EVENT))
        self.assertIsNone(error)
        self.assertEqual([row['stage_id'] for row in result['rows']],['1','2'])
        self.assertEqual(result['rows'][0]['series']['deviation'][0],10)
        self.assertEqual(result['rows'][1]['series']['deviation'][0],-1)

    def test_legacy_word_export_still_includes_existing_sections_by_default(self):
        async def run():
            response=await routes.download_docx({'rows':[], 'intro_desc':'Existing notes'})
            parts=[]
            async for part in response.body_iterator:parts.append(part)
            document=Document(io.BytesIO(b''.join(parts)))
            text='\n'.join(p.text for p in document.paragraphs)
            self.assertIn('Executive Summary',text)
            self.assertIn('Existing notes',text)
        asyncio.run(run())

if __name__=='__main__':unittest.main()
