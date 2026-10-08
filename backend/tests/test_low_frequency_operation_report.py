import asyncio
import base64
import io
import unittest
from unittest.mock import AsyncMock, Mock, patch
from types import SimpleNamespace
from docx import Document
from PIL import Image
from fastapi import HTTPException
from routes import frequency_routes as routes
from services.low_frequency_operation_report import render, SECTIONS


def fixture():
    metrics = {'frequency_minutes': 2, 'adverse_minutes': 1, 'adverse_pct': 50, 'average_od_ui_mw': 12, 'maximum_od_ui_mw': 20, 'message_count': None}
    row = {'entity': 'Bihar', 'entity_id': 'state', 'thresholds': {level: dict(metrics) for level in ('49.50','49.70','49.90')}}
    event = {'event_id': 'one', 'event_name': 'Event one', 'event_type': 'low', 'start_time': '2026-10-02T04:00:00', 'end_time': '2026-10-02T04:05:00', 'messages_complete': False,
             'warnings': ['SECRET traceback API failure'], 'chronology': [], 'report_entities': [{'entity_id':'state','display_name':'Bihar','group':'State'}, {'entity_id':'gen','display_name':'Station','group':'ISGS'}],
             'threshold_analysis': {'summary': {'minimum_frequency':49.4,'minimum_timestamp':'2026-10-02T04:01:00','thresholds':{level:dict(metrics) for level in ('49.50','49.70','49.90')}}, 'overall_performance':{'State':[row], 'ISGS':[{**row,'entity':'Station','entity_id':'gen'}], 'IPP':[{**row,'entity':'No UI','thresholds':{level:{**metrics,'adverse_minutes':0} for level in ('49.50','49.70','49.90')}}]}}}
    image = io.BytesIO(); Image.new('RGB',(1400,520),'white').save(image,format='PNG')
    charts = [{'event_id':'one','entity_id':entity,'title':title,'kind':kind,'is_state':state,'image':base64.b64encode(image.getvalue()).decode()} for entity,title,kind,state in [('', 'System frequency','System Frequency',False),('state','Bihar frequency vs deviation','Annexure - Deviation / Frequency',True),('state','Bihar drawal vs schedule','Annexure - State Schedule / Actual',True),('gen','Station frequency vs deviation','Annexure - Deviation / Frequency',False),('gen','Station actual vs schedule','Annexure - Generator Schedule / Actual',False)]]
    return event, charts


class OperationReportTests(unittest.TestCase):
    def test_all_formats_sections_grouping_no_diagnostics_and_same_images(self):
        event, charts = fixture(); options = {'operation_sections': list(SECTIONS),'action_summary':'Operator action <reviewed>'}
        html = render(event,options,charts,'html')
        for title in ('1 Executive Summary','2 State Performance','3 Central Sector','4 Action &','5 Action of','Annexure 1','Annexure 2'): self.assertIn(title,html)
        self.assertNotIn('SECRET',html); self.assertNotIn('No UI',html)
        self.assertEqual(html.count('<img '),5)
        self.assertIn('rowspan="3"',html)
        self.assertIn('Total frequency duration: 2 min',html)
        self.assertLess(html.index('Freq &lt;49.5'),html.index('Freq &lt;49.7'))
        word = Document(render(event,options,charts,'docx'))
        self.assertEqual(len(word.inline_shapes),5)
        self.assertEqual(word.tables[1].cell(0,0).text,'Reporting Period')
        self.assertIn('Total frequency duration',word.tables[1].cell(1,3).text)
        self.assertTrue(any('ERLDC' in p.text for s in word.sections for p in s.footer.paragraphs))
        self.assertNotIn('SECRET','\n'.join(p.text for p in word.paragraphs))
        self.assertTrue(render(event,options,charts,'pdf').read().startswith(b'%PDF'))
        only = render(event,{'operation_sections':['defence']},charts,'html')
        self.assertNotIn('Annexure',only); self.assertNotIn('Executive Summary',only)
        self.assertIn('Data not available',only)

    def test_operator_selection_includes_state_generator_html_charts_only(self):
        from services.low_frequency_operation_report import report_blocks
        event,charts=fixture()
        event['report_entities'].append({'entity_id':'local','display_name':'State Plant','group':None})
        local=[{**chart,'entity_id':'local','title':'State Plant '+chart['kind']} for chart in charts if chart.get('entity_id')=='gen']
        blocks,_=report_blocks(event,{'operation_sections':list(SECTIONS),'operation_entity_ids':['local']},charts+local)
        images=[chart for kind,chart in blocks if kind=='image']
        self.assertEqual(len(images),3)
        self.assertEqual({chart['entity_id'] for chart in images if chart['kind']!='System Frequency'},{'local'})
        self.assertTrue(any('State Plant' in value for kind,value in blocks if kind=='heading'))
        self.assertFalse(any('Bihar CRMS' in value for kind,value in blocks if kind=='heading'))
        html=render(event,{'operation_sections':['states']},charts,'html')
        self.assertIn('Low Frequency Report 02-Oct-26 04:00-04:05 hrs.',html)
        rows=next(value[1] for kind,value in report_blocks(event,{'operation_sections':['states']},charts)[0] if kind=='table')
        self.assertEqual(len(rows),1)
        self.assertEqual(len(rows[0]),15)

    def test_section_notes_and_threshold_statistics_follow_every_annexure_curve(self):
        from services.low_frequency_operation_report import report_blocks
        event,charts=fixture()
        notes={key:'Operator notes for '+key for key in SECTIONS}
        blocks,_=report_blocks(event,{'operation_section_notes':notes},charts)
        for value in notes.values(): self.assertIn(('paragraph',value),blocks)
        for at,(kind,chart) in enumerate(blocks):
            if kind!='image' or chart['kind']=='System Frequency':continue
            next_kind,(headers,rows,grouped)=blocks[at+1]
            self.assertEqual(next_kind,'table')
            self.assertEqual(headers[-2:],['OD/UI Duration (Min)','OD/UI Duration (%)'])
            self.assertEqual(rows[0],['<49.9 Hz',2,1,50])
        document=Document(render(event,{'operation_section_notes':notes},charts,'docx'))
        self.assertEqual(str(document.tables[0].cell(0,0).paragraphs[0].runs[0].font.color.rgb),'FFFFFF')

    def test_state_generator_annexure_uses_the_existing_timestamp_weighted_calculation(self):
        import numpy as np
        from services.frequency_threshold_analysis import timeline,calculate
        times,frequency,cadence=timeline(['2026-10-05T16:34:30','2026-10-05T16:35:00','2026-10-05T16:35:30'],[49.8,49.6,49.4])
        dataset={'times':times,'frequency':frequency,'cadence':cadence,'entities':[{'entity_id':'local','display_name':'State Plant','group':None,'deviation':np.array([-10,20,-20])}]}
        result=calculate(dataset,[('2026-10-05T16:34:30','2026-10-05T16:36:00')],[],include_annexure_entities=True)
        values=result['annexure_performance']['local']['thresholds']
        self.assertEqual(values['49.90']['adverse_minutes'],1)
        self.assertAlmostEqual(values['49.90']['adverse_pct'],66.666667,places=5)
        self.assertEqual(values['49.50']['adverse_minutes'],.5)
        self.assertEqual(result['overall_performance']['ISGS'],[])

    def test_export_rejects_cross_event_and_stale_capture(self):
        event,charts = fixture()
        async def run():
            with patch.object(routes,'_saved_report_events',new=AsyncMock(return_value=(None,[event],[{'updated_at':'revision'}]))):
                payload = routes.FrequencySavedReportPayload(event_ids=['one'],format='html',operation_report=True,operation_charts=charts,chart_event_id='other',chart_revision='revision')
                with self.assertRaises(HTTPException): await routes.export_saved_frequency_events(payload)
                payload.chart_event_id='one'; payload.chart_revision='stale'
                with self.assertRaises(HTTPException): await routes.export_saved_frequency_events(payload)
                payload.chart_revision='revision'
                response=await routes.export_saved_frequency_events(payload)
                self.assertIn('Annexure 2',response['html'])
        asyncio.run(run())

    def test_export_analysis_cache_reads_raw_event_and_calculates_once(self):
        from services import frequency_event_reporting as reporting
        from services import frequency_threshold_analysis as thresholds
        event = {'event_id':'cache-test','updated_at':'revision','name':'Cache event','event_type':'low','start_time':'2026-10-02T04:00:00','end_time':'2026-10-02T04:05:00',
                 'data_points':[{'plant_id':'STATE_TEST','plant_name':'State','type':'State','series':{'timestamps':['2026-10-02T04:00:00','2026-10-02T04:01:00'],'frequency':[49.4,49.6],'deviation':[10,20]}}]}
        collection = SimpleNamespace(find_one=Mock(side_effect=lambda query, projection: {'event_id':'cache-test','updated_at':'revision'} if 'updated_at' in projection else event))
        db=SimpleNamespace(db={routes.EVENT_COLLECTION:collection},map_collection=SimpleNamespace(find=lambda *args:[]))
        async def run():
            with patch.object(routes,'_build_frequency_message_timeline',new=AsyncMock(return_value={'rows':[],'messages_complete':True})),patch.object(thresholds,'calculate',wraps=thresholds.calculate) as calculate:
                first,_=await reporting.saved_event_analysis(db,'cache-test')
                second,document=await reporting.saved_event_analysis(db,'cache-test')
                self.assertIs(first,second); self.assertIs(document,event)
                calculate.assert_called_once()
                self.assertEqual(sum('updated_at' not in call.args[1] for call in collection.find_one.call_args_list),1)
        try: asyncio.run(run())
        finally: reporting._cache.pop(('cache-test','revision'),None)

    def test_operation_chart_sources_never_fall_back_to_outside_event_series(self):
        doc={'event_id':'one','event_type':'low','data_points':[{'plant_id':'STATE_TEST','plant_name':'State','type':'State','series':{'timestamps':['2026-10-01T04:00:00'],'frequency':[49.4],'deviation':[20],'actual':[30],'schedule':[10]}}]}
        with patch.object(routes,'lookup_rtg_capacity_on_bar',side_effect=AssertionError('Must reuse saved capacity')),patch.object(routes,'lookup_unit_fuel_name',return_value=''),patch.object(routes,'build_source_status',return_value={}):
            from datetime import datetime
            result,error=routes.build_saved_event_response(None,'one',[],datetime(2026,10,2,4),datetime(2026,10,2,5),saved_document=doc,refresh_capacity=False)
        self.assertIsNone(error)
        for key in ('timestamps','frequency','actual','schedule','deviation'): self.assertEqual(result['rows'][0]['series'][key],[])

    def test_chronology_deduplicates_messages_across_recipients_without_losing_values(self):
        from services.low_frequency_operation_report import report_blocks
        event,charts=fixture()
        row={'timestamp':'2026-10-02T04:01:00','message_no':'M1','message_details':'Reduce deviation','message_type':'Alert','frequency_hz':49.6,'state':'Bihar','deviation_mw':10}
        event['chronology']=[row,row,{**row,'state':'Station','deviation_mw':-3}]
        blocks,_=report_blocks(event,{'operation_sections':['actions']},charts)
        rows=next(value[1] for kind,value in blocks if kind=='table')
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0][2],'Bihar, Station')
        self.assertIn('Bihar: 10',rows[0][3]);self.assertIn('Station: -3',rows[0][3])

    def test_unknown_deviation_duration_is_not_reported_as_zero(self):
        from services.low_frequency_operation_report import report_blocks
        event,charts=fixture()
        event['threshold_analysis']['overall_performance']['State'][0]['thresholds']['49.90'].update(adverse_minutes=0,adverse_pct=None,unknown_deviation_minutes=2)
        blocks,_=report_blocks(event,{'operation_sections':['states']},charts)
        row=next(value[1][0] for kind,value in blocks if kind=='table')
        self.assertEqual(row[11],'Data not available')

    def test_prepared_event_workspace_restores_missing_scalar_deviation(self):
        doc={'event_id':'one','event_type':'low','data_points':[{'plant_id':'STATE_TEST','plant_name':'State','type':'State','series':{'timestamps':['2026-10-02T04:00:00','2026-10-02T04:01:00'],'frequency':[49.4,49.5],'actual':[30,50],'schedule':[10,20]}}]}
        with patch.object(routes,'lookup_unit_fuel_name',return_value=''),patch.object(routes,'build_source_status',return_value={}):
            from datetime import datetime
            response,error=routes.build_saved_event_response(None,'one',[],datetime(2026,10,2,4),datetime(2026,10,2,5),saved_document=doc,refresh_capacity=False)
        self.assertIsNone(error)
        self.assertEqual(response['rows'][0]['deviation'],25)
        self.assertEqual(response['rows'][0]['series']['deviation'],[20,30])

    def test_saved_message_context_uses_point_lookup_when_crms_is_offline(self):
        event,_=fixture()
        event['data_points']=[{'plant_id':'state','plant_name':'Bihar','type':'State','series':{'timestamps':['2026-10-02T04:01:00'],'frequency':[49.5],'deviation':[20]},'crms_messages':[{'timestamp':'2026-10-02T04:01:00','message_no':'M1','remarks':'Reduce drawal','issued_to':['Bihar']}]}]
        db=SimpleNamespace(map_collection=SimpleNamespace(find=lambda *args:[]))
        from services.frequency_event_reporting import timeline_aliases
        async def run():
            with patch.object(routes,'fetch_crms_frequency_messages',new=AsyncMock(side_effect=RuntimeError('offline'))):
                result=await routes._build_frequency_message_timeline(routes.FrequencyMessageTimelinePayload(event_id='one',ranges=[routes.FrequencyMessageRange(start_time=event['start_time'],end_time=event['end_time'])]),event_context=event,aliases_override=timeline_aliases(db,event),db=db)
            self.assertEqual(result['rows'][0]['deviation_mw'],20)
        asyncio.run(run())

    def test_frequency_only_event_can_capture_its_existing_system_plot(self):
        event,_=fixture()
        doc={**event,'updated_at':'revision','data_points':[{'plant_id':'SYSTEM_FREQUENCY','is_frequency':True,'series':{'timestamps':['2026-10-02T04:00:00','2026-10-02T04:01:00'],'frequency':[49.8,49.7]}}]}
        db=SimpleNamespace(map_collection=SimpleNamespace(find=lambda *args:[]))
        async def run():
            with patch.object(routes,'_saved_report_events',new=AsyncMock(return_value=(db,[event],[doc]))):
                response=await routes.export_saved_frequency_events(routes.FrequencySavedReportPayload(event_ids=['one'],format='html',operation_report=True))
                self.assertTrue(response['capture_required'])
                self.assertEqual(len(response['rows']),1)
                self.assertTrue(response['rows'][0]['is_frequency'])
                self.assertEqual(response['rows'][0]['series']['frequency'],[49.8,49.7])
        asyncio.run(run())


if __name__ == '__main__': unittest.main()
