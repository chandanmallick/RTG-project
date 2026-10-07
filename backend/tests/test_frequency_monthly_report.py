import asyncio
import io
import unittest
from unittest.mock import patch
from docx import Document
from test_frequency_threshold_analysis import fixture, USER, DB
from services import frequency_analysis_sessions as sessions
from services.frequency_monthly_report import build_blocks, performance_columns
from routes.frequency_analysis_routes import ResultPayload, export_result


class MonthlyReportTests(unittest.TestCase):
    def setUp(self):sessions._sessions.clear()

    def result(self):
        token=sessions._put(USER,fixture(),{'filename':'Report test','crms_enabled':False})['session_token']
        return sessions.analyse(token,USER,[('2026-10-03T17:00:00','2026-10-03T17:01:00'),('2026-10-03T17:02:00','2026-10-03T17:03:00')],db=DB)

    def test_template_exports_text_scope_and_annexures(self):
        result=self.result()
        payload=ResultPayload(session_token=result['session_token'],result_token=result['result_token'],layout='monthly-template',include_chronology=False,report_text={'executive_summary':'Reviewed by operator <safe>','adms_ufr_remarks':'UFR remarks entered by operator'})
        session,stored=sessions.get_result(payload.session_token,payload.result_token,USER)
        blocks=build_blocks(session,stored,payload)
        self.assertTrue(any('1440' in str(b) for b in blocks if b[0]=='text'))
        self.assertTrue(any(b[0]=='image' and 'Event 2' in b[1][0] for b in blocks))
        async def contents(response):
            return b''.join([chunk async for chunk in response.body_iterator])
        with patch('services.frequency_monthly_report.build_blocks',return_value=blocks):
            for fmt in ('docx','pdf'):
                payload.format=fmt;data=asyncio.run(contents(export_result(payload,USER)))
                if fmt=='docx':
                    document=Document(io.BytesIO(data));text='\n'.join(p.text for p in document.paragraphs)
                    self.assertIn('Reviewed by operator <safe>',text)
                    self.assertIn('UFR remarks entered by operator',text)
                    self.assertIn('Annexure 1',text);self.assertNotIn('6,791',text)
                    self.assertTrue(any(s.page_width>s.page_height for s in document.sections))
                    self.assertEqual(len(performance_columns(False)),15)
                    for table in document.tables:
                        header_width=sum(c.tcPr.tcW.w for c in table.rows[0]._tr.tc_lst)
                        body_width=sum(c.tcPr.tcW.w for c in table.rows[-1]._tr.tc_lst)
                        self.assertEqual(header_width,body_width)
                else:self.assertTrue(data.startswith(b'%PDF'))

    def test_owner_and_format_validation(self):
        result=self.result();payload=ResultPayload(session_token=result['session_token'],result_token=result['result_token'],layout='monthly-template',format='html')
        with self.assertRaises(ValueError):export_result(payload,USER)
        payload.format='pdf'
        with self.assertRaises(Exception) as caught:export_result(payload,{'employeeId':'other'})
        self.assertEqual(caught.exception.status_code,404)
