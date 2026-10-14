import math
import subprocess
import tempfile
import unittest
from datetime import date, datetime, time, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from routes import psp_routes
from services.curve_frequency_service import extract_curve_frequency_rows, load_curve_frequency_range


class CurveFrequencyTests(unittest.TestCase):
    def test_frequency_below_45_is_a_gap_and_45_is_valid(self):
        rows = self.samples()
        rows[0] = (rows[0][0],44.999)
        rows[1] = (rows[1][0],45)
        points,stats = extract_curve_frequency_rows(rows,'2026-10-05')
        self.assertIsNone(points[0]['frequency'])
        self.assertEqual(points[1]['frequency'],45)
        self.assertEqual(stats['missing_readings'],1)

    def samples(self):
        day = datetime(2026, 10, 5)
        return [(day + timedelta(seconds=index * 30), 50.0) for index in range(2880)]

    def test_exact_timeline_and_missing_readings(self):
        rows = self.samples()
        rows[0] = (date(2026, 10, 5), 49.90)
        rows[1] = (time(0, 0, 30), 50.05)
        rows[2] = (60 / 86400, None)
        rows[3] = (None, math.nan)
        points, stats = extract_curve_frequency_rows(rows, "2026-10-05")
        self.assertEqual(len(points), 2880)
        self.assertEqual(points[0], {"timestamp": "2026-10-05T00:00:00", "frequency": 49.90})
        self.assertEqual(points[-1]["timestamp"], "2026-10-05T23:59:30")
        self.assertIsNone(points[2]["frequency"])
        self.assertEqual(stats, {"missing_readings": 2, "derived_timestamps": 1})

    def test_wrong_date_and_short_range_rejected(self):
        with self.assertRaises(ValueError):
            extract_curve_frequency_rows(self.samples()[:-1], "2026-10-05")
        with self.assertRaises(ValueError):
            extract_curve_frequency_rows(self.samples(), "2026-10-06")

    def test_range_validation_and_partial_availability(self):
        for start, end in [("invalid", "2026-10-05"), ("2026-10-06", "2026-10-05"), ("2026-10-01", "2026-11-01")]:
            with self.assertRaises(ValueError):
                load_curve_frequency_range(start, end)
        point = {"timestamp": "2026-10-05T00:00:00", "frequency": 50}
        with patch.object(psp_routes, "get_psp_config_with_curve_defaults", return_value={}), patch.object(psp_routes, "read_curve_file_series", side_effect=[({"frequency": [point]}, {"available": True, "file": "curve.xlsm"}), ({}, {"available": False, "message": "Missing file"})]) as reader:
            result = load_curve_frequency_range("2026-10-05", "2026-10-06")
        self.assertTrue(result["success"])
        self.assertEqual(result["points"], [point])
        self.assertEqual(result["diagnostics"][0]["date"], "2026-10-06")
        self.assertEqual(reader.call_count, 2)
        self.assertTrue(reader.call_args.kwargs["frequency_only"])

    def test_existing_curve_reader_exact_cells_and_openpyxl_fallback(self):
        original_run = subprocess.run
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "curve_05102026.xlsm"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "30SEC"
            workbook.create_sheet("WRONG")["D8"] = 99
            sheet["C7"], sheet["D7"] = "Header", 999
            for index, (stamp, value) in enumerate(self.samples(), start=8):
                sheet.cell(index, 3, stamp)
                sheet.cell(index, 4, value)
            sheet["D8"] = 49.91
            sheet["D2887"] = 50.06
            sheet["D2888"] = 888
            workbook.save(path)
            workbook.close()
            config = {"curve_file_dir": directory, "curve_sheet_name": "WRONG"}
            for force_fallback in (False, True):
                def run(command, **kwargs):
                    if force_fallback:
                        command = [*command[:2], "import sys; sys.modules['python_calamine'] = None\n" + command[2]]
                    return original_run(command, **kwargs)
                with patch.object(psp_routes.subprocess, "run", side_effect=run):
                    result, meta = psp_routes.read_curve_file_series("2026-10-05", [], config, frequency_only=True)
                self.assertTrue(meta["available"], meta)
                self.assertEqual(len(result["frequency"]), 2880)
                self.assertEqual(result["frequency"][0]["frequency"], 49.91)
                self.assertEqual(result["frequency"][-1]["frequency"], 50.06)
                if force_fallback:
                    self.assertEqual(meta["reader"], "openpyxl")

    def test_calculate_curve_frequency_stats_all_bands(self):
        from services.curve_frequency_service import calculate_curve_frequency_stats
        points = [
            {"timestamp": "2026-10-05T00:00:00", "frequency": 49.40},  # < 49.5, < 49.7, < 49.9
            {"timestamp": "2026-10-05T00:00:30", "frequency": 49.65},  # < 49.7, < 49.9
            {"timestamp": "2026-10-05T00:01:00", "frequency": 49.85},  # < 49.9
            {"timestamp": "2026-10-05T00:01:30", "frequency": 49.95},  # Normal
            {"timestamp": "2026-10-05T00:02:00", "frequency": 50.00},  # Normal
            {"timestamp": "2026-10-05T00:02:30", "frequency": 50.08},  # > 50.05
            {"timestamp": "2026-10-05T00:03:00", "frequency": 50.15},  # > 50.05, > 50.10
            {"timestamp": "2026-10-05T00:03:30", "frequency": 50.25},  # > 50.05, > 50.10, > 50.20
            {"timestamp": "2026-10-05T00:04:00", "frequency": 50.35},  # > 50.05, > 50.10, > 50.20, > 50.30
            {"timestamp": "2026-10-05T00:04:30", "frequency": 50.55},  # > 50.05, > 50.10, > 50.20, > 50.30, > 50.50
            {"timestamp": "2026-10-05T00:05:00", "frequency": None},   # Missing
        ]
        stats = calculate_curve_frequency_stats(points, sample_seconds=30)
        self.assertEqual(stats["total_samples"], 11)
        self.assertEqual(stats["valid_samples"], 10)
        self.assertEqual(stats["missing_samples"], 1)
        self.assertEqual(stats["covered_minutes"], 5.0)
        self.assertEqual(stats["max_frequency"], 50.55)
        self.assertEqual(stats["max_timestamp"], "2026-10-05T00:04:30")
        self.assertEqual(stats["min_frequency"], 49.40)
        self.assertEqual(stats["min_timestamp"], "2026-10-05T00:00:00")

        by_id = {r["id"]: r for r in stats["rows"]}
        self.assertEqual(by_id["normal"]["count"], 2)
        self.assertEqual(by_id["normal"]["percent"], 20.0)
        self.assertEqual(by_id["lt49_9"]["count"], 3)
        self.assertEqual(by_id["lt49_7"]["count"], 2)
        self.assertEqual(by_id["lt49_5"]["count"], 1)
        self.assertEqual(by_id["gt50_05"]["count"], 5)
        self.assertEqual(by_id["gt50_10"]["count"], 4)
        self.assertEqual(by_id["gt50_20"]["count"], 3)
        self.assertEqual(by_id["gt50_30"]["count"], 2)
        self.assertEqual(by_id["gt50_50"]["count"], 1)

    def test_export_curve_frequency_excel_bytes(self):
        from openpyxl import load_workbook
        from services.curve_frequency_service import export_curve_frequency_excel_bytes
        points = [
            {"timestamp": "2026-10-05T00:00:00", "frequency": 49.80},
            {"timestamp": "2026-10-05T00:00:30", "frequency": 50.00},
            {"timestamp": "2026-10-05T00:01:00", "frequency": 50.20},
        ]
        buf = export_curve_frequency_excel_bytes("2026-10-05", "2026-10-05", points=points)
        wb = load_workbook(buf)
        self.assertIn("Frequency Statistics", wb.sheetnames)
        self.assertIn("30SEC Frequency Data", wb.sheetnames)
        ws1 = wb["Frequency Statistics"]
        self.assertIn("FREQUENCY OVERVIEW & STATISTICS", str(ws1["A1"].value))
        ws2 = wb["30SEC Frequency Data"]
        self.assertEqual(ws2["C2"].value, 49.8)
        self.assertEqual(ws2["C3"].value, 50.0)
        self.assertEqual(ws2["C4"].value, 50.2)


if __name__ == "__main__":
    unittest.main()

