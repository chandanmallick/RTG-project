import io
import unittest
from datetime import datetime,timedelta
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
import numpy as np
from openpyxl import Workbook,load_workbook
from fastapi import HTTPException
from services import frequency_analysis_sessions as sessions
from services.frequency_threshold_analysis import calculate,timeline,seconds,chart_points,merge_ranges
from routes import frequency_routes
from routes.frequency_analysis_routes import ResultPayload,export_result

USER={'employeeId':'test-owner'}
MAPPINGS=[{'plant_id':'state','plant_name':'Bihar','type':'State','is_state':True,'scada_key':'10001','scada_schedule_key':'10002','crms_utility_name':'BSPTCL'},
          {'plant_id':'gen','plant_name':'Unit','type':'ISGS','is_state':False,'scada_key':'20001','scada_schedule_key':'20002','crms_utility_name':'UNIT'},
          {'plant_id':'local','plant_name':'State Generator','type':'State','is_state':False,'scada_key':'20001','scada_schedule_key':'20002'}]
DB=SimpleNamespace(map_collection=SimpleNamespace(find=lambda *args,**kwargs:MAPPINGS))


def fixture():
    stamps=['2026-10-03T17:00:00','2026-10-03T17:00:30','2026-10-03T17:01:00','2026-10-03T17:02:00','2026-10-03T17:02:30']
    times,freq,cadence=timeline(stamps,[49.8,49.6,49.4,49.3,49.9])
    entities=[{'entity_id':'state::0','display_name':'Bihar','group':'State','deviation':np.array([10,-5,20,np.nan,0]),'point':MAPPINGS[0],'mapping':MAPPINGS[0]},
              {'entity_id':'gen::1','display_name':'Unit','group':'ISGS','deviation':np.array([-10,5,-20,-8,0]),'point':MAPPINGS[1],'mapping':MAPPINGS[1]}]
    return {'times':times,'frequency':freq,'cadence':cadence,'entities':entities,'warnings':[]}


def workbook():
    wb=Workbook();ws=wb.active;ws.title='Sheet1'
    ws.append(['Date','FREQUENCY','Bihar Actual','Bihar Schedule','Unit Actual','Unit Schedule'])
    ws.append(['DATE & TIME','04245232','10001','10002','20001','20002'])
    for at in range(180):ws.append([datetime(2026,10,3,17)+timedelta(seconds=30*at),49.4 if at%3==0 else 49.8,110,100,70,80])
    buffer=io.BytesIO();wb.save(buffer);return buffer.getvalue()


class ThresholdAnalysisTests(unittest.TestCase):
    def setUp(self):sessions._sessions.clear()

    def test_nested_duration_gaps_and_missing_deviation(self):
        result=calculate(fixture(),[('2026-10-03T17:00:00','2026-10-03T17:03:00')],[])
        summary=result['summary']
        self.assertEqual(summary['selected_minutes'],3)
        self.assertEqual(summary['covered_frequency_minutes'],2.5)
        self.assertEqual(summary['uncovered_minutes'],.5)
        self.assertEqual([summary['thresholds'][level]['frequency_minutes'] for level in ['49.90','49.70','49.50']],[2,1.5,1])
        self.assertEqual(summary['thresholds']['49.90']['occurrences'],2)
        self.assertEqual(summary['thresholds']['49.90']['longest_minutes'],1.5)
        self.assertEqual(summary['minimum_timestamp'],'2026-10-03T17:02:00')
        state=result['overall_performance']['State'][0]['thresholds']['49.90']
        self.assertIsNone(state['adverse_pct'])
        self.assertEqual(state['adverse_minutes'],1)
        self.assertEqual(state['unknown_deviation_minutes'],.5)
        self.assertAlmostEqual(state['average_od_ui_mw'],25/3,places=5)
        gen=result['overall_performance']['ISGS'][0]['thresholds']
        self.assertEqual(gen['49.90']['adverse_pct'],75)
        self.assertEqual(gen['49.50']['maximum_od_ui_mw'],-20)

    def test_overlap_boundaries_and_block_splitting(self):
        result=calculate(fixture(),[('2026-10-03T17:00:00','2026-10-03T17:02:00'),('2026-10-03T17:01:00','2026-10-03T17:03:00')],[])
        self.assertEqual(result['summary']['selected_minutes'],3)
        self.assertEqual(result['summary']['thresholds']['49.90']['frequency_minutes'],2)
        times,freq,cadence=timeline(['2026-10-03T17:14:50','2026-10-03T17:15:20','2026-10-03T17:15:50'],[49.8]*3)
        data=fixture();data.update(times=times,frequency=freq,cadence=cadence);data['entities']=data['entities'][:1];data['entities'][0]['deviation']=np.array([10,20,30])
        result=calculate(data,[('2026-10-03T17:14:50','2026-10-03T17:16:20')],[])
        rows=result['performance']['State']
        self.assertEqual(len(rows),2)
        self.assertAlmostEqual(rows[0]['thresholds']['49.90']['frequency_minutes'],10/60,places=5)
        self.assertAlmostEqual(rows[1]['thresholds']['49.90']['frequency_minutes'],80/60,places=5)

    def test_message_dedup_and_half_open_periods(self):
        messages=[{'entity_id':'state::0','timestamp':'2026-10-03T17:00:00','message_no':'M1'}]*2
        result=calculate(fixture(),[('2026-10-03T17:00:00','2026-10-03T17:03:00')],messages)
        self.assertEqual(result['overall_performance']['State'][0]['message_count'],1)
        unknown=calculate(fixture(),[('2026-10-03T17:00:00','2026-10-03T17:03:00')],messages,False)
        self.assertIsNone(unknown['performance']['State'][0]['message_count'])
        self.assertEqual(sessions.daily_ranges('2026-10-03','2026-10-03',[{'start':'17:00','end':'19:00'},{'start':'18:00','end':'21:00'}]),[('2026-10-03T17:00:00','2026-10-03T21:00:00')])
        self.assertEqual(sessions.daily_ranges('2026-10-03','2026-10-03',[{'start':'23:00','end':'24:00'}])[0][1],'2026-10-04T00:00:00')

    def test_upload_parser_once_cached_crms_and_owner_scope(self):
        with patch.object(frequency_routes,'parse_scada_file',wraps=frequency_routes.parse_scada_file) as parser:
            metadata=sessions.upload_session(workbook(),'base.xlsx',USER,db=DB)
            self.assertEqual(parser.call_count,1)
        self.assertEqual(metadata['row_count'],180)
        self.assertEqual(metadata['sampling_seconds'],30)
        self.assertEqual(metadata['entity_counts']['State'],1)
        self.assertEqual(metadata['entity_counts']['ISGS'],1)
        token=metadata['session_token']
        messages=[{'timestamp':'2026-10-03 17:00:00','issued_to':['BSPTCL','UNIT'],'message_no':'M1','remarks':'Reduce adverse deviation'}, {'timestamp':'2026-10-03 18:00:00','issued_to':['BSPTCL'],'message_no':'END'}]
        with patch.object(frequency_routes,'fetch_crms_frequency_messages',new=AsyncMock(return_value=(messages,0))) as fetch,patch.object(frequency_routes,'_nearest_saved_event_values',side_effect=AssertionError('No DB measurements')),patch.object(frequency_routes,'_raw_timeline_values',side_effect=AssertionError('No raw Mongo data')):
            result=sessions.analyse(token,USER,[('2026-10-03T17:00:00','2026-10-03T18:00:00')],db=DB)
            again=sessions.analyse(token,USER,[('2026-10-03T17:00:00','2026-10-03T18:00:00')],db=DB)
            smaller=sessions.analyse(token,USER,[('2026-10-03T17:15:00','2026-10-03T17:45:00')],db=DB)
            self.assertEqual(fetch.call_count,1)
        self.assertEqual(result['result_token'],again['result_token'])
        self.assertNotIn('performance',result)
        self.assertNotIn('series',result)
        self.assertEqual(result['chronology_count'],2)
        self.assertEqual(sessions.result_page(token,result['result_token'],USER,'Chronology')['total'],2)
        with self.assertRaises(HTTPException):sessions.get_session(token,{'employeeId':'other'})
        with patch.object(sessions.time,'monotonic',return_value=sessions._sessions[token]['accessed']+sessions.TTL+1):
            with self.assertRaises(HTTPException):sessions.get_session(token,USER)

    def test_nested_export_workbook_and_existing_word_engine(self):
        metadata=sessions._put(USER,fixture(),{'filename':'Long Period'})
        with patch.object(frequency_routes,'fetch_crms_frequency_messages',new=AsyncMock(return_value=([],0))):
            result=sessions.analyse(metadata['session_token'],USER,[('2026-10-03T17:00:00','2026-10-03T17:03:00')],db=DB)
        payload=ResultPayload(session_token=result['session_token'],result_token=result['result_token'])
        response=export_result(payload,USER)
        async def body():
            chunks=[]
            async for chunk in response.body_iterator:chunks.append(chunk)
            return b''.join(chunks)
        import asyncio
        wb=load_workbook(io.BytesIO(asyncio.run(body())))
        self.assertIn('Overall Statistics',wb.sheetnames)
        self.assertIn('State Overall',wb.sheetnames)
        self.assertIn('State Performance',wb.sheetnames)
        self.assertEqual(wb['State Performance'].freeze_panes,'F3')
        self.assertEqual(wb['State Performance']['F1'].value,'Freq <49.90')
        self.assertEqual(wb['State Performance']['K1'].value,'Freq <49.70')
        self.assertEqual(wb['State Performance']['P1'].value,'Freq <49.50')
        payload.format='docx'
        response=export_result(payload,USER)
        from docx import Document
        document=Document(io.BytesIO(asyncio.run(body())))
        self.assertTrue(any(len(table.columns)==10 for table in document.tables))

    def test_envelope_keeps_minimum_and_od_peak(self):
        data=fixture()
        chart=chart_points(data,[('2026-10-03T17:00:00','2026-10-03T17:03:00')],'state::0',limit=10)
        self.assertEqual(min(point['frequency'] for point in chart['points'] if point['frequency'] is not None),49.3)
        self.assertEqual(max(point['od_mw'] for point in chart['points'] if point['od_mw'] is not None),20)

    def test_invalid_data_and_empty_denominator(self):
        with self.assertRaises(ValueError):timeline(['2026-10-03T17:00:00']*2,[49.8]*2)
        data=fixture();data['frequency']=np.full(5,49.9)
        result=calculate(data,[('2026-10-03T17:00:00','2026-10-03T17:03:00')],[])
        self.assertEqual(result['summary']['low_frequency_events'],0)
        self.assertIsNone(result['overall_performance']['State'][0]['thresholds']['49.90']['adverse_pct'])

    def test_different_sampling_intervals_and_timezone_alignment(self):
        times,freq,cadence=timeline(['2026-10-03T11:30:00Z','2026-10-03T11:31:00Z'],[49.4,49.6])
        self.assertEqual(cadence,60)
        self.assertEqual(times[0],seconds('2026-10-03T17:00:00'))
        data=fixture();data['times']=times;data['frequency']=freq;data['cadence']=cadence
        data['entities']=data['entities'][:1];data['entities'][0]['deviation']=np.array([10,20])
        result=calculate(data,[('2026-10-03T17:00:00','2026-10-03T17:02:00')],[])
        self.assertEqual(result['summary']['thresholds']['49.90']['frequency_minutes'],2)
        self.assertEqual(result['summary']['thresholds']['49.50']['frequency_minutes'],1)

    def test_crms_failure_and_chart_do_not_trigger_reparse_or_source_fallback(self):
        metadata=sessions.upload_session(workbook(),'base.xlsx',USER,db=DB)
        with patch.object(frequency_routes,'fetch_crms_frequency_messages',new=AsyncMock(side_effect=RuntimeError('offline'))),patch.object(frequency_routes,'parse_scada_file',side_effect=AssertionError('Must reuse parsed session')):
            result=sessions.analyse(metadata['session_token'],USER,[('2026-10-03T17:00:00','2026-10-03T18:00:00')],db=DB)
            sessions.result_chart(result['session_token'],result['result_token'],USER,result['states'][0]['entity_id'])
            sessions.result_page(result['session_token'],result['result_token'],USER,'State')
        self.assertFalse(result['messages_complete'])
        self.assertIsNone(result['overall_performance']['State'][0]['message_count'])
        self.assertTrue(any('CRMS' in warning for warning in result['warnings']))

    def test_large_envelope_is_bounded_and_preserves_extrema_and_gaps(self):
        count=100000
        data=fixture();data['times']=seconds('2026-10-03T00:00:00')+np.arange(count)*30
        data['frequency']=np.full(count,50.0);data['frequency'][43000]=49.3
        data['frequency'][40000:40003]=np.nan
        data['entities']=data['entities'][:1];data['entities'][0]['deviation']=np.zeros(count);data['entities'][0]['deviation'][88000]=999
        chart=chart_points(data,[('2026-10-03T00:00:00',sessions.iso(data['times'][-1]+30))],'state::0')
        self.assertLessEqual(len(chart['points']),2400)
        self.assertEqual(min(p['frequency'] for p in chart['points'] if p['frequency'] is not None),49.3)
        self.assertEqual(max(p['od_mw'] for p in chart['points'] if p['od_mw'] is not None),999)
        self.assertTrue(any(p['frequency'] is None for p in chart['points'][:-1]))

if __name__=='__main__':unittest.main()
