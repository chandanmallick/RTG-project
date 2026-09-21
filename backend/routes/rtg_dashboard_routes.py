import io
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
from fastapi import APIRouter, Body, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from typing import Optional

from services.db_handler import MongoService
from routes.dso_report_routes import (
    CURRENT_CRMS_GENERATOR_OUTAGES_URL,
    fetch_current_crms_generator_outages,
    outage_fuel_group,
)
from routes.pipeline_routes import build_unit_lookup, compact_key, make_legacy_session, normalize_text, to_float

from services.rtg_dashboard_service import (
    RTGDashboardService
)

router = APIRouter(
    prefix="/api/rtg-dashboard",
    tags=["RTG Dashboard"]
)

CRMS_OUTAGE_CACHE_COLLECTION = "rtg_current_crms_outages"
CRMS_TRANSMISSION_CACHE_COLLECTION = "rtg_current_crms_transmission_outages"
ISGS_SCHEDULE_CHECK_COLLECTION = "rtg_isgs_schedule_checks"
ISGS_SCHEDULE_CHECK_CONFIG_ID = "ISGS_SCHEDULE_RECONCILIATION"


def _number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _serialise_schedule_check(document):
    if not document:
        return None
    result = dict(document)
    result.pop("_id", None)
    for key in ("checked_at", "rtg_snapshot_time"):
        if isinstance(result.get(key), datetime):
            result[key] = result[key].isoformat()
    return result


def get_isgs_schedule_check_config(db=None):
    db = db or MongoService()
    stored = db.pipeline_config_collection.find_one({"config_type": ISGS_SCHEDULE_CHECK_CONFIG_ID}) or {}
    interval = int(stored.get("interval_minutes", 15))
    if interval not in (0, 15, 30):
        interval = 15
    return {
        "interval_minutes": interval,
        "tolerance_mw": max(0.0, float(stored.get("tolerance_mw", 1.0))),
    }


def run_isgs_schedule_check(force_rtg_refresh=False):
    """Compare the current IST block's WBES and RTG schedules for mapped ISGS rows."""
    from concurrent.futures import ThreadPoolExecutor
    from routes.frequency_routes import fetch_rtg_schedule_raw, fetch_wbes_schedule_raw, normalize_wbes_identifier

    db = MongoService()
    config = get_isgs_schedule_check_config(db)
    now = datetime.now(ZoneInfo("Asia/Kolkata"))
    date_iso = now.date().isoformat()
    date_dmy = now.strftime("%d-%m-%Y")
    block_index = min(95, (now.hour * 60 + now.minute) // 15)

    diagnostics = []
    snapshot = db.rtg_dashboard_collection.find_one(
        {"snapshot_date": date_iso}, sort=[("snapshot_time", -1)]
    ) or db.rtg_dashboard_collection.find_one({}, sort=[("snapshot_time", -1)])
    rtg_by_id = {
        str(row.get("plant_id") or "").strip(): row
        for row in (snapshot or {}).get("data", [])
        if row.get("plant_id")
    }

    # utility_type distinguishes true inter-state stations from legacy rows whose
    # report bucket was once labelled ISGS despite being state-sector stations.
    mappings = list(db.map_collection.find({
        "$or": [
            {"utility_type": {"$regex": "^ISGS$", "$options": "i"}},
            {"type": {"$regex": "^ISGS$", "$options": "i"}, "wbes_name": {"$nin": [None, ""]}},
        ]
    }, {"_id": 0}))
    acronyms = sorted({
        normalize_wbes_identifier(row.get("wbes_name") or row.get("wbes_acronym"))
        for row in mappings
        if normalize_wbes_identifier(row.get("wbes_name") or row.get("wbes_acronym"))
    })
    wbes_raw = fetch_wbes_schedule_raw(date_dmy, acronyms, diagnostics, force_refresh=True)
    wbes_by_name = {
        normalize_wbes_identifier(item.get("Acronym")): item
        for item in wbes_raw
        if normalize_wbes_identifier(item.get("Acronym"))
    }
    plant_ids = sorted({str(row.get("rtg_plant_id") or row.get("plant_id") or "").strip() for row in mappings if row.get("rtg_plant_id") or row.get("plant_id")})

    def load_rtg_schedule(plant_id):
        fresh = fetch_rtg_schedule_raw(date_iso, plant_id, force_refresh=force_rtg_refresh)
        if fresh:
            return plant_id, fresh
        # A transient source failure should be identified, but a cached series
        # remains preferable to the dashboard aggregate (which is not block-wise).
        cached = fetch_rtg_schedule_raw(date_iso, plant_id, force_refresh=False)
        return plant_id, cached

    with ThreadPoolExecutor(max_workers=min(12, max(1, len(plant_ids)))) as executor:
        rtg_schedule_by_id = dict(executor.map(load_rtg_schedule, plant_ids))

    rows = []
    tolerance = config["tolerance_mw"]
    for mapping in mappings:
        plant_id = str(mapping.get("rtg_plant_id") or mapping.get("plant_id") or "").strip()
        acronym = normalize_wbes_identifier(mapping.get("wbes_name") or mapping.get("wbes_acronym"))
        plant_name = str(mapping.get("plant_name") or mapping.get("STAGE_NAME") or acronym or plant_id).strip()
        rtg_row = rtg_by_id.get(plant_id)
        rtg_series = (rtg_schedule_by_id.get(plant_id) or {}).get("schedule") or []
        wbes_row = wbes_by_name.get(acronym)
        schedule = ((wbes_row or {}).get("NetScheduleSummary") or {}).get("TotalNetSchdAmount") or []
        raw_wbes_value = _number(schedule[block_index]) if block_index < len(schedule) else None
        # WBES represents generator injection with a negative sign, while RTG
        # stores generation schedules as positive MW.
        wbes_value = -raw_wbes_value if raw_wbes_value is not None else None
        rtg_value = _number(rtg_series[block_index]) if block_index < len(rtg_series) else None
        difference = round(rtg_value - wbes_value, 3) if rtg_value is not None and wbes_value is not None else None
        if not acronym:
            status, remark = "UNMAPPED", "WBES name is not mapped"
        elif wbes_value is None:
            status, remark = "NO_WBES_DATA", "WBES schedule is unavailable for the current block"
        elif rtg_value is None:
            status, remark = "NO_RTG_DATA", "RTG schedule is unavailable"
        elif abs(difference) > tolerance:
            status, remark = "MISMATCH", f"Difference exceeds {tolerance:g} MW tolerance"
        else:
            status, remark = "MATCHED", "Schedule matched"
        rows.append({
            "plant_id": plant_id,
            "plant_name": plant_name,
            "stage_name": mapping.get("STAGE_NAME") or mapping.get("stage_name") or "",
            "wbes_name": acronym,
            "wbes_schedule_mw": wbes_value,
            "rtg_schedule_mw": rtg_value,
            "rtg_schedule_updated_at": (rtg_row or {}).get("schedule_last_updated"),
            "difference_mw": difference,
            "status": status,
            "remark": remark,
        })

    rows.sort(key=lambda row: (row["status"] == "MATCHED", -(abs(row["difference_mw"]) if row["difference_mw"] is not None else 10**9), row["plant_name"]))
    mismatches = [row for row in rows if row["status"] != "MATCHED"]
    document = {
        "checked_at": now.replace(tzinfo=None),
        "date": date_iso,
        "block": block_index + 1,
        "block_start": f"{(block_index * 15) // 60:02d}:{(block_index * 15) % 60:02d}",
        "interval_minutes": config["interval_minutes"],
        "tolerance_mw": tolerance,
        "total": len(rows),
        "matched": len(rows) - len(mismatches),
        "mismatch_count": len(mismatches),
        "rows": rows,
        "mismatches": mismatches,
        "diagnostics": diagnostics,
        "rtg_snapshot_time": (snapshot or {}).get("snapshot_time"),
    }
    db.db[ISGS_SCHEDULE_CHECK_COLLECTION].insert_one(document)
    return _serialise_schedule_check(document)


def run_scheduled_isgs_schedule_check():
    db = MongoService()
    config = get_isgs_schedule_check_config(db)
    interval = config["interval_minutes"]
    if interval == 0:
        return {"skipped": True, "reason": "Manual mode"}
    latest = db.db[ISGS_SCHEDULE_CHECK_COLLECTION].find_one({}, sort=[("checked_at", -1)])
    if latest and isinstance(latest.get("checked_at"), datetime):
        if datetime.now() - latest["checked_at"] < timedelta(minutes=interval - 1):
            return {"skipped": True, "reason": "Configured interval has not elapsed"}
    return run_isgs_schedule_check(force_rtg_refresh=False)


@router.get("/isgs-schedule-check")
async def get_isgs_schedule_check():
    db = MongoService()
    latest = db.db[ISGS_SCHEDULE_CHECK_COLLECTION].find_one({}, sort=[("checked_at", -1)])
    return {"success": True, "config": get_isgs_schedule_check_config(db), "report": _serialise_schedule_check(latest)}


@router.post("/isgs-schedule-check/run")
async def check_isgs_schedules_now():
    try:
        return {"success": True, "report": run_isgs_schedule_check(force_rtg_refresh=True)}
    except Exception as exc:
        return {"success": False, "message": str(exc)}


@router.put("/isgs-schedule-check/config")
async def update_isgs_schedule_check_config(payload: dict = Body(...)):
    interval = int(payload.get("interval_minutes") or 0)
    if interval not in (0, 15, 30):
        return {"success": False, "message": "Interval must be Manual, 15 minutes or 30 minutes."}
    tolerance = max(0.0, float(payload.get("tolerance_mw", 1.0)))
    db = MongoService()
    db.pipeline_config_collection.update_one(
        {"config_type": ISGS_SCHEDULE_CHECK_CONFIG_ID},
        {"$set": {"config_type": ISGS_SCHEDULE_CHECK_CONFIG_ID, "interval_minutes": interval, "tolerance_mw": tolerance, "updated_at": datetime.utcnow()}},
        upsert=True,
    )
    return {"success": True, "config": {"interval_minutes": interval, "tolerance_mw": tolerance}}


def _first_value(row, *keys):
    return next((row.get(key) for key in keys if row.get(key) not in (None, "")), "")


def _parse_crms_datetime(date_value, time_value=""):
    text = " ".join(filter(None, [str(date_value or "").strip(), str(time_value or "").strip()])).strip()
    if not text:
        return None
    normalized = text.replace("T", " ").replace("/", "-")
    for pattern in (
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d-%m-%Y %H:%M:%S",
        "%d-%m-%Y %H:%M", "%Y-%m-%d", "%d-%m-%Y",
    ):
        try:
            return datetime.strptime(normalized, pattern)
        except ValueError:
            pass
    return None


def _transmission_element_type(entity_value: str) -> str:
    entity = normalize_text(entity_value)
    if entity == "AC_TRANSMISSION_LINE_CIRCUIT":
        return "line"
    if entity == "LINE_REACTOR":
        return "line_reactor"
    if entity in {"BUS_REACTOR", "BUS REACTOR"}:
        return "bus_reactor"
    return ""


def _fetch_current_crms_transmission_rows():
    http = make_legacy_session()
    response = http.get(
        CURRENT_CRMS_GENERATOR_OUTAGES_URL,
        params={"entityFeatureName": ""},
        timeout=120,
    )
    response.raise_for_status()
    payload = response.json()
    rows = payload.get("data") if isinstance(payload, dict) else payload
    if isinstance(rows, dict):
        rows = rows.get("currentTrElementOutages") or rows.get("rows") or rows.get("data") or []
    return (rows if isinstance(rows, list) else []), response.url


def _current_crms_transmission_payload():
    db = MongoService()
    cache = db.db[CRMS_TRANSMISSION_CACHE_COLLECTION]
    try:
        raw_rows, source_url = _fetch_current_crms_transmission_rows()
        now = datetime.now()
        details = []
        type_summary = {
            "line": {"label": "Transmission Line", "total": 0, "over_15_days": 0},
            "reactor": {"label": "Reactors (Line + Bus)", "total": 0, "over_90_days": 0},
        }
        seen = set()
        target_category = "ERLDC/NLDC INITIATED - VOLTAGE REGULATION"

        for raw in raw_rows:
            category = str(_first_value(raw, "outageCategory", "OutageCategory", "OUTAGE_CATEGORY", "OUTAGE_TYPE", "TYPE") or "").strip()
            if normalize_text(category) != target_category:
                continue
            entity = normalize_text(_first_value(raw, "ENTITY_NAME", "entityName", "EntityName", "ENTITY_TYPE", "entityFeatureName"))
            element_type = _transmission_element_type(entity)
            if not element_type:
                continue
            name = _first_value(raw, "ELEMENT_NAME", "ELEMENTNAME", "elementName", "ElementName", "LINE_NAME", "lineName")
            if not name:
                continue
            # CRMS current-element data supplies the authoritative ISO openingTime.
            # It must win over legacy date/time fields, which are often blank.
            opening_time = _first_value(raw, "openingTime", "OPENING_TIME", "opening_time")
            outage_date = _first_value(raw, "OUTAGE_DATE", "outageDate", "outage_date")
            outage_time = _first_value(raw, "OUTAGE_TIME", "outageTime", "outage_time", "Tripped_Time", "trippedTime")
            outage_at = _parse_crms_datetime(opening_time) or _parse_crms_datetime(outage_date, outage_time)
            if not outage_at:
                outage_at = _parse_crms_datetime(_first_value(raw, "OUTAGE_DATE_TIME", "outageDateTime", "outage_datetime"))
            days_out = max(0, (now - outage_at).total_seconds() / 86400) if outage_at else 0
            identity = (compact_key(name), outage_at.isoformat() if outage_at else "")
            if identity in seen:
                continue
            seen.add(identity)
            age_threshold = 15 if element_type == "line" else 90
            details.append({
                "line_name": str(name).strip(),
                "entity_name": entity or "Transmission Element",
                "element_type": element_type,
                "owner": _first_value(raw, "OWNER", "owner", "Owners", "owners", "RequestingEntity", "requestingEntity"),
                "state_name": _first_value(raw, "STATE_NAME", "stateName", "STATE", "state"),
                "outage_category": category,
                "reason": _first_value(raw, "REASON", "reason", "OUT_REASON", "outReason"),
                "outage_at": outage_at.isoformat() if outage_at else "",
                "outage_date": str(opening_time or outage_date or ""),
                "outage_time": str(outage_time or ""),
                "days_out": round(days_out, 1),
                "over_15_days": days_out > 15,
                "age_threshold_days": age_threshold,
                "over_age_threshold": days_out > age_threshold,
                "expected_revival": " ".join(filter(None, [
                    str(_first_value(raw, "EXPECTED_REVIVAL_DATE", "expectedRevivalDate") or "").strip(),
                    str(_first_value(raw, "EXPECTED_REVIVAL_TIME", "expectedRevivalTime", "exprecteTimeOfRestoration", "expectedRestorationTime") or "").strip(),
                ])).strip(),
            })
            summary_key = "line" if element_type == "line" else "reactor"
            type_summary[summary_key]["total"] += 1
            if element_type == "line" and days_out > 15:
                type_summary[summary_key]["over_15_days"] += 1
            if element_type != "line" and days_out > 90:
                type_summary[summary_key]["over_90_days"] += 1

        details.sort(key=lambda item: (not item["over_age_threshold"], -item["days_out"], item["line_name"]))
        payload = {
            "success": True,
            "source": "CRMS current transmission-element outage API",
            "source_url": source_url,
            "filter": "ERLDC/NLDC Initiated - Voltage Regulation",
            "fetched_at": datetime.utcnow().isoformat(),
            "total": len(details),
            "over_15_days": sum(1 for item in details if item["over_15_days"]),
            "over_90_days": sum(1 for item in details if item["element_type"] != "line" and item["days_out"] > 90),
            "type_summary": type_summary,
            "rows": details,
            "cached": False,
        }
        cache.update_one({"_id": "latest"}, {"$set": payload}, upsert=True)
        return payload
    except Exception as exc:
        cached = cache.find_one({"_id": "latest"}, {"_id": 0})
        if cached:
            cached.update({"success": True, "cached": True, "warning": str(exc)})
            return cached
        return {"success": False, "message": str(exc), "total": 0, "over_15_days": 0, "rows": []}


def _current_crms_outage_payload(force_refresh=False):
    db = MongoService()
    cache = db.db[CRMS_OUTAGE_CACHE_COLLECTION]
    try:
        rows, source_url = fetch_current_crms_generator_outages()
        unit_lookup = build_unit_lookup(db)
        groups = {
            "thermal_state": [],
            "hydro_state": [],
            "thermal_central": [],
            "hydro_central": [],
        }
        seen = set()

        for raw in rows:
            name = _first_value(
                raw, "ELEMENT_NAME", "ELEMENTNAME", "elementName", "ElementName",
                "UNIT_NAME", "Unit_Name", "entityName",
            )
            if not name:
                continue
            unit = unit_lookup.get(normalize_text(name)) or unit_lookup.get(compact_key(name)) or {}
            fuel = _first_value(raw, "FUEL_TYPE", "FuelName", "FUEL", "fuel_type", "fuel") or \
                unit.get("fuel_type") or unit.get("FuelName") or unit.get("fuel")
            fuel_group = outage_fuel_group(fuel)
            if fuel_group not in {"THERMAL", "HYDRO"}:
                continue

            utility_type = str(unit.get("utility_type") or _first_value(raw, "UTILITY_TYPE", "utilityType") or "").strip()
            utility_key = normalize_text(utility_type)
            sector = "STATE" if utility_key in {"STATE", "STATE_IPP"} else "CENTRAL"
            capacity = to_float(
                _first_value(raw, "INSTALLED_CAPACITY", "installedCapacity", "Installed_Capacity", "CAPACITY", "unit_capacity")
                or unit.get("installed_capacity")
            )
            outage_type = str(_first_value(raw, "TYPE", "OUTAGE_TYPE", "outageCategory", "OutageCategory") or "").strip()
            identity = (compact_key(name), outage_type.upper(), capacity)
            if identity in seen:
                continue
            seen.add(identity)

            detail = {
                "unit_name": str(name).strip(),
                "station_name": unit.get("Generating_Station_Name") or _first_value(raw, "GENERATING_STATION_NAME", "stationName"),
                "state_name": unit.get("state_name") or _first_value(raw, "STATE_NAME", "stateName"),
                "utility_type": utility_type or "Unmapped",
                "fuel": fuel_group.title(),
                "sector": sector.title(),
                "outage_type": outage_type or "Outage",
                "capacity_mw": round(capacity, 3),
                "reason": _first_value(raw, "OUT_REASON", "REASON", "reason", "Reason"),
                "outage_date": _first_value(raw, "OUTAGE_DATE", "outageDate"),
                "outage_time": _first_value(raw, "OUTAGE_TIME", "outageTime", "Tripped_Time", "trippedTime"),
                "expected_revival": " ".join(filter(None, [
                    str(_first_value(raw, "EXPECTED_REVIVAL_DATE", "expectedRevivalDate") or "").strip(),
                    str(_first_value(raw, "EXPECTED_REVIVAL_TIME", "expectedRevivalTime", "exprecteTimeOfRestoration", "expectedRestorationTime") or "").strip(),
                ])).strip(),
                "mapped": bool(unit),
            }
            groups[f"{fuel_group.lower()}_{sector.lower()}"].append(detail)

        for details in groups.values():
            details.sort(key=lambda item: (-item["capacity_mw"], item["state_name"], item["unit_name"]))

        summary = {
            key: {
                "mw": round(sum(item["capacity_mw"] for item in details), 3),
                "units": len(details),
            }
            for key, details in groups.items()
        }
        payload = {
            "success": True,
            "source": "CRMS current transmission-element outage API filtered to GENERATING_UNIT",
            "source_url": source_url,
            "fetched_at": datetime.utcnow().isoformat(),
            "row_count": len(rows),
            "summary": summary,
            "groups": groups,
            "cached": False,
        }
        cache.update_one({"_id": "latest"}, {"$set": payload}, upsert=True)
        payload.pop("_id", None)
        return payload
    except Exception as exc:
        cached = cache.find_one({"_id": "latest"}, {"_id": 0})
        if cached:
            cached.update({"success": True, "cached": True, "warning": str(exc)})
            return cached
        return {"success": False, "message": str(exc), "summary": {}, "groups": {}}


@router.get("/current-crms-outages")
async def get_current_crms_outages(refresh: bool = False):
    return _current_crms_outage_payload(force_refresh=refresh)


@router.get("/current-crms-transmission-outages")
async def get_current_crms_transmission_outages(refresh: bool = False):
    return _current_crms_transmission_payload()


@router.get("/historical/options")
async def get_historical_options():
    return {"success": True, **RTGDashboardService.historical_options()}


@router.get("/historical/matrix")
async def get_historical_matrix(
    start_date: str,
    end_date: str,
    metrics: list[str] = Query(default=[]),
    selections: list[str] = Query(default=[]),
    plant_wise: bool = False,
    interval_minutes: int = 15,
):
    try:
        data = RTGDashboardService.fetch_historical_matrix(
            start_date, end_date, metrics, selections, plant_wise, interval_minutes
        )
        return {"success": True, **data}
    except Exception as exc:
        return {"success": False, "message": str(exc)}


@router.get("/historical/download")
async def download_historical_matrix(
    start_date: str,
    end_date: str,
    metrics: list[str] = Query(default=[]),
    selections: list[str] = Query(default=[]),
    plant_wise: bool = False,
    interval_minutes: int = 15,
):
    data = RTGDashboardService.fetch_historical_matrix(
        start_date, end_date, metrics, selections, plant_wise, interval_minutes
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RTG Historical Data"
    headers = ["Date", "Time", *[column["label"] for column in data["columns"]]]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0057B7")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in data["rows"]:
        sheet.append([
            row["date"], row["time"],
            *[row.get(column["key"]) for column in data["columns"]],
        ])
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 10
    for index in range(3, len(headers) + 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 26
    metadata = workbook.create_sheet("Selection")
    metadata.append(["Setting", "Value"])
    metadata.append(["Date range", f'{start_date} to {end_date}'])
    metadata.append(["Interval", f'{data["interval_minutes"]} minutes'])
    metadata.append(["View", "Plant-wise" if plant_wise else "Entity-wise"])
    metadata.append(["Entities", ", ".join(data["selections"])])
    metadata.append(["Data fields", ", ".join(data["metrics"])])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"RTG_Historical_{start_date}_to_{end_date}_{interval_minutes}min.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/summary")
async def get_summary():

    db = MongoService()

    latest = db.rtg_dashboard_collection.find_one(

        {},

        sort=[("snapshot_time",-1)]
    )

    if not latest:

        return {
            "success": False
        }

    data = latest["data"]

    installed = sum(
        x.get("installed_capacity",0)
        for x in data
    )

    actual = sum(
        x.get("actual_gen",0)
        for x in data
    )

    outage = sum(

        x.get("planned_outage",0)
        +
        x.get("forced_outage",0)
        +
        x.get("fuel_shortage",0)
        +
        x.get("rsd",0)
        +
        x.get("commercial_issues",0)

        for x in data
    )

    return {

        "success": True,

        "snapshot_time":
            latest["snapshot_time"],

        "total_plants":
            len(data),

        "installed_capacity":
            installed,

        "actual_generation":
            actual,

        "outage_capacity":
            outage,

        "data":
            data
    }


@router.post("/refresh")
async def refresh_dashboard():

    try:

        count = (
            RTGDashboardService
            .fetch_snapshot()
        )

        return {

            "success": True,

            "message":
                f"{count} records fetched"
        }

    except Exception as e:

        return {

            "success": False,

            "message":
                str(e)
        }
    


@router.get("/live")
async def get_live_dashboard():

    try:

        data = (
            RTGDashboardService
            .fetch_live_data()
        )

        return {
            "success": True,
            "data": data
        }

    except Exception as e:

        try:
            db = MongoService()
            latest = db.rtg_dashboard_collection.find_one(
                {},
                sort=[("snapshot_time", -1)]
            )
            if latest and "data" in latest:
                return {
                    "success": True,
                    "data": latest["data"],
                    "is_cached": True,
                    "snapshot_time": latest["snapshot_time"],
                    "message": f"Live fetch failed ({str(e)}). Loaded cached data from snapshot."
                }
        except Exception as db_err:
            pass

        return {
            "success": False,
            "message": str(e)
        }


@router.get("/trend/today")
async def get_today_trend(
    date_str: Optional[str] = None
):

    try:

        data = (
            RTGDashboardService
            .fetch_today_trend(date_str)
        )

        return {
            "success": True,
            "data": data
        }

    except Exception as e:

        return {
            "success": False,
            "message": str(e)
        }


@router.get("/trend/snapshot")
async def get_snapshot_trend(
    date_str: Optional[str] = None,
    historical_only: bool = False,
):

    print(
        f"RTG snapshot trend API date_str={date_str!r}",
        flush=True
    )

    try:

        data = (
            RTGDashboardService
            .fetch_snapshot_trend(date_str, historical_only=historical_only)
        )

        return {
            "success": True,
            **data
        }

    except Exception as e:

        return {
            "success": False,
            "message": str(e)
        }
