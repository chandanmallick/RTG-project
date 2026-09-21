from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import base64
import json
import os
from pathlib import Path
from time import monotonic
from typing import Any
from urllib.parse import quote

import requests
import urllib3
from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from services.db_handler import MongoService
from services.token_service import TokenService
from crew_legacy.admin_logic.notification_service import _runtime_env_value, replacement_mail_settings
from crew_legacy.api.mail_settings import _persist_backend_credentials


urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

router = APIRouter(prefix="/api/plant-deviation", tags=["Plant Deviation from I/C - MOP"])

GENERATOR_MASTER_URL = (
    "https://rtgapi.grid-india.in/sendData/generator/filtered_details/"
    "?region_name=ERLDC"
)
DEFAULT_CAP_ON_BAR_URL = "https://rtgapi.grid-india.in/sendData/cap-on-bar/"
DEFAULT_SCADA_URL = "https://rtgapi.grid-india.in/sendData/scada-data/"
STATE_UTILITY_TYPES = {"STATE", "STATE_IPP"}
REGIONAL_IPP_UTILITY_TYPES = {"IPP", "REGIONAL_IPP"}
ISGS_UTILITY_TYPES = {"ISGS"}
THERMAL_FUELS = {"THERMAL", "COAL"}
HYDRO_FUELS = {"HYDEL", "HYDRO", "HYDROELECTRIC"}
REPORT_CACHE: dict[str, tuple[float, dict]] = {}
CACHE_SECONDS = 300
EDIT_COLLECTION = "plant_deviation_mop_edits"
DAY_AHEAD_COLLECTION = "plant_deviation_day_ahead_reports"
ALL_INDIA_COLLECTION = "plant_deviation_all_india_reports"
DEFAULT_MOP_06_HRS_DIR = Path(
    r"\\10.3.95.200\HTTP-Access\Control_Room_Report\Reserve Margin Report_Intra state\MOP_06 HRS Report"
)
MOP_06_HRS_CACHE: dict[str, tuple[float, dict]] = {}
DEFAULT_ALL_INDIA_DIR = Path(__file__).resolve().parents[1] / "data" / "plant_deviation" / "all_india"
ALL_INDIA_REGIONS = ("NR", "SR", "WR")


class PlantDeviationEdit(BaseModel):
    report_date: str
    running_capacity_mw: float | None = None
    max_generation_1900_2400_mw: float | None = None
    max_generation_1900_2400_time: str | None = None
    generation_max_min_1900_2400: str | None = None
    reason_for_not_attaining_full_generation: str | None = None


class PlantDeviationBatchItem(BaseModel):
    plant_id: str
    running_capacity_mw: float | None = None
    max_generation_1900_2400_mw: float | None = None
    max_generation_1900_2400_time: str | None = None
    generation_max_min_1900_2400: str | None = None
    reason_for_not_attaining_full_generation: str | None = None


class PlantDeviationBatchEdit(BaseModel):
    report_date: str
    edits: list[PlantDeviationBatchItem]


class DayAheadCoalStockItem(BaseModel):
    source_row: int
    plant_name: str
    coal_stock_days_left: float | None = None
    daily_coal_requirement: float | None = None


class DayAheadReportSave(BaseModel):
    report_date: str
    rows: list[DayAheadCoalStockItem]


def _text(value: Any) -> str:
    return str(value or "").strip()


def _key(value: Any) -> str:
    return _text(value).upper().replace("-", "_").replace(" ", "_")


def _compact(value: Any) -> str:
    return "".join(character for character in _text(value).upper() if character.isalnum())


def _number(value: Any) -> float | None:
    try:
        if value is None or value == "":
            return None
        return round(float(value), 3)
    except (TypeError, ValueError):
        return None


def _parse_report_date(value: str | None) -> date:
    if not value:
        return date.today() - timedelta(days=1)
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="report_date must be in YYYY-MM-DD format.") from exc


def _mop_06_hrs_directory() -> Path:
    return Path(os.getenv("MOP_06_HRS_REPORT_DIR", str(DEFAULT_MOP_06_HRS_DIR)))


def _mop_06_hrs_file(selected_date: date) -> Path:
    return _mop_06_hrs_directory() / f"ER_Plantwise Deviation from IC_{selected_date.strftime('%d-%m-%Y')}.xlsx"


def _day_ahead_thermal_report(report_date: str | None = None, force_refresh: bool = False) -> dict:
    """Read the daily MOP workbook and retain Thermal stations with non-zero margin."""
    selected_date = _parse_report_date(report_date)
    source_file = _mop_06_hrs_file(selected_date)
    if not source_file.is_file():
        raise HTTPException(
            status_code=404,
            detail=f"Day-ahead MOP file not found: {source_file.name}",
        )

    modified_at = source_file.stat().st_mtime
    cache_key = selected_date.isoformat()
    cached = MOP_06_HRS_CACHE.get(cache_key)
    if not force_refresh and cached and cached[0] == modified_at:
        return cached[1]

    try:
        workbook = load_workbook(source_file, read_only=True, data_only=True)
    except PermissionError as exc:
        raise HTTPException(status_code=503, detail="The daily MOP workbook is currently locked or inaccessible.") from exc
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Unable to read daily MOP workbook: {exc}") from exc

    try:
        sheet_name = next(
            (name for name in workbook.sheetnames if _compact(name) == "ERTHERMAL"),
            None,
        )
        if not sheet_name:
            raise HTTPException(status_code=422, detail="ER_Thermal sheet was not found in the daily MOP workbook.")
        sheet = workbook[sheet_name]
        output = []
        category = "State"
        section = "STATE - NOT MAPPED"
        ignored_zero_rows = 0
        all_thermal_outage_capacity = 0.0
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=4, max_col=11, values_only=True),
            start=4,
        ):
            values = list(values) + [None] * (11 - len(values))
            station = _text(values[0])
            marker = _text(values[5])
            installed = _number(values[1])
            outage = _number(values[7])
            margin = _number(values[6])

            if marker.upper() == "REGIONAL IPP":
                category, section = "Regional IPP", "REGIONAL IPP"
                continue
            if marker.upper() == "ISGS":
                category, section = "ISGS", "ISGS"
                continue
            if station.upper() == "STATE ENTITIES GENERATION":
                category, section = "State", "STATE - NOT MAPPED"
                continue

            # Section headings have text in column A but no numeric capacity.
            if station and installed is None:
                if station.upper() not in {"GRAND TOTAL", "TOTAL"}:
                    section = station.upper()
                continue
            if not station or installed is None:
                continue
            all_thermal_outage_capacity += outage or 0
            if margin is None or abs(margin) < 1e-9:
                ignored_zero_rows += 1
                continue

            output.append({
                "serial": len(output) + 1,
                "source_row": row_number,
                "category": category,
                "section": section,
                "plant_name": station,
                "installed_capacity_mw": installed,
                "capacity_on_bar_mw": _number(values[2]),
                "max_generation_1900_2400_mw": _number(values[3]),
                "max_generation_1900_2400_time": _text(values[4]),
                "generation_max_min_1900_2400": _text(values[5]),
                "running_units_margin_mw": margin,
                "outage_capacity_mw": outage,
                "reason_for_not_attaining_full_generation": _text(values[8]),
                "loading_factor_pct": _number(values[9]),
                "expected_revival_time": _text(values[10]),
            })
    finally:
        workbook.close()

    saved_document = MongoService().db[DAY_AHEAD_COLLECTION].find_one(
        {"report_date": selected_date.isoformat()},
        {"_id": 0},
    ) or {}
    saved_rows = {
        (_compact(row.get("plant_name")), row.get("source_row")): row
        for row in saved_document.get("rows", [])
    }
    saved_rows_by_name = {
        _compact(row.get("plant_name")): row
        for row in saved_document.get("rows", [])
    }
    output.sort(key=lambda row: row.get("running_units_margin_mw") or 0, reverse=True)
    for serial, row in enumerate(output, 1):
        saved = saved_rows.get(
            (_compact(row["plant_name"]), row["source_row"]),
        ) or saved_rows_by_name.get(_compact(row["plant_name"]), {})
        row["serial"] = serial
        row["coal_stock_days_left"] = _number(saved.get("coal_stock_days_left"))
        row["daily_coal_requirement"] = _number(saved.get("daily_coal_requirement"))

    report = {
        "success": True,
        "report_name": "Day-ahead Thermal Margin Consolidation",
        "report_date": selected_date.isoformat(),
        "source_file": source_file.name,
        "source_directory": str(source_file.parent),
        "sheet_name": sheet_name,
        "filter": "Margin on running units != 0 MW",
        "sort": "Margin on running units (descending)",
        "saved_at": saved_document.get("saved_at"),
        "rows": output,
        "summary": {
            "thermal_margin_stations": len(output),
            "zero_margin_stations_excluded": ignored_zero_rows,
            "total_margin_mw": round(sum(row["running_units_margin_mw"] for row in output), 3),
            "total_outage_capacity_mw": round(sum((row["outage_capacity_mw"] or 0) for row in output), 3),
            "all_thermal_outage_capacity_mw": round(all_thermal_outage_capacity, 3),
            "total_running_capacity_mw": round(sum((row["capacity_on_bar_mw"] or 0) for row in output), 3),
            "total_installed_capacity_mw": round(sum((row["installed_capacity_mw"] or 0) for row in output), 3),
        },
    }
    MOP_06_HRS_CACHE[cache_key] = (modified_at, report)
    return report


def _all_india_report_directory(selected_date: date) -> Path:
    root = Path(os.getenv("PLANT_DEVIATION_ALL_INDIA_DIR", str(DEFAULT_ALL_INDIA_DIR)))
    return root / f"{selected_date:%Y}" / f"{selected_date:%m}"


def _all_india_attachment_name(selected_date: date) -> str:
    return f"Plantwise Deviation from IC_{selected_date:%d-%m-%Y}_consolidated.xlsm"


def _all_india_subject(selected_date: date) -> str:
    return f"Consolidated details for thermal and hydro stations_{selected_date:%d.%m.%y}"


def _graph_access_token() -> tuple[str, str]:
    tenant_id = _runtime_env_value("PLANT_REPORT_GRAPH_TENANT_ID")
    client_id = _runtime_env_value("PLANT_REPORT_GRAPH_CLIENT_ID")
    auth_mode = _runtime_env_value("PLANT_REPORT_GRAPH_AUTH_MODE").strip().lower()
    auth_mode = auth_mode if auth_mode in {"application", "delegated"} else "delegated"

    if auth_mode == "delegated":
        refresh_token = _runtime_env_value("PLANT_REPORT_GRAPH_REFRESH_TOKEN")
        if not all((tenant_id, client_id, refresh_token)):
            raise HTTPException(
                status_code=409,
                detail=(
                    "Connect your Microsoft mailbox in Admin > Mail & 2FA Settings > "
                    "Consolidated Report Inbox."
                ),
            )
        try:
            response = requests.post(
                f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token",
                data={
                    "grant_type": "refresh_token",
                    "client_id": client_id,
                    "refresh_token": refresh_token,
                    "scope": (
                        "openid profile offline_access "
                        "https://graph.microsoft.com/Mail.Read"
                    ),
                },
                timeout=25,
            )
        except requests.RequestException as exc:
            raise HTTPException(status_code=502, detail=f"Microsoft delegated sign-in refresh failed: {exc}") from exc
        payload = response.json() if response.content else {}
        token = payload.get("access_token")
        if response.status_code != 200 or not token:
            error = payload.get("error_description") or payload.get("error") or "delegated authentication failed"
            raise HTTPException(
                status_code=401,
                detail=(
                    "The connected Microsoft mailbox session has expired or was revoked. "
                    "Reconnect it in Admin > Mail & 2FA Settings > Consolidated Report Inbox. "
                    f"Microsoft response: {str(error)[:240]}"
                ),
            )
        rotated_refresh_token = str(payload.get("refresh_token") or "").strip()
        if rotated_refresh_token and rotated_refresh_token != refresh_token:
            _persist_backend_credentials(
                {"PLANT_REPORT_GRAPH_REFRESH_TOKEN": rotated_refresh_token},
                "Rotated delegated Microsoft Graph report-mailbox token",
            )
        return token, "delegated"

    client_secret = _runtime_env_value("PLANT_REPORT_GRAPH_CLIENT_SECRET")
    if not all((tenant_id, client_id, client_secret)):
        raise HTTPException(status_code=409, detail="Configure the Consolidated Thermal/Hydro Report Tenant ID, Application ID and Client Secret in Mail Settings.")
    try:
        response = requests.post(
            f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token",
            data={
                "grant_type": "client_credentials",
                "client_id": client_id,
                "client_secret": client_secret,
                "scope": "https://graph.microsoft.com/.default",
            },
            timeout=25,
        )
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"Microsoft Graph authentication request failed: {exc}") from exc
    payload = response.json() if response.content else {}
    token = payload.get("access_token")
    if response.status_code != 200 or not token:
        error = payload.get("error_description") or payload.get("error") or "authentication failed"
        if "AADSTS7000215" in str(error):
            raise HTTPException(
                status_code=502,
                detail=(
                    "Invalid Microsoft Graph client secret. Open Microsoft Entra "
                    "ID > App registrations > Certificates & secrets and paste the "
                    "secret Value (not the Secret ID) in Admin > Mail Settings > "
                    "Consolidated Report Inbox, then save it again."
                ),
            )
        raise HTTPException(status_code=502, detail=f"Microsoft Graph authentication failed: {str(error)[:300]}")
    try:
        encoded_claims = token.split(".")[1]
        encoded_claims += "=" * (-len(encoded_claims) % 4)
        claims = json.loads(base64.urlsafe_b64decode(encoded_claims).decode("utf-8"))
        application_roles = set(claims.get("roles") or [])
        if not application_roles.intersection({"Mail.Read", "Mail.ReadWrite"}):
            raise HTTPException(
                status_code=403,
                detail=(
                    "Microsoft Graph authentication succeeded, but this application "
                    "token has no Mail.Read application role. In Microsoft Entra ID, "
                    "add Microsoft Graph > Application permissions > Mail.Read and "
                    "select Grant admin consent. Delegated Mail.Read is not sufficient."
                ),
            )
    except HTTPException:
        raise
    except (IndexError, ValueError, TypeError, UnicodeDecodeError):
        # If Microsoft changes the access-token format, allow Graph itself to
        # make the authorization decision instead of rejecting a valid token.
        pass
    return token, "application"


def _fetch_all_india_attachment(selected_date: date) -> tuple[Path, dict]:
    inbox_settings = replacement_mail_settings()
    if not inbox_settings.get("plantReportInboxEnabled"):
        raise HTTPException(status_code=409, detail="Enable the Consolidated Thermal/Hydro Report inbox in Mail Settings.")
    mailbox = _text(inbox_settings.get("reportMailbox"))
    if not mailbox:
        raise HTTPException(status_code=409, detail="Configure the Microsoft 365 report mailbox in Mail Settings.")
    token, auth_mode = _graph_access_token()
    headers = {"Authorization": f"Bearer {token}"}
    subject = _all_india_subject(selected_date)
    expected_attachment = _all_india_attachment_name(selected_date)
    mailbox_path = quote(mailbox, safe="")
    mailbox_resource = "me" if auth_mode == "delegated" else f"users/{mailbox_path}"
    try:
        message_response = requests.get(
            f"https://graph.microsoft.com/v1.0/{mailbox_resource}/mailFolders/inbox/messages",
            headers=headers,
            params={
                "$filter": f"subject eq '{subject.replace(chr(39), chr(39) * 2)}' and hasAttachments eq true",
                "$select": "id,subject,receivedDateTime,from,hasAttachments",
                "$top": "50",
            },
            timeout=35,
        )
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"Microsoft Graph inbox request failed: {exc}") from exc
    if message_response.status_code != 200:
        try:
            graph_error = (message_response.json().get("error") or {})
        except (ValueError, AttributeError):
            graph_error = {}
        message = graph_error.get("message") or "Mailbox access failed"
        if message_response.status_code == 403:
            message = (
                "Mailbox authorization was denied. Reconnect your own mailbox and consent "
                "to delegated Mail.Read in Mail Settings"
                if auth_mode == "delegated"
                else (
                    "Mailbox authorization was denied. Confirm that Mail.Read is an "
                    "Application permission with admin consent. If it is already granted, "
                    "ask the Exchange administrator to verify the application's mailbox "
                    "scope includes the configured report mailbox"
                )
            )
        raise HTTPException(
            status_code=502,
            detail=f"Microsoft Graph mailbox request failed (HTTP {message_response.status_code}): {str(message)[:400]}.",
        )
    messages = message_response.json().get("value") or []
    messages.sort(key=lambda item: item.get("receivedDateTime") or "", reverse=True)
    for message in messages:
        message_id = message.get("id")
        if not message_id:
            continue
        attachments_response = requests.get(
            f"https://graph.microsoft.com/v1.0/{mailbox_resource}/messages/{quote(message_id, safe='')}/attachments",
            headers=headers,
            params={"$select": "id,name,size,isInline,contentType"},
            timeout=35,
        )
        if attachments_response.status_code != 200:
            continue
        attachment = next(
            (
                item for item in attachments_response.json().get("value") or []
                if _text(item.get("name")).casefold() == expected_attachment.casefold()
                and not item.get("isInline")
            ),
            None,
        )
        if not attachment:
            continue
        attachment_id = attachment.get("id")
        content_response = requests.get(
            f"https://graph.microsoft.com/v1.0/{mailbox_resource}/messages/{quote(message_id, safe='')}/attachments/{quote(attachment_id, safe='')}/$value",
            headers=headers,
            timeout=90,
        )
        if content_response.status_code != 200:
            raise HTTPException(status_code=502, detail=f"Microsoft Graph could not download the report attachment (HTTP {content_response.status_code}).")
        destination_directory = _all_india_report_directory(selected_date)
        destination_directory.mkdir(parents=True, exist_ok=True)
        destination = destination_directory / expected_attachment
        destination.write_bytes(content_response.content)
        return destination, {
            "mailbox": mailbox,
            "subject": message.get("subject") or subject,
            "received_at": message.get("receivedDateTime"),
            "message_id": message_id,
            "attachment_id": attachment_id,
        }
    raise HTTPException(
        status_code=404,
        detail=f"No Inbox message with subject '{subject}' and attachment '{expected_attachment}' was found.",
    )


def _parse_all_india_workbook(source_file: Path, selected_date: date, mail: dict | None = None) -> dict:
    try:
        workbook = load_workbook(source_file, read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Unable to read the All India workbook: {exc}") from exc
    try:
        missing = [name for name in ALL_INDIA_REGIONS if name not in workbook.sheetnames]
        if missing:
            raise HTTPException(status_code=422, detail=f"Required sheet(s) missing: {', '.join(missing)}")
        rows = []
        zero_margin_excluded = 0
        all_outage = 0.0
        for region in ALL_INDIA_REGIONS:
            sheet = workbook[region]
            section = "NOT MAPPED"
            for row_number, values in enumerate(sheet.iter_rows(min_row=1, max_col=11, values_only=True), start=1):
                values = list(values) + [None] * (11 - len(values))
                station = _text(values[0])
                installed = _number(values[1])
                if station and installed is None:
                    normalized = _compact(station)
                    if normalized not in {"STATIONCONSTITUENTS", "GRANDTOTAL", "TOTAL", "ER", "NR", "SR", "WR"}:
                        section = station.upper()
                    continue
                if not station or installed is None or station.upper() in {"TOTAL", "GRAND TOTAL"}:
                    continue
                margin = _number(values[6])
                outage = _number(values[7])
                all_outage += outage or 0
                if margin is None or abs(margin) < 1e-9:
                    zero_margin_excluded += 1
                    continue
                rows.append({
                    "serial": 0,
                    "source_row": row_number,
                    "region": region,
                    "section": section,
                    "plant_name": station,
                    "installed_capacity_mw": installed,
                    "capacity_on_bar_mw": _number(values[2]),
                    "max_generation_1900_2400_mw": _number(values[3]),
                    "max_generation_1900_2400_time": _text(values[4]),
                    "generation_max_min_1900_2400": _text(values[5]),
                    "running_units_margin_mw": margin,
                    "outage_capacity_mw": outage,
                    "reason_for_not_attaining_full_generation": _text(values[8]),
                    "loading_factor_pct": _number(values[9]),
                    "expected_revival_time": _text(values[10]),
                })
    finally:
        workbook.close()
    rows.sort(key=lambda row: row["running_units_margin_mw"], reverse=True)
    for serial, row in enumerate(rows, 1):
        row["serial"] = serial
    report = {
        "success": True,
        "report_name": "All India Partial Outage - Margin Consolidation",
        "report_date": selected_date.isoformat(),
        "source_file": source_file.name,
        "source_sheets": list(ALL_INDIA_REGIONS),
        "filter": "Margin on running units != 0 MW",
        "sort": "Margin on running units (descending)",
        "mail": mail or {},
        "rows": rows,
        "summary": {
            "stations": len(rows),
            "zero_margin_stations_excluded": zero_margin_excluded,
            "total_margin_mw": round(sum(row["running_units_margin_mw"] for row in rows), 3),
            "total_outage_capacity_mw": round(sum((row["outage_capacity_mw"] or 0) for row in rows), 3),
            "all_outage_capacity_mw": round(all_outage, 3),
            "regions": {
                region: len([row for row in rows if row["region"] == region])
                for region in ALL_INDIA_REGIONS
            },
        },
    }
    MongoService().db[ALL_INDIA_COLLECTION].update_one(
        {"report_date": selected_date.isoformat()},
        {"$set": {**report, "updated_at": datetime.now(timezone.utc)}},
        upsert=True,
    )
    return report


def _all_india_report(report_date: str | None = None, fetch_mail: bool = False) -> dict:
    selected_date = _parse_report_date(report_date)
    source_file = _all_india_report_directory(selected_date) / _all_india_attachment_name(selected_date)
    mail = None
    if fetch_mail or not source_file.is_file():
        source_file, mail = _fetch_all_india_attachment(selected_date)
    return _parse_all_india_workbook(source_file, selected_date, mail)


def sync_er_day_ahead_report(report_date: str | None = None) -> dict:
    """Scheduled ER network-file ingestion; keeps any manually entered coal values."""
    selected_date = _parse_report_date(report_date or date.today().isoformat())
    report = _day_ahead_thermal_report(selected_date.isoformat(), force_refresh=True)
    synced_at = datetime.now(timezone.utc)
    MongoService().db[DAY_AHEAD_COLLECTION].update_one(
        {"report_date": selected_date.isoformat()},
        {"$set": {
            "report_date": selected_date.isoformat(),
            "source_file": report["source_file"],
            "filter": report["filter"],
            "sort": report["sort"],
            "summary": report["summary"],
            "rows": report["rows"],
            "auto_fetched_at": synced_at,
        }},
        upsert=True,
    )
    MOP_06_HRS_CACHE.pop(selected_date.isoformat(), None)
    return {"success": True, "report_date": selected_date.isoformat(), "rows": len(report["rows"]), "synced_at": synced_at}


def _rtg_access() -> tuple[str, str, str]:
    mongo = MongoService()
    config = mongo.pipeline_config_collection.find_one({"config_type": "RTG"}) or {}
    required = ("rtg_token_url", "rtg_username", "rtg_password")
    if not all(config.get(field) for field in required):
        raise RuntimeError("RTG authentication is not configured in pipeline_config.")
    token = TokenService.get_token(
        config["rtg_token_url"],
        config["rtg_username"],
        config["rtg_password"],
    )
    cap_url = _text(config.get("rtg_cap_on_bar_url")) or DEFAULT_CAP_ON_BAR_URL
    scada_url = _text(config.get("rtg_scada_url")) or DEFAULT_SCADA_URL
    return token, cap_url.rstrip("/") + "/", scada_url.rstrip("/") + "/"


def _rtg_get_json(url: str, token: str, timeout: int = 30) -> Any:
    response = requests.get(
        url,
        headers={"Authorization": f"Token {token}", "Content-Type": "application/json"},
        timeout=timeout,
        verify=False,
    )
    response.raise_for_status()
    return response.json()


def _cap_summary(payload: Any) -> dict:
    series_source = payload.get("cap_on_bar") if isinstance(payload, dict) else payload
    if isinstance(series_source, list):
        series = [_number(value) for value in series_source]
        series = [value for value in series if value is not None]
    else:
        value = _number(series_source)
        series = [value] if value is not None else []
    return {
        "capacity_on_bar_mw": series[-1] if series else None,
        "capacity_on_bar_min_mw": min(series) if series else None,
        "capacity_on_bar_max_mw": max(series) if series else None,
        "capacity_on_bar_points": len(series),
        "capacity_on_bar_interval_minutes": (
            payload.get("interval_minutes") if isinstance(payload, dict) else None
        ),
        # Kept for the SCADA comparison phase. The page only displays the
        # latest and range values.
        "capacity_on_bar_series": series,
    }


def _fetch_capacities(plants: list[dict], source_date: str, token: str, base_url: str) -> dict[str, dict]:
    results: dict[str, dict] = {}

    def fetch(plant_id: str) -> tuple[str, dict]:
        payload = _rtg_get_json(f"{base_url}{source_date}/{plant_id}/", token, timeout=15)
        return plant_id, _cap_summary(payload)

    with ThreadPoolExecutor(max_workers=16) as executor:
        futures = {
            executor.submit(fetch, _text(plant.get("plant_id"))): _text(plant.get("plant_id"))
            for plant in plants
            if _text(plant.get("plant_id"))
        }
        for future in as_completed(futures):
            plant_id = futures[future]
            try:
                result_id, summary = future.result()
                results[result_id] = summary
            except Exception:
                results[plant_id] = _cap_summary(None)
    return results


def _scada_summary(payload: Any) -> dict:
    if isinstance(payload, dict):
        values = next(
            (
                payload.get(key)
                for key in ("actual_gen", "actual", "data", "scada", "values")
                if isinstance(payload.get(key), list)
            ),
            [],
        )
    elif isinstance(payload, list):
        values = payload
    else:
        values = []
    series = [_number(value) for value in values]
    series = [value for value in series if value is not None]
    if not series:
        evening = []
        interval_minutes = None
    else:
        interval_minutes = 1440 / len(series)
        start_index = max(0, min(len(series), round((19 * 60) / interval_minutes)))
        evening = series[start_index:]
    maximum = max(evening) if evening else None
    minimum = min(evening) if evening else None
    maximum_time = ""
    if maximum is not None and interval_minutes:
        maximum_index = start_index + evening.index(maximum)
        minute_of_day = min(1439, round(maximum_index * interval_minutes))
        maximum_time = f"{minute_of_day // 60:02d}:{minute_of_day % 60:02d}"
    return {
        "max_generation_1900_2400_mw": maximum,
        "max_generation_1900_2400_time": maximum_time,
        "min_generation_1900_2400_mw": minimum,
        "generation_max_min_1900_2400": (
            f"{maximum:g} / {minimum:g}"
            if maximum is not None and minimum is not None
            else ""
        ),
        "generation_evening_points": len(evening),
        "scada_total_points": len(series),
        "scada_interval_minutes": interval_minutes,
    }


def _fetch_scada_statistics(plants: list[dict], selected_date: str, token: str, base_url: str) -> dict[str, dict]:
    results: dict[str, dict] = {}

    def fetch(plant_id: str) -> tuple[str, dict]:
        payload = _rtg_get_json(f"{base_url}{selected_date}/{plant_id}/", token, timeout=15)
        return plant_id, _scada_summary(payload)

    with ThreadPoolExecutor(max_workers=16) as executor:
        futures = {
            executor.submit(fetch, _text(plant.get("plant_id"))): _text(plant.get("plant_id"))
            for plant in plants
            if _text(plant.get("plant_id"))
        }
        for future in as_completed(futures):
            plant_id = futures[future]
            try:
                result_id, summary = future.result()
                results[result_id] = summary
            except Exception:
                results[plant_id] = _scada_summary(None)
    return results


def _expected_revival_text(row: dict) -> str:
    expected_date = _text(
        row.get("EXPECTED_REVIVAL_DATE")
        or row.get("expectedRevivalDate")
        or row.get("expected_revival_date")
    )
    expected_time = _text(
        row.get("EXPECTED_REVIVAL_TIME")
        or row.get("expectedRevivalTime")
        or row.get("expected_revival_time")
    )
    if expected_date or expected_time:
        return " ".join(value for value in (expected_date, expected_time) if value)
    return _text(
        row.get("exprecteTimeOfRestoration")
        or row.get("expectedRestorationTime")
        or row.get("expected_restoration_time")
    )


def _fetch_expected_revivals(selected_date: str) -> tuple[dict[str, dict], dict]:
    mongo = MongoService()
    from routes.pipeline_routes import fetch_generation_outage_history_rows

    rows, from_cache, source_url = fetch_generation_outage_history_rows(
        mongo,
        selected_date,
        selected_date,
    )
    unit_lookup: dict[str, dict] = {}
    for unit in mongo.unit_collection.find(
        {},
        {"_id": 0, "Unit_Name": 1, "plant_id": 1},
    ):
        unit_name = _text(unit.get("Unit_Name"))
        if not unit_name:
            continue
        unit_lookup[_key(unit_name)] = unit
        unit_lookup[_compact(unit_name)] = unit

    details_by_plant: dict[str, list[dict]] = {}
    direct_matches = 0
    unit_name_matches = 0
    unmatched = 0
    for outage in rows:
        expected_revival = _expected_revival_text(outage)
        if not expected_revival:
            continue
        unit_name = _text(
            outage.get("ELEMENT_NAME")
            or outage.get("Unit_Name")
            or outage.get("unit_name")
            or outage.get("elementName")
        )
        plant_id = _text(
            outage.get("RTG_PLANT_ID")
            or outage.get("rtg_plant_id")
            or outage.get("PLANT_ID")
            or outage.get("plant_id")
        )
        match_source = "crms_plant_id"
        if plant_id:
            direct_matches += 1
        else:
            unit = unit_lookup.get(_key(unit_name)) or unit_lookup.get(_compact(unit_name)) or {}
            plant_id = _text(unit.get("plant_id"))
            match_source = "units_data.Unit_Name"
            if plant_id:
                unit_name_matches += 1
            else:
                unmatched += 1
                continue
        details_by_plant.setdefault(plant_id, []).append({
            "unit_name": unit_name,
            "expected_revival": expected_revival,
            "match_source": match_source,
            "outage_type": _text(outage.get("TYPE") or outage.get("OUTAGE_TYPE")),
            "reason": _text(outage.get("OUT_REASON") or outage.get("REASON")),
        })

    output: dict[str, dict] = {}
    for plant_id, details in details_by_plant.items():
        unique_pairs = []
        seen_pairs = set()
        for detail in details:
            pair = (detail["unit_name"], detail["expected_revival"])
            if pair in seen_pairs:
                continue
            seen_pairs.add(pair)
            unique_pairs.append(detail)
        unique_times = list(dict.fromkeys(detail["expected_revival"] for detail in unique_pairs))
        if len(unique_times) == 1:
            display = unique_times[0]
        else:
            display = "; ".join(
                f"{detail['unit_name']}: {detail['expected_revival']}"
                for detail in unique_pairs
            )
        output[plant_id] = {
            "expected_revival_time": display,
            "expected_revival_details": unique_pairs,
        }

    return output, {
        "crms_records": len(rows),
        "expected_revival_records": sum(len(value) for value in details_by_plant.values()),
        "plants_with_expected_revival": len(output),
        "direct_plant_id_matches": direct_matches,
        "unit_name_fallback_matches": unit_name_matches,
        "unmatched_expected_revivals": unmatched,
        "from_cache": from_cache,
        "source_url": source_url,
    }


def _apply_saved_edits(output: dict[str, list[dict]], selected_date: date) -> None:
    mongo = MongoService()
    collection = mongo.db[EDIT_COLLECTION]
    date_text = selected_date.isoformat()
    current = {
        _text(document.get("plant_id")): document
        for document in collection.find({"report_date": date_text}, {"_id": 0})
    }
    plant_ids = [
        _text(row.get("plant_id"))
        for rows in output.values()
        for row in rows
        if _text(row.get("plant_id"))
    ]
    previous_reasons: dict[str, str] = {}
    if plant_ids:
        previous_date_text = (selected_date - timedelta(days=1)).isoformat()
        previous_documents = collection.find(
            {
                "plant_id": {"$in": plant_ids},
                "report_date": previous_date_text,
                "reason_for_not_attaining_full_generation": {"$nin": [None, ""]},
            },
            {
                "_id": 0,
                "plant_id": 1,
                "report_date": 1,
                "reason_for_not_attaining_full_generation": 1,
            },
        )
        for document in previous_documents:
            plant_id = _text(document.get("plant_id"))
            if plant_id and plant_id not in previous_reasons:
                previous_reasons[plant_id] = _text(
                    document.get("reason_for_not_attaining_full_generation")
                )

    editable_map = {
        "running_capacity_mw": "capacity_on_bar_mw",
        "max_generation_1900_2400_mw": "max_generation_1900_2400_mw",
        "max_generation_1900_2400_time": "max_generation_1900_2400_time",
        "generation_max_min_1900_2400": "generation_max_min_1900_2400",
    }
    for rows in output.values():
        for row in rows:
            plant_id = _text(row.get("plant_id"))
            saved = current.get(plant_id, {})
            edited_fields = []
            for saved_key, row_key in editable_map.items():
                if saved_key in saved and saved.get(saved_key) is not None:
                    row[row_key] = saved.get(saved_key)
                    edited_fields.append(saved_key)

            if "reason_for_not_attaining_full_generation" in saved:
                row["reason_for_not_attaining_full_generation"] = _text(
                    saved.get("reason_for_not_attaining_full_generation")
                )
                reason_source = "selected_date"
            elif previous_reasons.get(plant_id):
                row["reason_for_not_attaining_full_generation"] = previous_reasons[plant_id]
                reason_source = "previous_saved_day"
            else:
                reason_source = "crms"

            installed = _number(row.get("installed_capacity_mw"))
            running = _number(row.get("capacity_on_bar_mw"))
            maximum = _number(row.get("max_generation_1900_2400_mw"))
            row["outage_capacity_mw"] = (
                round(installed - running, 3)
                if installed is not None and running is not None
                else None
            )
            ex_bus = round(running * 0.93, 3) if running is not None else None
            row["ex_bus_running_capacity_mw"] = ex_bus
            row["running_units_margin_mw"] = (
                round(ex_bus - maximum, 3)
                if ex_bus is not None and maximum is not None
                else None
            )
            row["loading_factor_pct"] = (
                round((maximum / ex_bus) * 100, 2)
                if maximum is not None and ex_bus not in (None, 0)
                else None
            )
            row["manual_edit_fields"] = edited_fields
            row["reason_source"] = reason_source


def build_static_report(report_date: str | None = None, force_refresh: bool = False) -> dict:
    selected_date = _parse_report_date(report_date)
    capacity_date = selected_date - timedelta(days=1)
    cache_key = capacity_date.isoformat()
    cached = REPORT_CACHE.get(cache_key)
    if not force_refresh and cached and monotonic() - cached[0] < CACHE_SECONDS:
        return cached[1]

    token, cap_base_url, scada_base_url = _rtg_access()
    master_rows = _rtg_get_json(GENERATOR_MASTER_URL, token)
    if not isinstance(master_rows, list):
        raise RuntimeError("RTG generator master returned an unexpected response.")
    with ThreadPoolExecutor(max_workers=3) as executor:
        capacity_future = executor.submit(
            _fetch_capacities,
            master_rows,
            capacity_date.isoformat(),
            token,
            cap_base_url,
        )
        scada_future = executor.submit(
            _fetch_scada_statistics,
            master_rows,
            selected_date.isoformat(),
            token,
            scada_base_url,
        )
        expected_revival_future = executor.submit(
            _fetch_expected_revivals,
            selected_date.isoformat(),
        )
        capacity_by_plant = capacity_future.result()
        scada_by_plant = scada_future.result()
        expected_revival_by_plant, expected_revival_summary = expected_revival_future.result()

    output = {"thermal": [], "hydro": []}
    skipped = Counter()
    for plant in master_rows:
        utility_type = _text(plant.get("utility_type"))
        utility_key = _key(utility_type)
        state_name = _text(plant.get("state_name"))
        if utility_key in STATE_UTILITY_TYPES:
            section = state_name.upper() if state_name else "STATE - NOT MAPPED"
            category = "State"
        elif utility_key in REGIONAL_IPP_UTILITY_TYPES:
            section = "REGIONAL IPP"
            category = "Regional IPP"
        elif utility_key in ISGS_UTILITY_TYPES:
            section = "ISGS"
            category = "ISGS"
        else:
            skipped["unsupported_utility_type"] += 1
            continue

        fuel_type = _text(plant.get("fuel_type"))
        fuel_key = _key(fuel_type)
        if fuel_key in THERMAL_FUELS:
            fuel_group = "thermal"
        elif fuel_key in HYDRO_FUELS:
            fuel_group = "hydro"
        else:
            skipped["unsupported_fuel_type"] += 1
            continue

        plant_id = _text(plant.get("plant_id"))
        capacity = capacity_by_plant.get(plant_id, _cap_summary(None))
        scada = scada_by_plant.get(plant_id, _scada_summary(None))
        expected_revival = expected_revival_by_plant.get(
            plant_id,
            {"expected_revival_time": "", "expected_revival_details": []},
        )
        outage_reasons = list(
            dict.fromkeys(
                detail.get("reason", "")
                for detail in expected_revival.get("expected_revival_details", [])
                if detail.get("reason")
            )
        )
        installed_capacity = _number(plant.get("installed_capacity"))
        capacity_on_bar = capacity.get("capacity_on_bar_mw")
        max_generation = scada.get("max_generation_1900_2400_mw")
        outage_capacity = (
            round(installed_capacity - capacity_on_bar, 3)
            if installed_capacity is not None and capacity_on_bar is not None
            else None
        )
        running_margin = (
            round((capacity_on_bar * 0.93) - max_generation, 3)
            if capacity_on_bar is not None and max_generation is not None
            else None
        )
        ex_bus_running_capacity = (
            round(capacity_on_bar * 0.93, 3)
            if capacity_on_bar is not None
            else None
        )
        loading_factor = (
            round((max_generation / ex_bus_running_capacity) * 100, 2)
            if max_generation is not None
            and ex_bus_running_capacity is not None
            and ex_bus_running_capacity != 0
            else None
        )
        row = {
            "section": section,
            "category": category,
            "state_name": state_name.upper(),
            "utility_type": utility_type,
            "plant_name": _text(plant.get("plant_name")),
            "plant_id": plant_id,
            "fuel_name": fuel_type,
            "installed_capacity_mw": installed_capacity,
            "effective_capacity_mw": _number(plant.get("effective_capacity")),
            "owner_name": _text(plant.get("owner_name")),
            "wbes_acronym": _text(plant.get("wbes_acronym")),
            "scada_point": _text(plant.get("scada_point")),
            **capacity,
            **scada,
            "outage_capacity_mw": outage_capacity,
            "running_units_margin_mw": running_margin,
            "ex_bus_running_capacity_mw": ex_bus_running_capacity,
            "loading_factor_pct": loading_factor,
            "reason_for_not_attaining_full_generation": "; ".join(outage_reasons),
            **expected_revival,
        }
        output[fuel_group].append(row)

    _apply_saved_edits(output, selected_date)

    def sort_key(row: dict) -> tuple:
        category_order = {"State": 0, "Regional IPP": 1, "ISGS": 2}
        return (
            category_order.get(row.get("category"), 99),
            row["section"],
            row["plant_name"].upper(),
        )

    for group in output:
        output[group].sort(key=sort_key)
        for serial, row in enumerate(output[group], 1):
            row["serial"] = serial

    available_capacities = sum(
        1
        for group in output.values()
        for row in group
        if row["capacity_on_bar_mw"] is not None
    )
    available_scada = sum(
        1
        for group in output.values()
        for row in group
        if row["max_generation_1900_2400_mw"] is not None
    )
    margin_categories = {
        "State/State IPP": 0.0,
        "Regional IPP": 0.0,
        "ISGS": 0.0,
    }
    category_labels = {
        "State": "State/State IPP",
        "Regional IPP": "Regional IPP",
        "ISGS": "ISGS",
    }
    for rows in output.values():
        for row in rows:
            label = category_labels.get(row.get("category"))
            margin = _number(row.get("running_units_margin_mw"))
            if label and margin is not None:
                # A unit running above 93% ex-bus capacity has no available
                # upward reserve; it must not reduce reserve on other units.
                margin_categories[label] += max(0.0, margin)
    margin_categories = {
        label: round(value, 2)
        for label, value in margin_categories.items()
    }
    eastern_total = round(sum(margin_categories.values()), 2)
    margin_available = {
        "report_date": selected_date.isoformat(),
        "calculation": "Sum of max(0, Running Capacity * 0.93 - Max Generation 19:00-24:00)",
        "regions": [
            {
                "region": region,
                "total_label": total_label,
                "categories": (
                    dict(margin_categories)
                    if region == "ER"
                    else {label: 0.0 for label in margin_categories}
                ),
                "total_mw": eastern_total if region == "ER" else 0.0,
            }
            for region, total_label in (
                ("NR", "Northern Region Total"),
                ("WR", "Western Region Total"),
                ("SR", "Southern Region Total"),
                ("ER", "Eastern Region Total"),
                ("NER", "N Eastern Region Total"),
            )
        ],
        "all_india": {
            "region": "All India",
            "total_label": "All India Total",
            "categories": dict(margin_categories),
            "total_mw": eastern_total,
        },
    }
    report = {
        "success": True,
        "report_name": "Plant Deviation from I/C - MOP",
        "phase": "static_with_capacity_on_bar",
        "report_date": selected_date.isoformat(),
        "capacity_on_bar_date": capacity_date.isoformat(),
        "source": {
            "generator_master": GENERATOR_MASTER_URL,
            "capacity_on_bar": f"{cap_base_url}{{D-1 date}}/{{plant_id}}/",
            "rtg_scada": f"{scada_base_url}{{selected date}}/{{plant_id}}/",
            "expected_revival": "CRMS GenOutagesHistoryData",
            "expected_revival_fallback": "units_data.Unit_Name -> plant_id",
            "fuel_key": "fuel_type",
            "utility_key": "utility_type",
            "state_key": "state_name",
            "installed_capacity_key": "installed_capacity",
        },
        "summary": {
            "master_records": len(master_rows),
            "thermal_records": len(output["thermal"]),
            "hydro_records": len(output["hydro"]),
            "capacity_on_bar_available": available_capacities,
            "capacity_on_bar_missing": len(output["thermal"]) + len(output["hydro"]) - available_capacities,
            "scada_evening_available": available_scada,
            "scada_evening_missing": len(output["thermal"]) + len(output["hydro"]) - available_scada,
            **expected_revival_summary,
            "skipped": dict(skipped),
        },
        "margin_available": margin_available,
        **output,
    }
    REPORT_CACHE[cache_key] = (monotonic(), report)
    return report


@router.get("/static")
def get_static_report(
    report_date: str | None = Query(default=None),
    refresh: bool = Query(default=False),
):
    try:
        return build_static_report(report_date, force_refresh=refresh)
    except HTTPException:
        raise
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"RTG API request failed: {exc}") from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Unable to prepare plant table: {exc}") from exc


@router.get("/day-ahead")
def get_day_ahead_thermal_report(
    report_date: str | None = Query(default=None),
    refresh: bool = Query(default=False),
):
    return _day_ahead_thermal_report(report_date, force_refresh=refresh)


@router.get("/all-india")
def get_all_india_partial_outage_report(
    report_date: str | None = Query(default=None),
    fetch_mail: bool = Query(default=False),
):
    return _all_india_report(report_date, fetch_mail=fetch_mail)


@router.put("/day-ahead/save")
def save_day_ahead_thermal_report(payload: DayAheadReportSave):
    selected_date = _parse_report_date(payload.report_date)
    report = _day_ahead_thermal_report(selected_date.isoformat())
    source_rows = {
        (row["source_row"], _compact(row["plant_name"])): row
        for row in report["rows"]
    }
    saved_rows = []
    for item in payload.rows:
        values = item.model_dump() if hasattr(item, "model_dump") else item.dict()
        key = (values["source_row"], _compact(values["plant_name"]))
        source = source_rows.get(key)
        if not source:
            raise HTTPException(
                status_code=422,
                detail=f"{values['plant_name']}: station is not present in the selected source report.",
            )
        for field in ("coal_stock_days_left", "daily_coal_requirement"):
            value = values.get(field)
            if value is not None and value < 0:
                raise HTTPException(status_code=422, detail=f"{values['plant_name']}: {field} cannot be negative.")
        saved_rows.append({
            **source,
            "coal_stock_days_left": values.get("coal_stock_days_left"),
            "daily_coal_requirement": values.get("daily_coal_requirement"),
        })

    saved_at = datetime.now(timezone.utc)
    MongoService().db[DAY_AHEAD_COLLECTION].update_one(
        {"report_date": selected_date.isoformat()},
        {"$set": {
            "report_date": selected_date.isoformat(),
            "source_file": report["source_file"],
            "filter": report["filter"],
            "sort": report["sort"],
            "summary": report["summary"],
            "rows": saved_rows,
            "saved_at": saved_at,
        }},
        upsert=True,
    )
    MOP_06_HRS_CACHE.pop(selected_date.isoformat(), None)
    return {
        "success": True,
        "report_date": selected_date.isoformat(),
        "saved_rows": len(saved_rows),
        "saved_at": saved_at,
    }


@router.put("/static/edits/{plant_id}")
def save_static_report_edit(plant_id: str, payload: PlantDeviationEdit):
    selected_date = _parse_report_date(payload.report_date)
    if payload.max_generation_1900_2400_time:
        try:
            datetime.strptime(payload.max_generation_1900_2400_time, "%H:%M")
        except ValueError as exc:
            raise HTTPException(status_code=422, detail="Maximum generation time must be HH:MM.") from exc

    values = (
        payload.model_dump(exclude_unset=True)
        if hasattr(payload, "model_dump")
        else payload.dict(exclude_unset=True)
    )
    values.pop("report_date", None)
    values.update({
        "plant_id": _text(plant_id),
        "report_date": selected_date.isoformat(),
        "updated_at": datetime.now(timezone.utc),
    })
    mongo = MongoService()
    mongo.db[EDIT_COLLECTION].update_one(
        {"plant_id": _text(plant_id), "report_date": selected_date.isoformat()},
        {"$set": values},
        upsert=True,
    )
    REPORT_CACHE.clear()
    return {
        "success": True,
        "plant_id": _text(plant_id),
        "report_date": selected_date.isoformat(),
    }


@router.put("/static/edits")
def save_static_report_edits(payload: PlantDeviationBatchEdit):
    selected_date = _parse_report_date(payload.report_date)
    date_text = selected_date.isoformat()
    mongo = MongoService()
    collection = mongo.db[EDIT_COLLECTION]
    saved = 0
    for item in payload.edits:
        if item.max_generation_1900_2400_time:
            try:
                datetime.strptime(item.max_generation_1900_2400_time, "%H:%M")
            except ValueError as exc:
                raise HTTPException(
                    status_code=422,
                    detail=f"{item.plant_id}: maximum generation time must be HH:MM.",
                ) from exc
        values = (
            item.model_dump(exclude_unset=True)
            if hasattr(item, "model_dump")
            else item.dict(exclude_unset=True)
        )
        plant_id = _text(values.pop("plant_id", ""))
        values.pop("report_date", None)
        if not plant_id:
            continue
        values.update({
            "plant_id": plant_id,
            "report_date": date_text,
            "updated_at": datetime.now(timezone.utc),
        })
        collection.update_one(
            {"plant_id": plant_id, "report_date": date_text},
            {"$set": values},
            upsert=True,
        )
        saved += 1
    REPORT_CACHE.clear()
    return {"success": True, "report_date": date_text, "saved": saved}


def _write_sheet(workbook: Workbook, title: str, rows: list[dict], report: dict) -> None:
    sheet = workbook.create_sheet(title=title)
    black = "000000"
    grey = "A6A6A6"
    light_grey = "D9D9D9"
    pale_green = "E2F0D9"
    yellow = "FFFF00"
    pale_yellow = "FFF2CC"
    white = "FFFFFF"
    border = Border(
        left=Side(style="thin", color=black),
        right=Side(style="thin", color=black),
        top=Side(style="thin", color=black),
        bottom=Side(style="thin", color=black),
    )

    sheet.merge_cells("A1:K1")
    sheet["A1"] = "ER"
    sheet["A1"].fill = PatternFill("solid", fgColor=pale_green)
    sheet["A1"].font = Font(color=black, bold=True, size=12)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")

    merged_headers = {
        "A2:A3": "Station/Constituents",
        "B2": "Installed\nCapacity (A)",
        "C2": "Running\nCapacity (B=A-E)",
        "D2:E2": "Max generation between 1900 to 2400\n(C)",
        "F2:F3": "Range of generation between 1900hrs to 2400hrs\n(Max(MW)&Min(MW))",
        "G2": "Margin remained available on running units\n(Actual margin or D=B*0.93-C)",
        "H2": "Outage Capacity\n(E)",
        "I2:I3": "Reason for not attaining full generation",
        "J2": "Loading Factor\n[max gen/ex-bus running capacity]\n(corres. to max gen. achieved b/w 1900 to 2400hrs)",
        "K2:K3": "Expected revival dates of units\nunder outage",
    }
    for cell_range, label in merged_headers.items():
        if ":" in cell_range:
            sheet.merge_cells(cell_range)
        sheet[cell_range.split(":")[0]] = label
    for coordinate, label in {
        "B3": "(MW)", "C3": "(MW)", "D3": "(MW)", "E3": "Hrs",
        "G3": "(MW)", "H3": "(MW)", "J3": "%",
    }.items():
        sheet[coordinate] = label

    for row_cells in sheet.iter_rows(min_row=2, max_row=3, min_col=1, max_col=11):
        for cell in row_cells:
            cell.fill = PatternFill("solid", fgColor=grey)
            cell.font = Font(color=black, bold=True, size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
    for row_number in (2, 3):
        sheet.cell(row_number, 7).fill = PatternFill("solid", fgColor=yellow)
        sheet.cell(row_number, 10).fill = PatternFill("solid", fgColor=pale_yellow)
        sheet.cell(row_number, 11).fill = PatternFill("solid", fgColor=white)
    sheet.row_dimensions[2].height = 64
    sheet.row_dimensions[3].height = 25

    row_index = 4
    for category, category_label in (
        ("State", "State Entities Generation"),
        ("Regional IPP", "Regional IPP Generation"),
        ("ISGS", "ISGS Generation"),
    ):
        category_rows = [row for row in rows if row.get("category") == category]
        if not category_rows:
            continue
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=11)
        category_cell = sheet.cell(row_index, 1, category_label)
        category_cell.fill = PatternFill("solid", fgColor=light_grey)
        category_cell.font = Font(color=black, bold=True, size=11)
        category_cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in range(1, 12):
            sheet.cell(row_index, column).border = border
        row_index += 1

        sections = list(dict.fromkeys(row.get("section") or "NOT MAPPED" for row in category_rows))
        for section in sections:
            section_rows = [
                row for row in category_rows
                if (row.get("section") or "NOT MAPPED") == section
            ]
            sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=11)
            section_cell = sheet.cell(row_index, 1, section)
            section_cell.font = Font(color=black, bold=True, size=10)
            section_cell.alignment = Alignment(horizontal="center", vertical="center")
            for column in range(1, 12):
                sheet.cell(row_index, column).border = border
            row_index += 1

            for row in section_rows:
                values = [
                    row.get("plant_name", ""),
                    row.get("installed_capacity_mw"),
                    row.get("capacity_on_bar_mw"),
                    row.get("max_generation_1900_2400_mw"),
                    row.get("max_generation_1900_2400_time", ""),
                    row.get("generation_max_min_1900_2400", ""),
                    row.get("running_units_margin_mw"),
                    row.get("outage_capacity_mw"),
                    row.get("reason_for_not_attaining_full_generation", ""),
                    row.get("loading_factor_pct"),
                    row.get("expected_revival_time", ""),
                ]
                for column_index, value in enumerate(values, 1):
                    cell = sheet.cell(row_index, column_index, "" if value is None else value)
                    cell.border = border
                    cell.font = Font(color=black, size=10)
                    cell.alignment = Alignment(
                        horizontal="left" if column_index in (1, 9, 11) else "center",
                        vertical="center",
                        wrap_text=True,
                    )
                    if column_index == 7:
                        cell.fill = PatternFill("solid", fgColor=yellow)
                    elif column_index == 10:
                        cell.fill = PatternFill("solid", fgColor=pale_yellow)
                row_index += 1

            totals = {
                2: sum((row.get("installed_capacity_mw") or 0) for row in section_rows),
                3: sum((row.get("capacity_on_bar_mw") or 0) for row in section_rows),
                7: sum((row.get("running_units_margin_mw") or 0) for row in section_rows),
                8: sum((row.get("outage_capacity_mw") or 0) for row in section_rows),
            }
            sheet.cell(row_index, 1, "Total")
            for column in range(1, 12):
                cell = sheet.cell(row_index, column)
                if column in totals:
                    cell.value = round(totals[column], 2)
                cell.border = border
                cell.font = Font(color=black, bold=True, size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if column == 7:
                    cell.fill = PatternFill("solid", fgColor=yellow)
                elif column == 10:
                    cell.fill = PatternFill("solid", fgColor=pale_yellow)
            row_index += 1

    widths = [34, 14, 15, 18, 12, 22, 21, 15, 42, 23, 30]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A4"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:3"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3


def _write_margin_available_sheet(workbook: Workbook, report: dict) -> None:
    sheet = workbook.create_sheet(title="Margin Available")
    black = "000000"
    navy = "08103A"
    header_blue = "D9EAF7"
    pale_green = "E2F0D9"
    light_grey = "F2F2F2"
    white = "FFFFFF"
    border = Border(
        left=Side(style="thin", color=black),
        right=Side(style="thin", color=black),
        top=Side(style="thin", color=black),
        bottom=Side(style="thin", color=black),
    )
    margin_data = report.get("margin_available") or {}
    report_date = margin_data.get("report_date") or report.get("report_date") or ""
    try:
        report_date_display = date.fromisoformat(report_date).strftime("%d-%m-%Y")
    except (TypeError, ValueError):
        report_date_display = report_date

    sheet.merge_cells("A1:C1")
    sheet["A1"] = f"Margin available as on {report_date_display} (between 1900 to 2400hrs)"
    sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    sheet["A1"].font = Font(color=white, bold=True, size=13)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 25

    sheet["A2"] = "Region"
    sheet["B2"] = "Category"
    sheet["C2"] = "Reserve Available (MW)"
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor=header_blue)
        cell.font = Font(color=black, bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    regions = list(margin_data.get("regions") or [])
    all_india = margin_data.get("all_india") or {}
    if all_india:
        regions.append(all_india)
    row_index = 3
    for region_data in regions:
        start_row = row_index
        categories = region_data.get("categories") or {}
        for label in ("State/State IPP", "Regional IPP", "ISGS"):
            sheet.cell(row_index, 2, label)
            sheet.cell(row_index, 3, round(float(categories.get(label) or 0), 2))
            row_index += 1
        sheet.cell(row_index, 2, region_data.get("total_label") or f"{region_data.get('region', '')} Total")
        sheet.cell(row_index, 3, round(float(region_data.get("total_mw") or 0), 2))
        total_row = row_index
        row_index += 1

        sheet.merge_cells(start_row=start_row, start_column=1, end_row=total_row, end_column=1)
        region_cell = sheet.cell(start_row, 1, region_data.get("region") or "")
        region_cell.font = Font(color=black, bold=True, size=11)
        region_cell.alignment = Alignment(horizontal="center", vertical="center")
        for current_row in range(start_row, total_row + 1):
            for column in range(1, 4):
                cell = sheet.cell(current_row, column)
                cell.border = border
                cell.font = Font(color=black, bold=current_row == total_row, size=10)
                cell.alignment = Alignment(
                    horizontal="right" if column == 3 else "left",
                    vertical="center",
                )
                if current_row == total_row:
                    cell.fill = PatternFill("solid", fgColor=pale_green)
                elif region_data.get("region") == "All India":
                    cell.fill = PatternFill("solid", fgColor=light_grey)
        region_cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 25
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _write_day_ahead_sheet(workbook: Workbook, report: dict) -> None:
    sheet = workbook.create_sheet(title="Thermal Margin Report")
    headers = [
        "Station/Constituents",
        "Installed Capacity (A) MW",
        "Running Capacity (B=A-E) MW",
        "Max generation 1900-2400 (C) MW",
        "Time",
        "Generation range 1900-2400 Max/Min",
        "Margin on running units (D=B*0.93-C) MW",
        "Outage Capacity (E) MW",
        "Reason for not attaining full generation",
        "Loading Factor %",
        "Expected revival dates",
        "No. of days left",
        "Daily coal requirement",
    ]
    navy, white, pale_yellow, pale_green = "08103A", "FFFFFF", "FFF2CC", "E2F0D9"
    thin_border = Border(
        left=Side(style="thin", color="A6A6A6"),
        right=Side(style="thin", color="A6A6A6"),
        top=Side(style="thin", color="A6A6A6"),
        bottom=Side(style="thin", color="A6A6A6"),
    )
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    sheet.row_dimensions[1].height = 55

    keys = [
        "plant_name", "installed_capacity_mw", "capacity_on_bar_mw",
        "max_generation_1900_2400_mw", "max_generation_1900_2400_time",
        "generation_max_min_1900_2400", "running_units_margin_mw",
        "outage_capacity_mw", "reason_for_not_attaining_full_generation",
        "loading_factor_pct", "expected_revival_time", "coal_stock_days_left",
        "daily_coal_requirement",
    ]
    for source in report["rows"]:
        sheet.append([source.get(key) for key in keys])
        row_index = sheet.max_row
        for column_index, cell in enumerate(sheet[row_index], 1):
            cell.border = thin_border
            cell.font = Font(color="000000", size=10)
            cell.alignment = Alignment(
                horizontal="left" if column_index in (1, 9, 11) else "center",
                vertical="center",
                wrap_text=True,
            )
            if column_index == 7:
                cell.fill = PatternFill("solid", fgColor=pale_yellow)
            elif column_index in (12, 13):
                cell.fill = PatternFill("solid", fgColor=pale_green)

    total_row = sheet.max_row + 1
    sheet.cell(total_row, 1, f"TOTAL ({len(report['rows'])} stations)")
    totals = {
        2: report["summary"]["total_installed_capacity_mw"],
        3: report["summary"]["total_running_capacity_mw"],
        7: report["summary"]["total_margin_mw"],
        8: report["summary"]["total_outage_capacity_mw"],
    }
    for column_index in range(1, 14):
        cell = sheet.cell(total_row, column_index)
        if column_index in totals:
            cell.value = totals[column_index]
        cell.font = Font(color="000000", bold=True, size=10)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [34, 16, 17, 18, 11, 23, 21, 16, 40, 16, 28, 16, 22]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:M{sheet.max_row - 1}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


@router.get("/day-ahead/excel")
def download_day_ahead_thermal_report(report_date: str | None = Query(default=None)):
    report = _day_ahead_thermal_report(report_date)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Day-ahead Thermal Margin Consolidation", None])
    summary_sheet.append(["Report date", report["report_date"]])
    summary_sheet.append(["Source file", report["source_file"]])
    summary_sheet.append(["Filter", report["filter"]])
    summary_sheet.append(["Sort", report["sort"]])
    summary_sheet.append(["Thermal stations with non-zero margin", report["summary"]["thermal_margin_stations"]])
    summary_sheet.append(["Total margin on running units (MW)", report["summary"]["total_margin_mw"]])
    summary_sheet.append(["Outage capacity in displayed rows (MW)", report["summary"]["total_outage_capacity_mw"]])
    summary_sheet.append(["Total outage capacity - all Thermal rows (MW)", report["summary"]["all_thermal_outage_capacity_mw"]])
    summary_sheet.merge_cells("A1:B1")
    summary_sheet["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    summary_sheet["A1"].fill = PatternFill("solid", fgColor="08103A")
    summary_sheet["A1"].alignment = Alignment(horizontal="center")
    for row in summary_sheet.iter_rows(min_row=2, max_row=9, min_col=1, max_col=2):
        for cell in row:
            cell.border = Border(
                left=Side(style="thin", color="000000"),
                right=Side(style="thin", color="000000"),
                top=Side(style="thin", color="000000"),
                bottom=Side(style="thin", color="000000"),
            )
        row[0].font = Font(bold=True)
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 65
    _write_day_ahead_sheet(workbook, report)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f'attachment; filename="ER_Thermal_Margin_Consolidated_{report["report_date"]}.xlsx"'
            )
        },
    )


@router.get("/all-india/excel")
def download_all_india_partial_outage_report(report_date: str | None = Query(default=None)):
    report = _all_india_report(report_date, fetch_mail=False)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["All India Partial Outage - Margin Consolidation", None])
    summary.append(["Report date", report["report_date"]])
    summary.append(["Source attachment", report["source_file"]])
    summary.append(["Source sheets", ", ".join(report["source_sheets"])])
    summary.append(["Filter", report["filter"]])
    summary.append(["Sort", report["sort"]])
    summary.append(["Stations", report["summary"]["stations"]])
    summary.append(["Total margin (MW)", report["summary"]["total_margin_mw"]])
    summary.append(["Total outage - all source rows (MW)", report["summary"]["all_outage_capacity_mw"]])
    summary.merge_cells("A1:B1")
    summary["A1"].fill = PatternFill("solid", fgColor="312E81")
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    summary["A1"].alignment = Alignment(horizontal="center")
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 70

    sheet = workbook.create_sheet("NR SR WR Consolidated")
    headers = ["Region", "Station/Constituents", "Section", "Installed MW", "Running MW", "Max generation MW", "Time", "Generation Max/Min", "Margin on running units MW", "Outage MW", "Reason for not attaining full generation", "Loading Factor %", "Expected revival"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="312E81")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in report["rows"]:
        sheet.append([
            row["region"], row["plant_name"], row["section"], row["installed_capacity_mw"],
            row["capacity_on_bar_mw"], row["max_generation_1900_2400_mw"],
            row["max_generation_1900_2400_time"], row["generation_max_min_1900_2400"],
            row["running_units_margin_mw"], row["outage_capacity_mw"],
            row["reason_for_not_attaining_full_generation"], row["loading_factor_pct"],
            row["expected_revival_time"],
        ])
        sheet.cell(sheet.max_row, 9).fill = PatternFill("solid", fgColor="FDE68A")
        sheet.cell(sheet.max_row, 9).font = Font(color="713F12", bold=True)
    widths = [10, 34, 24, 14, 14, 18, 11, 22, 22, 14, 40, 16, 28]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:M{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="All_India_Partial_Outage_{report["report_date"]}.xlsx"'},
    )


@router.get("/static/excel")
def download_static_report(report_date: str | None = Query(default=None)):
    try:
        report = build_static_report(report_date)
        workbook = Workbook()
        workbook.remove(workbook.active)
        _write_sheet(workbook, "Thermal", report["thermal"], report)
        _write_sheet(workbook, "Hydro", report["hydro"], report)
        _write_margin_available_sheet(workbook, report)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": (
                    f'attachment; filename="Plant_Deviation_IC_MOP_{report["report_date"]}.xlsx"'
                )
            },
        )
    except HTTPException:
        raise
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"RTG API request failed: {exc}") from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Unable to export plant table: {exc}") from exc
