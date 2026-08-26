import io
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from bson import ObjectId
from fastapi import APIRouter, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, Table, TableStyle

from services.db_handler import MongoService


router = APIRouter(prefix="/api/old-logbook", tags=["Old Logbook"])

OLD_LOGBOOK_DB = "Old_logbook"
VIOLATION_COLLECTION = "Violation_Message"

COLLECTION_CONFIG: Dict[str, Dict[str, Any]] = {
    "shutdown": {
        "collection": "shutdown",
        "label": "Shutdown",
        "element": "ElementName",
        "outage_date": "ActualOutageDate",
        "outage_time": "ActualOutageTime",
        "revival_date": "ActualRestoreDate",
        "revival_time": "ActualRestoreTime",
        "reason_fields": ["Reason"],
        "planned_fields": ["PlannedOutage", "PlannedRestore"],
        "element_type": "EntityId",
    },
    "tripping": {
        "collection": "Tripping",
        "label": "Tripping",
        "element": "Name",
        "outage_date": "TripDate",
        "outage_time": "TripTime",
        "revival_date": "RevivalDate",
        "revival_time": "RevivalTime",
        "reason_fields": ["EndRelayReasonOne", "EndRelayReasonTwo"],
        "reason_label": "Relay",
    },
    "outage": {
        "collection": "Outage",
        "label": "Outage",
        "element": "Name",
        "outage_date": "LogDate",
        "outage_time": "LogTime",
        "revival_date": "RestoreDate",
        "revival_time": "RestoreTime",
        "reason_fields": ["Reason"],
    },
}

ELEMENT_TYPE_LABELS = {
    "AC_TRANSMISSION_LINE_CIRCUIT": "Transmission Line",
    "TRANSFORMER": "Transformer",
    "BUS_REACTOR": "Bus Reactor",
    "LINE_REACTOR": "Line Reactor",
    "BUS": "Bus",
    "BAY": "Bay",
    "GENERATING_UNIT": "Generating Unit",
    "AUTO_RECLOSER": "Auto Recloser",
    "HVDC_POLE": "HVDC Pole",
    "STATCOM": "STATCOM",
}

OLD_LOGBOOK_ELEMENT_TYPE_MAP = {
    "14": "AC_TRANSMISSION_LINE_CIRCUIT",
    "TRANSMISSION LINE": "AC_TRANSMISSION_LINE_CIRCUIT",
    "9": "TRANSFORMER",
    "TRANSFORMER": "TRANSFORMER",
    "4": "BUS_REACTOR",
    "BUS REACTOR": "BUS_REACTOR",
    "5": "LINE_REACTOR",
    "LINE REACTOR": "LINE_REACTOR",
    "16": "BUS",
    "BUS": "BUS",
    "25": "BAY",
    "BAY": "BAY",
    "8": "GENERATING_UNIT",
    "GENERATING UNIT": "GENERATING_UNIT",
    "26": "AUTO_RECLOSER",
    "AUTO RECLOSER": "AUTO_RECLOSER",
    "15": "HVDC_POLE",
    "HVDC POLE": "HVDC_POLE",
    "STATCOM": "STATCOM",
}


def to_jsonable(value: Any) -> Any:
    if isinstance(value, ObjectId):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, list):
        return [to_jsonable(item) for item in value]
    if isinstance(value, dict):
        return {str(key): to_jsonable(val) for key, val in value.items()}
    return value


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(to_jsonable(value)).strip()
    if text.lower() in {"none", "nan", "nat", "null"}:
        return ""
    return text


def combine_datetime(doc: dict, date_key: str, time_key: str) -> str:
    date_part = clean_text(doc.get(date_key))
    time_part = clean_text(doc.get(time_key))
    return " ".join(part for part in [date_part, time_part] if part)


def combine_fields(doc: dict, fields: List[str], separator: str = " | ") -> str:
    values = [clean_text(doc.get(field)) for field in fields]
    return separator.join(value for value in values if value)


def normalize_element_type_key(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    compact = text.upper().replace("_", " ").replace("-", " ")
    compact = " ".join(compact.split())
    return OLD_LOGBOOK_ELEMENT_TYPE_MAP.get(compact, OLD_LOGBOOK_ELEMENT_TYPE_MAP.get(text.upper(), text.upper()))


def display_element_type(value: Any) -> str:
    key = normalize_element_type_key(value)
    return ELEMENT_TYPE_LABELS.get(key, clean_text(value))


def parse_logbook_date(value: Any):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if not text:
        return None
    text = text.replace("T", " ").split(" ")[0]
    for fmt in (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y/%m/%d",
        "%d.%m.%Y",
        "%d-%b-%Y",
        "%d/%b/%Y",
        "%d-%B-%Y",
        "%d/%B/%Y",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def violation_datetime(doc: dict) -> datetime:
    parsed_date = parse_logbook_date(doc.get("CreatedDate")) or date.min
    time_text = clean_text(doc.get("CreatedTime"))
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            parsed_time = datetime.strptime(time_text, fmt).time()
            return datetime.combine(parsed_date, parsed_time)
        except ValueError:
            continue
    return datetime.combine(parsed_date, datetime.min.time())


def normalize_violation(doc: dict) -> dict:
    raw = to_jsonable(doc)
    timestamp = violation_datetime(doc)
    return {
        "id": clean_text(raw.get("_id")),
        "request_id": clean_text(raw.get("RequestId")),
        "logbook_id": clean_text(raw.get("LogbookId")),
        "message_id": raw.get("Id", ""),
        "date": timestamp.date().isoformat() if timestamp.date() != date.min else "",
        "display_date": clean_text(raw.get("CreatedDate")),
        "time": clean_text(raw.get("CreatedTime")),
        "constituent": clean_text(raw.get("Constituent")),
        "violation_type": clean_text(raw.get("ViolationType")),
        "sub_violation_type": clean_text(raw.get("SubViolationType")),
        "message": clean_text(raw.get("Message")),
        "frequency": raw.get("Frequency", ""),
        "schedule_mw": raw.get("ScheduleMW", ""),
        "actual_mw": raw.get("ActualMW", ""),
        "deviation_mw": raw.get("ActualDeviationMW", ""),
        "ace_mw": raw.get("AreaControlErrorMW", ""),
        "desired": clean_text(raw.get("Desired")),
        "raw": raw,
    }


def violation_search_query(search: str, constituent: str) -> dict:
    clauses = []
    if constituent:
        clauses.append({"Constituent": constituent})
    if search:
        pattern = re.escape(search.strip())
        clauses.append({
            "$or": [
                {field: {"$regex": pattern, "$options": "i"}}
                for field in (
                    "Message", "Constituent", "ViolationType", "SubViolationType",
                    "CreatedDate", "CreatedTime", "RequestId", "LogbookId", "Desired",
                )
            ]
        })
    return {"$and": clauses} if clauses else {}


def violation_bucket_key(day: date, grouping: str) -> tuple:
    if grouping == "monthly":
        return day.strftime("%Y-%m"), day.strftime("%b %Y")
    if grouping == "weekly":
        week_start = day.fromordinal(day.toordinal() - day.weekday())
        week_end = week_start.fromordinal(week_start.toordinal() + 6)
        return week_start.isoformat(), f"{week_start.strftime('%d %b')} - {week_end.strftime('%d %b %Y')}"
    return day.isoformat(), day.strftime("%d %b %Y")


def violation_summary(docs: List[dict], grouping: str) -> List[dict]:
    buckets: Dict[str, dict] = {}
    for doc in docs:
        day = parse_logbook_date(doc.get("CreatedDate"))
        if not day:
            continue
        key, label = violation_bucket_key(day, grouping)
        bucket = buckets.setdefault(key, {
            "key": key,
            "label": label,
            "count": 0,
            "constituents": set(),
            "types": Counter(),
        })
        bucket["count"] += 1
        constituent = clean_text(doc.get("Constituent"))
        if constituent:
            bucket["constituents"].add(constituent)
        violation_type = clean_text(doc.get("ViolationType")) or "Unspecified"
        bucket["types"][violation_type] += 1
    return [
        {
            "key": item["key"],
            "label": item["label"],
            "count": item["count"],
            "constituent_count": len(item["constituents"]),
            "types": dict(item["types"].most_common()),
        }
        for item in sorted(buckets.values(), key=lambda row: row["key"], reverse=True)
    ]


def violation_matrix(docs: List[dict], grouping: str) -> dict:
    column_totals = Counter()
    buckets: Dict[str, dict] = {}
    for doc in docs:
        day = parse_logbook_date(doc.get("CreatedDate"))
        if not day:
            continue
        key, label = violation_bucket_key(day, grouping)
        constituent = clean_text(doc.get("Constituent")) or "Unspecified"
        bucket = buckets.setdefault(key, {"key": key, "label": label, "values": Counter(), "total": 0})
        bucket["values"][constituent] += 1
        bucket["total"] += 1
        column_totals[constituent] += 1
    columns = [
        {"key": name, "label": name, "total": total}
        for name, total in sorted(column_totals.items(), key=lambda item: item[0].lower())
    ]
    rows = [
        {
            "key": item["key"],
            "label": item["label"],
            "values": dict(item["values"]),
            "total": item["total"],
        }
        for item in sorted(buckets.values(), key=lambda row: row["key"], reverse=True)
    ]
    return {"grouping": grouping, "columns": columns, "rows": rows, "grand_total": sum(column_totals.values())}


def filtered_violations(
    collection,
    start_date: date,
    end_date: date,
    search: str,
    constituent: str,
    violation_types: Optional[List[str]] = None,
    sub_violation_types: Optional[List[str]] = None,
) -> List[dict]:
    selected_types = {clean_text(value) for value in (violation_types or []) if clean_text(value)}
    selected_subtypes = {clean_text(value) for value in (sub_violation_types or []) if clean_text(value)}
    docs = []
    for doc in collection.find(violation_search_query(search, constituent)):
        created = parse_logbook_date(doc.get("CreatedDate"))
        if not created or not start_date <= created <= end_date:
            continue
        violation_type = clean_text(doc.get("ViolationType")) or "Unspecified"
        subtype = clean_text(doc.get("SubViolationType")) or "__NONE__"
        if selected_types and violation_type not in selected_types:
            continue
        if selected_subtypes and subtype not in selected_subtypes:
            continue
        docs.append(doc)
    docs.sort(key=violation_datetime, reverse=True)
    return docs


def segregated_violation_matrices(docs: List[dict], grouping: str) -> List[dict]:
    groups: Dict[tuple, List[dict]] = {}
    for doc in docs:
        violation_type = clean_text(doc.get("ViolationType")) or "Unspecified"
        subtype = clean_text(doc.get("SubViolationType")) or "No subtype"
        groups.setdefault((violation_type, subtype), []).append(doc)
    result = []
    for (violation_type, subtype), group_docs in sorted(groups.items()):
        result.append({
            "key": f"{violation_type}|{subtype}",
            "label": f"{violation_type} · {subtype}",
            "violation_type": violation_type,
            "sub_violation_type": subtype,
            "matrix": violation_matrix(group_docs, grouping),
        })
    return result


def in_date_range(doc: dict, date_key: str, start_date: date, end_date: date) -> bool:
    parsed = parse_logbook_date(doc.get(date_key))
    return parsed is not None and start_date <= parsed <= end_date


def normalize_record(doc: dict, kind: str) -> dict:
    config = COLLECTION_CONFIG[kind]
    raw = to_jsonable(doc)
    reason_fields = config.get("reason_fields", [])
    element_type_raw = clean_text(raw.get(config.get("element_type", "Type")))
    normalized = {
        "id": clean_text(raw.get("_id")),
        "kind": kind,
        "kind_label": config["label"],
        "element_name": clean_text(raw.get(config["element"])),
        "revival_time": combine_datetime(raw, config["revival_date"], config["revival_time"]),
        "outage_time": combine_datetime(raw, config["outage_date"], config["outage_time"]),
        "reason": combine_fields(raw, reason_fields),
        "reason_label": config.get("reason_label", "Reason"),
        "sub_category": clean_text(raw.get("Nature")),
        "element_type": display_element_type(element_type_raw),
        "element_type_key": normalize_element_type_key(element_type_raw),
        "audit_history": raw.get("AuditHistory", ""),
        "planned_period": "",
        "raw": raw,
    }
    if config.get("planned_fields"):
        normalized["planned_period"] = combine_fields(raw, config["planned_fields"], " to ")
    return normalized


def analyze_attributes(docs: List[dict]) -> List[dict]:
    stats: Dict[str, Dict[str, Any]] = {}
    total = len(docs)
    for doc in docs:
        for key, value in doc.items():
            text = clean_text(value)
            if key not in stats:
                stats[key] = {
                    "name": key,
                    "present_count": 0,
                    "filled_count": 0,
                    "types": set(),
                    "samples": [],
                }
            item = stats[key]
            item["present_count"] += 1
            item["types"].add(type(value).__name__)
            if text:
                item["filled_count"] += 1
                if len(item["samples"]) < 3 and text not in item["samples"]:
                    item["samples"].append(text[:160])

    result = []
    for item in stats.values():
        result.append({
            "name": item["name"],
            "present_count": item["present_count"],
            "filled_count": item["filled_count"],
            "missing_count": max(0, total - item["present_count"]),
            "types": sorted(item["types"]),
            "samples": item["samples"],
        })
    return sorted(result, key=lambda row: row["name"].lower())


def old_logbook_query(search: str, config: dict) -> dict:
    search = (search or "").strip()
    if not search:
        return {}
    fields = [
        config["element"],
        "Nature",
        "Type",
        "EntityId",
        "Reason",
        "EndRelayReasonOne",
        "EndRelayReasonTwo",
        "AuditHistory",
    ]
    return {
        "$or": [
            {field: {"$regex": search, "$options": "i"}}
            for field in fields
        ]
    }


def iter_filtered_docs(collection, config: dict, search: str, element_type: str, start_date: date, end_date: date):
    element_type_filter = normalize_element_type_key(element_type)
    for doc in collection.find(old_logbook_query(search, config)):
        if not in_date_range(doc, config["outage_date"], start_date, end_date):
            continue
        doc_element_type = normalize_element_type_key(doc.get(config.get("element_type", "Type")))
        if element_type_filter and doc_element_type != element_type_filter:
            continue
        yield doc


def collection_payload(db, kind: str, search: str, element_type: str, limit: int, skip: int, start_date: date, end_date: date) -> dict:
    config = COLLECTION_CONFIG[kind]
    collection = db[config["collection"]]
    page_start = max(0, skip)
    page_end = page_start + max(1, min(limit, 5000))
    total_count = 0
    docs = []
    analysis_docs = []
    for doc in iter_filtered_docs(collection, config, search, element_type, start_date, end_date):
        if len(analysis_docs) < 1000:
            analysis_docs.append(doc)
        if page_start <= total_count < page_end:
            docs.append(doc)
        total_count += 1
    rows = [normalize_record(doc, kind) for doc in docs]
    elements = sorted({
        row["element_name"]
        for row in rows
        if row.get("element_name")
    })
    element_types = sorted({
        row["element_type_key"]
        for row in rows
        if row.get("element_type_key")
    })

    return {
        "kind": kind,
        "label": config["label"],
        "collection": config["collection"],
        "total_count": total_count,
        "returned_count": len(rows),
        "rows": rows,
        "elements": elements,
        "element_types": element_types,
        "attributes": analyze_attributes([to_jsonable(doc) for doc in analysis_docs]),
    }


def export_row(row: dict) -> dict:
    audit = row.get("audit_history")
    if isinstance(audit, (dict, list)):
        audit = json.dumps(to_jsonable(audit), ensure_ascii=False)
    return {
        "Outage Type": row.get("kind_label", ""),
        "Element Name": row.get("element_name", ""),
        "Revival Time": row.get("revival_time", ""),
        "Outage Time": row.get("outage_time", ""),
        "Reason / Relay": row.get("reason", ""),
        "Sub Category": row.get("sub_category", ""),
        "Element Type": row.get("element_type", ""),
        "Planned Period": row.get("planned_period", ""),
        "AuditHistory": clean_text(audit),
    }


@router.get("/historical-outages")
async def get_historical_outages(
    kind: str = Query("all", pattern="^(all|shutdown|tripping|outage)$"),
    start_date: date = Query(...),
    end_date: date = Query(...),
    search: str = "",
    element_type: str = "",
    limit: int = Query(1000, ge=1, le=5000),
    skip: int = Query(0, ge=0),
):
    if end_date < start_date:
        start_date, end_date = end_date, start_date
    service = MongoService()
    db = service.client[OLD_LOGBOOK_DB]
    selected = list(COLLECTION_CONFIG.keys()) if kind == "all" else [kind]
    sections = {
        item: collection_payload(db, item, search, element_type, limit, skip, start_date, end_date)
        for item in selected
    }
    return {
        "success": True,
        "database": OLD_LOGBOOK_DB,
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat(),
        "sections": sections,
    }


@router.get("/historical-outages/export")
async def export_historical_outages(
    kind: str = Query("all", pattern="^(all|shutdown|tripping|outage)$"),
    start_date: date = Query(...),
    end_date: date = Query(...),
    search: str = "",
    element_type: str = "",
):
    if end_date < start_date:
        start_date, end_date = end_date, start_date
    service = MongoService()
    db = service.client[OLD_LOGBOOK_DB]
    selected = list(COLLECTION_CONFIG.keys()) if kind == "all" else [kind]

    workbook = Workbook()
    workbook.remove(workbook.active)
    columns = [
        "Outage Type",
        "Element Name",
        "Revival Time",
        "Outage Time",
        "Reason / Relay",
        "Sub Category",
        "Element Type",
        "Planned Period",
        "AuditHistory",
    ]

    for item in selected:
        config = COLLECTION_CONFIG[item]
        worksheet = workbook.create_sheet(config["label"][:31])
        worksheet.append(columns)
        collection = db[config["collection"]]
        for doc in iter_filtered_docs(collection, config, search, element_type, start_date, end_date):
            normalized = normalize_record(doc, item)
            row = export_row(normalized)
            worksheet.append([row.get(column, "") for column in columns])
        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells[:200])
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 14), 42)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"old_logbook_{kind}_{start_date.isoformat()}_to_{end_date.isoformat()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/violation-messages/export-pdf")
async def export_violation_matrices_pdf(
    start_date: date = Query(...),
    end_date: date = Query(...),
    search: str = "",
    constituent: str = "",
    violation_type: Optional[List[str]] = Query(default=None),
    sub_violation_type: Optional[List[str]] = Query(default=None),
    matrix_grouping: str = Query("daily", pattern="^(daily|monthly)$"),
):
    if end_date < start_date:
        start_date, end_date = end_date, start_date
    service = MongoService()
    collection = service.client[OLD_LOGBOOK_DB][VIOLATION_COLLECTION]
    docs = filtered_violations(
        collection, start_date, end_date, search, constituent,
        violation_type, sub_violation_type,
    )
    sections = segregated_violation_matrices(docs, matrix_grouping)

    output = io.BytesIO()
    page_width, page_height = landscape(A3)
    pdf = canvas.Canvas(output, pagesize=(page_width, page_height))
    pdf.setTitle("Violation Message Matrices")
    margin = 10 * mm
    logo_path = Path(__file__).resolve().parents[2] / "frontend" / "public" / "logo.png"
    header_line_1 = "Grid Controller of India Limited (GRID-INDIA)"
    header_line_2 = "EASTERN REGIONAL LOAD DESPATCH CENTRE, KOLKATA"

    def draw_page_header(section: dict):
        header_top = page_height - margin
        if logo_path.exists():
            pdf.drawImage(
                str(logo_path), margin, header_top - 22 * mm,
                width=34 * mm, height=18 * mm,
                preserveAspectRatio=True, anchor="c", mask="auto",
            )
        pdf.setFillColor(colors.HexColor("#003366"))
        pdf.setFont("Helvetica-Bold", 13)
        pdf.drawCentredString(page_width / 2, header_top - 6 * mm, header_line_1)
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawCentredString(page_width / 2, header_top - 12 * mm, header_line_2)
        pdf.setStrokeColor(colors.HexColor("#334155"))
        pdf.setLineWidth(0.7)
        pdf.line(margin, header_top - 23 * mm, page_width - margin, header_top - 23 * mm)

        pdf.setFillColor(colors.HexColor("#0B55B8"))
        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawCentredString(page_width / 2, header_top - 31 * mm, section["label"])
        pdf.setFillColor(colors.HexColor("#475569"))
        pdf.setFont("Helvetica", 7.5)
        matrix = section["matrix"]
        subtitle = (
            f"{start_date.isoformat()} to {end_date.isoformat()} · "
            f"{matrix_grouping.title()} grouping · {matrix['grand_total']:,} messages"
        )
        pdf.drawCentredString(page_width / 2, header_top - 36 * mm, subtitle)
        return header_top - 42 * mm

    for section_index, section in enumerate(sections):
        table_top = draw_page_header(section)
        matrix = section["matrix"]
        columns = matrix["columns"]
        table_data = [["Period", *[column["label"] for column in columns], "Total"]]
        for row in matrix["rows"]:
            table_data.append([
                row["label"],
                *[row["values"].get(column["key"], 0) for column in columns],
                row["total"],
            ])
        table_data.append(["Total", *[column["total"] for column in columns], matrix["grand_total"]])

        table_width = page_width - (2 * margin)
        table_height_limit = table_top - margin
        period_width = min(36 * mm, table_width * 0.14)
        total_width = min(18 * mm, table_width * 0.07)
        data_width = (table_width - period_width - total_width) / max(1, len(columns))
        header_height = min(15 * mm, max(7 * mm, table_height_limit * 0.08))
        data_row_height = min(6.5 * mm, (table_height_limit - header_height) / max(1, len(table_data) - 1))
        font_size = max(3.0, min(7.0, data_row_height * 0.42, data_width * 0.18))
        header_font_size = max(3.0, min(6.5, font_size))
        header_style = ParagraphStyle(
            "matrix_header",
            fontName="Helvetica-Bold",
            fontSize=header_font_size,
            leading=max(3.2, header_font_size + 0.4),
            textColor=colors.white,
            alignment=1,
            wordWrap="CJK",
        )
        table_data[0] = [Paragraph(str(value), header_style) for value in table_data[0]]
        row_heights = [header_height, *([data_row_height] * (len(table_data) - 1))]
        table = Table(
            table_data,
            colWidths=[period_width, *([data_width] * len(columns)), total_width],
            rowHeights=row_heights,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B55B8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#DDF1EA")),
            ("BACKGROUND", (-1, 1), (-1, -1), colors.HexColor("#DDF1EA")),
            ("ROWBACKGROUNDS", (0, 1), (-2, -2), [colors.white, colors.HexColor("#F6F9FD")]),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9DB7D5")),
            ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("FONTNAME", (-1, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 1), (-1, -1), font_size),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("ALIGN", (0, 0), (0, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 1.5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 1.5),
            ("TOPPADDING", (0, 0), (-1, -1), 0.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0.5),
        ]))
        rendered_width, rendered_height = table.wrap(table_width, table_height_limit)
        table.drawOn(pdf, margin, table_top - rendered_height)
        if section_index < len(sections) - 1:
            pdf.showPage()
    if not sections:
        empty_section = {"label": "VIOLATION MESSAGE MATRIX", "matrix": {"grand_total": 0}}
        table_top = draw_page_header(empty_section)
        pdf.setFillColor(colors.HexColor("#64748B"))
        pdf.setFont("Helvetica", 10)
        pdf.drawCentredString(page_width / 2, table_top - 15 * mm, "No violation messages match the selected filters.")
    pdf.save()
    output.seek(0)
    filename = f"violation_matrices_{start_date.isoformat()}_to_{end_date.isoformat()}.pdf"
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/violation-messages/meta")
async def get_violation_message_meta():
    service = MongoService()
    collection = service.client[OLD_LOGBOOK_DB][VIOLATION_COLLECTION]
    constituents = sorted(value for value in collection.distinct("Constituent") if clean_text(value))
    violation_types = sorted(clean_text(value) for value in collection.distinct("ViolationType") if clean_text(value))
    raw_subtypes = {clean_text(value) for value in collection.distinct("SubViolationType")}
    sub_violation_types = sorted(value for value in raw_subtypes if value)
    if "" in raw_subtypes:
        sub_violation_types.append("__NONE__")
    available_dates = [
        parsed
        for value in collection.distinct("CreatedDate")
        if (parsed := parse_logbook_date(value)) is not None
    ]
    return {
        "success": True,
        "total_count": collection.count_documents({}),
        "constituents": constituents,
        "violation_types": violation_types,
        "sub_violation_types": sub_violation_types,
        "min_date": min(available_dates).isoformat() if available_dates else "",
        "max_date": max(available_dates).isoformat() if available_dates else "",
    }


@router.get("/violation-messages")
async def get_violation_messages(
    start_date: date = Query(...),
    end_date: date = Query(...),
    search: str = "",
    constituent: str = "",
    violation_type: Optional[List[str]] = Query(default=None),
    sub_violation_type: Optional[List[str]] = Query(default=None),
    grouping: str = Query("daily", pattern="^(daily|weekly|monthly)$"),
    matrix_grouping: str = Query("daily", pattern="^(daily|monthly)$"),
    limit: int = Query(100, ge=1, le=1000),
    skip: int = Query(0, ge=0),
):
    if end_date < start_date:
        start_date, end_date = end_date, start_date
    service = MongoService()
    collection = service.client[OLD_LOGBOOK_DB][VIOLATION_COLLECTION]
    docs = filtered_violations(
        collection, start_date, end_date, search, constituent,
        violation_type, sub_violation_type,
    )
    type_counts = Counter(clean_text(doc.get("ViolationType")) or "Unspecified" for doc in docs)
    constituent_counts = Counter(clean_text(doc.get("Constituent")) or "Unspecified" for doc in docs)
    page_docs = docs[skip:skip + limit]
    return {
        "success": True,
        "database": OLD_LOGBOOK_DB,
        "collection": VIOLATION_COLLECTION,
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat(),
        "total_count": len(docs),
        "returned_count": len(page_docs),
        "rows": [normalize_violation(doc) for doc in page_docs],
        "summary": violation_summary(docs, grouping),
        "matrix": violation_matrix(docs, matrix_grouping),
        "segregated_matrices": segregated_violation_matrices(docs, matrix_grouping),
        "type_counts": dict(type_counts.most_common()),
        "constituent_counts": dict(constituent_counts.most_common()),
    }


@router.get("/violation-messages/export")
async def export_violation_messages(
    start_date: date = Query(...),
    end_date: date = Query(...),
    search: str = "",
    constituent: str = "",
    violation_type: Optional[List[str]] = Query(default=None),
    sub_violation_type: Optional[List[str]] = Query(default=None),
    grouping: str = Query("daily", pattern="^(daily|weekly|monthly)$"),
    matrix_grouping: str = Query("daily", pattern="^(daily|monthly)$"),
):
    if end_date < start_date:
        start_date, end_date = end_date, start_date
    service = MongoService()
    collection = service.client[OLD_LOGBOOK_DB][VIOLATION_COLLECTION]
    docs = filtered_violations(
        collection, start_date, end_date, search, constituent,
        violation_type, sub_violation_type,
    )

    workbook = Workbook()
    details = workbook.active
    details.title = "Violation Messages"
    columns = [
        "Date", "Time", "Constituent", "Violation Type", "Sub Violation Type", "Message",
        "Frequency", "Schedule MW", "Actual MW", "Deviation MW", "ACE MW", "Desired",
        "Message ID", "Request ID", "Logbook ID",
    ]
    details.append(columns)
    for doc in docs:
        row = normalize_violation(doc)
        details.append([
            row["date"], row["time"], row["constituent"], row["violation_type"],
            row["sub_violation_type"], row["message"], row["frequency"], row["schedule_mw"],
            row["actual_mw"], row["deviation_mw"], row["ace_mw"], row["desired"],
            row["message_id"], row["request_id"], row["logbook_id"],
        ])
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    detail_widths = [13, 10, 18, 23, 20, 58, 12, 14, 14, 14, 14, 22, 12, 38, 38]
    for index, width in enumerate(detail_widths, start=1):
        details.column_dimensions[details.cell(1, index).column_letter].width = width

    summary_sheet = workbook.create_sheet("Summary")
    summary_sheet.append(["Period", "Messages", "Constituents", "Violation Type Breakdown"])
    for item in violation_summary(docs, grouping):
        breakdown = ", ".join(f"{name}: {count}" for name, count in item["types"].items())
        summary_sheet.append([item["label"], item["count"], item["constituent_count"], breakdown])
    summary_sheet.freeze_panes = "A2"
    for column, width in zip(("A", "B", "C", "D"), (28, 14, 16, 65)):
        summary_sheet.column_dimensions[column].width = width

    segregated = segregated_violation_matrices(docs, matrix_grouping)
    matrix_sheet = workbook.create_sheet(f"{matrix_grouping.title()} Matrices")
    matrix_sheet.append([f"SEGREGATED {matrix_grouping.upper()} VIOLATION MATRICES"])
    for section in segregated:
        matrix = section["matrix"]
        matrix_columns = matrix["columns"]
        matrix_sheet.append([])
        matrix_sheet.append([section["label"]])
        matrix_sheet.append(["Period", *[item["label"] for item in matrix_columns], "Total"])
        for item in matrix["rows"]:
            matrix_sheet.append([
                item["label"],
                *[item["values"].get(column["key"], 0) for column in matrix_columns],
                item["total"],
            ])
        matrix_sheet.append(["Total", *[item["total"] for item in matrix_columns], matrix["grand_total"]])
    matrix_sheet.freeze_panes = "B2"
    matrix_sheet.column_dimensions["A"].width = 28
    for cells in matrix_sheet.iter_cols(min_col=2, max_col=matrix_sheet.max_column):
        matrix_sheet.column_dimensions[cells[0].column_letter].width = 15

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"violation_messages_{start_date.isoformat()}_to_{end_date.isoformat()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
