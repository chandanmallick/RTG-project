import io
import os
import posixpath
import re
import statistics
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape as xml_escape

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

from routes.dso_report_routes import master_collection
from services.db_handler import MongoService


router = APIRouter(prefix="/api/sri", tags=["System Reliability Report"])

SRI_COLLECTION = "sri_reports"
DEFAULT_SOURCE_PATH = r"\\10.3.95.200\HTTP-Access\ScadaData\Web_based_PSP"
DEFAULT_TEMPLATE_PATHS = (
    Path(r"C:\Users\crew-admin\Downloads\SRI\data\SRI 19072026.xlsx"),
    Path(r"C:\Users\crew-admin\Downloads\SRI 19072026.xlsx"),
)
STATE_ORDER = ("BIHAR", "JHARKHAND", "DVC", "ODISHA", "WB", "SIKKIM")
STATE_LABELS = {
    "BIHAR": "Bihar", "JHARKHAND": "Jharkhand", "DVC": "DVC",
    "ODISHA": "Odisha", "WB": "West Bengal", "SIKKIM": "Sikkim",
}
STATE_KEYS = {
    "BIHAR": "04247020", "DVC": "04247030", "JHARKHAND": "04247040",
    "ODISHA": "04247060", "WB": "04247010", "SIKKIM": "04247050",
}


def sri_collection():
    result = MongoService().db[SRI_COLLECTION]
    result.create_index("report_date", unique=True)
    return result


def source_root() -> Path:
    return Path(str(os.getenv("SRI_SOURCE_PATH") or DEFAULT_SOURCE_PATH).strip())


def source_file(report_date: str) -> Path:
    selected = date.fromisoformat(report_date)
    folder = source_root() / f"{selected.year:04d}" / f"{selected.month:02d}"
    expected = folder / f"ER_PSP_State_drwl_and_demand_{selected:%d%m%Y}.xlsx"
    if not expected.is_file():
        raise FileNotFoundError(f"SRI source workbook was not found: {expected}")
    return expected


def template_file() -> Path:
    configured = str(os.getenv("SRI_TEMPLATE_PATH") or "").strip()
    candidates = ([Path(configured)] if configured else []) + list(DEFAULT_TEMPLATE_PATHS)
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise FileNotFoundError(
        "SRI layout template was not found. Configure SRI_TEMPLATE_PATH to the supplied 'SRI 19072026.xlsx' file."
    )


def number(value):
    if value in (None, ""):
        return None
    try:
        parsed = float(value)
        return parsed if parsed == parsed else None
    except (TypeError, ValueError):
        return None


def excel_timestamp(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)):
        return datetime(1899, 12, 30) + timedelta(days=float(value))
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d-%m-%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(text[:19], pattern)
        except ValueError:
            continue
    return None


def read_state_drawl(path: Path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = next((item for item in workbook.worksheets if not item.title.startswith("_")), None)
    if sheet is None:
        raise ValueError("No data worksheet was found in the SRI source workbook.")
    key_columns = {}
    key_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True), ())
    for column, raw_key in enumerate(key_row[1:], 2):
        key = str(raw_key or "").strip()
        for state, expected_key in STATE_KEYS.items():
            if key == expected_key:
                key_columns[state] = column
    missing = [f"{STATE_LABELS[state]} ({STATE_KEYS[state]})" for state in STATE_ORDER if state not in key_columns]
    if missing:
        raise ValueError("Required state drawal SCADA columns were not found: " + ", ".join(missing))
    rows = []
    for values in sheet.iter_rows(min_row=3, values_only=True):
        timestamp = excel_timestamp(values[0] if values else None)
        if timestamp is None:
            continue
        state_values = {state: number(values[column - 1] if column <= len(values) else None) for state, column in key_columns.items()}
        if any(value is not None for value in state_values.values()):
            rows.append({"timestamp": timestamp, "values": state_values})
    workbook.close()
    rows.sort(key=lambda item: item["timestamp"])
    if not rows:
        raise ValueError("The SRI source workbook contains no usable drawal samples.")
    return rows


def sample_minutes(rows):
    differences = [
        (current["timestamp"] - previous["timestamp"]).total_seconds() / 60
        for previous, current in zip(rows, rows[1:])
        if 0 < (current["timestamp"] - previous["timestamp"]).total_seconds() <= 21600
    ]
    return max(1 / 60, min(float(statistics.median(differences)) if differences else 1.0, 60.0))


def hours_text(hours):
    total_minutes = int(round(max(0.0, float(hours or 0)) * 60))
    return f"{total_minutes // 60:02d}:{total_minutes % 60:02d}"


def violates(actual, limit):
    if actual is None or limit is None:
        return False
    return actual < limit if limit < 0 else actual > limit


def default_corrective(kind: str, violating_states: list[str]):
    prefix = "सुधारात्मक कार्यों की आवश्यकता/ Corrective actions required:"
    if not violating_states:
        return f"{prefix} NA"
    names = ", ".join(STATE_LABELS[state] for state in violating_states)
    return f"{prefix} {names} are advised to control drawal within {kind} limit."


def exceedance(actual, limit):
    if actual is None or limit is None:
        return 0.0
    return max(0.0, (limit - actual) if limit < 0 else (actual - limit))


def violation_events(rows, limits, interval_minutes):
    events = []
    nominal_delta = timedelta(minutes=interval_minutes * 1.5)
    for state in STATE_ORDER:
        for kind in ("TTC", "ATC"):
            limit = limits[state][kind.lower()]
            active = []

            def close_event():
                if not active:
                    return
                values = [item[1] for item in active]
                peak = min(values) if limit is not None and limit < 0 else max(values)
                maximum = max(exceedance(value, limit) for value in values)
                denominator = abs(limit or 0)
                events.append({
                    "state": state, "state_label": STATE_LABELS[state], "limit": kind,
                    "direction": "Export" if limit is not None and limit < 0 else "Import",
                    "start": active[0][0].isoformat(timespec="minutes"),
                    "end": active[-1][0].isoformat(timespec="minutes"),
                    "duration_hours": round(len(active) * interval_minutes / 60, 4),
                    "peak_drawl_mw": round(peak, 3), "applicable_limit_mw": limit,
                    "max_exceedance_mw": round(maximum, 3),
                    "max_exceedance_percent": round(maximum / denominator * 100 if denominator else 0, 4),
                })
                active.clear()

            previous = None
            for item in rows:
                timestamp, actual = item["timestamp"], item["values"].get(state)
                if violates(actual, limit):
                    if previous is not None and timestamp - previous > nominal_delta:
                        close_event()
                    active.append((timestamp, actual))
                    previous = timestamp
                else:
                    close_event()
                    previous = None
            close_event()
    return sorted(events, key=lambda item: (item["start"], item["state"], item["limit"]))


def calculate_report(report_date: str):
    path = source_file(report_date)
    rows = read_state_drawl(path)
    interval_minutes = sample_minutes(rows)
    limits_doc = master_collection().find_one({"config_type": "DSO_TTC_ATC"}, {"_id": 0}) or {}
    raw_limits = limits_doc.get("limits") or {}
    limits = {
        state: {"ttc": number((raw_limits.get(state) or {}).get("ttc")), "atc": number((raw_limits.get(state) or {}).get("atc"))}
        for state in STATE_ORDER
    }
    sections = {}
    plots = {}
    quality = []
    for state in STATE_ORDER:
        plots[state] = []
        for item in rows:
            actual = item["values"].get(state)
            plots[state].append({
                "timestamp": item["timestamp"].isoformat(timespec="minutes"),
                "actual_mw": round(actual, 3) if actual is not None else None,
                "ttc_mw": limits[state]["ttc"],
                "atc_mw": limits[state]["atc"],
                "ttc_violation": violates(actual, limits[state]["ttc"]),
                "atc_violation": violates(actual, limits[state]["atc"]),
            })
        monitored = sum(item["values"].get(state) is not None for item in rows)
        quality.append({
            "state": state, "state_label": STATE_LABELS[state], "monitored_samples": monitored,
            "missing_samples": len(rows) - monitored,
            "coverage_percent": round(monitored / len(rows) * 100 if rows else 0, 4),
        })
    for kind in ("TTC", "ATC"):
        result_rows = []
        violating_states = []
        limit_key = kind.lower()
        for state in STATE_ORDER:
            actual_values = [item["values"].get(state) for item in rows]
            monitored = sum(value is not None for value in actual_values)
            limit = limits[state][limit_key]
            count = sum(violates(value, limit) for value in actual_values)
            hours = count * interval_minutes / 60
            percent = count / monitored * 100 if monitored else 0
            if count:
                violating_states.append(state)
            result_rows.append({
                "state": state,
                "state_label": STATE_LABELS[state],
                "corridor": f"ER - {STATE_LABELS[state]}",
                "direction": "Export" if limit is not None and limit < 0 else "Import",
                "limit_mw": limit,
                "monitored_samples": monitored,
                "violation_samples": count,
                "violation_hours": round(hours, 4),
                "violation_time": hours_text(hours),
                "violation_percent": round(percent, 4),
                "intimation": "Review required" if count else "NA",
                "max_exceedance_mw": round(max((exceedance(value, limit) for value in actual_values), default=0), 3),
            })
        sections[kind] = {"rows": result_rows, "corrective_action": default_corrective(kind, violating_states)}
    events = violation_events(rows, limits, interval_minutes)
    summary = []
    for state in STATE_ORDER:
        ttc = next(row for row in sections["TTC"]["rows"] if row["state"] == state)
        atc = next(row for row in sections["ATC"]["rows"] if row["state"] == state)
        state_quality = next(row for row in quality if row["state"] == state)
        summary.append({
            "state": state, "state_label": STATE_LABELS[state],
            "ttc_hours": ttc["violation_hours"], "ttc_percent": ttc["violation_percent"],
            "max_ttc_exceedance_mw": ttc["max_exceedance_mw"],
            "atc_hours": atc["violation_hours"], "atc_percent": atc["violation_percent"],
            "max_atc_exceedance_mw": atc["max_exceedance_mw"],
            "limit_coverage_percent": state_quality["coverage_percent"],
        })
    schedule = [{
        "state": state, "state_label": STATE_LABELS[state],
        "direction": "Export" if (limits[state]["ttc"] or 0) < 0 else "Import",
        "date_range": datetime.strptime(report_date, "%Y-%m-%d").strftime("%d-%b-%Y"),
        "time_period": "00-24 hrs", "ttc_mw": limits[state]["ttc"], "atc_mw": limits[state]["atc"],
        "source": "DSO TTC/ATC Master",
    } for state in STATE_ORDER]
    return {
        "report_date": report_date,
        "title": f"SYSTEM RELIABILITY REPORT for {datetime.strptime(report_date, '%Y-%m-%d'):%d-%m-%Y}",
        "source_file": str(path),
        "sample_interval_minutes": round(interval_minutes, 4),
        "sample_count": len(rows),
        "coverage_start": rows[0]["timestamp"].isoformat(timespec="minutes"),
        "coverage_end": rows[-1]["timestamp"].isoformat(timespec="minutes"),
        "limits_updated_at": limits_doc.get("updated_at"),
        "limits": limits,
        "sections": sections,
        "plots": plots,
        "violation_summary": summary,
        "violation_events": events,
        "capability_schedule": schedule,
        "data_quality": quality,
    }


def apply_saved_edits(report: dict):
    saved = sri_collection().find_one({"report_date": report["report_date"]}, {"_id": 0}) or {}
    edits = saved.get("edits") or {}
    for kind in ("TTC", "ATC"):
        section_edits = edits.get(kind) or {}
        row_edits = section_edits.get("rows") or {}
        for row in report["sections"][kind]["rows"]:
            values = row_edits.get(row["state"]) or {}
            for field in ("violation_hours", "violation_percent", "intimation"):
                if field in values:
                    row[field] = values[field]
        if "corrective_action" in section_edits:
            report["sections"][kind]["corrective_action"] = section_edits["corrective_action"]
    for item in report.get("violation_summary") or []:
        for kind in ("TTC", "ATC"):
            row = next(entry for entry in report["sections"][kind]["rows"] if entry["state"] == item["state"])
            item[f"{kind.lower()}_hours"] = row["violation_hours"]
            item[f"{kind.lower()}_percent"] = row["violation_percent"]
    for kind in ("TTC", "ATC"):
        for row in report["sections"][kind]["rows"]:
            row["violation_time"] = hours_text(row["violation_hours"])
    report["has_saved_edits"] = bool(edits)
    report["saved_at"] = saved.get("updated_at")
    return report


def report_for_date(report_date: str):
    try:
        date.fromisoformat(report_date)
        return apply_saved_edits(calculate_report(report_date))
    except FileNotFoundError as exc:
        raise HTTPException(404, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except OSError as exc:
        raise HTTPException(502, f"SRI source folder is unavailable: {exc}") from exc


@router.get("/report")
def get_sri_report(report_date: str = Query(...)):
    return {"success": True, "report": report_for_date(report_date)}


@router.put("/report/{report_date}")
def save_sri_report(report_date: str, payload: dict):
    report = report_for_date(report_date)
    supplied = payload.get("sections") or {}
    edits = {}
    for kind in ("TTC", "ATC"):
        incoming = supplied.get(kind) or {}
        rows = {}
        for item in incoming.get("rows") or []:
            state = str(item.get("state") or "").upper()
            if state not in STATE_ORDER:
                continue
            rows[state] = {
                "violation_hours": max(0, float(item.get("violation_hours") or 0)),
                "violation_percent": max(0, min(100, float(item.get("violation_percent") or 0))),
                "intimation": str(item.get("intimation") or "NA").strip() or "NA",
            }
        edits[kind] = {
            "rows": rows,
            "corrective_action": str(incoming.get("corrective_action") or report["sections"][kind]["corrective_action"]).strip(),
        }
    updated_at = datetime.utcnow().isoformat()
    sri_collection().update_one(
        {"report_date": report_date},
        {"$set": {"report_date": report_date, "edits": edits, "updated_at": updated_at}},
        upsert=True,
    )
    return {"success": True, "report": apply_saved_edits(calculate_report(report_date))}


@router.delete("/report/{report_date}/edits")
def reset_sri_report(report_date: str):
    sri_collection().delete_one({"report_date": report_date})
    return {"success": True, "report": calculate_report(report_date)}


def build_excel(report: dict):
    workbook = Workbook()
    workbook.remove(workbook.active)
    navy, blue, pale, border_color = "17365D", "D9EAF7", "EDF4FB", "8EA9C1"
    labels = {
        "TTC": (
            "टीटीसी उल्लंघन की रिपोर्टिंग / Reporting of TTC Violation",
            "In case of violation of TTC, regional entities in the importing/exporting control area shall restrict drawal or injection so that TTC is maintained within limit and SLDCs shall take appropriate action.",
        ),
        "ATC": (
            "एटीसी उल्लंघन की रिपोर्टिंग / Reporting of ATC Violation",
            "In case of violation of ATC, regional entities in the importing/exporting control area shall restrict drawal or injection so that ATC is maintained within limit and SLDCs shall take appropriate action.",
        ),
    }
    for kind in ("TTC", "ATC"):
        sheet = workbook.create_sheet(f"{kind} Violation")
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells("A1:F1"); sheet["A1"] = "GRID CONTROLLER OF INDIA LIMITED\n(A Government of India Enterprise)\nEASTERN REGIONAL LOAD DESPATCH CENTRE, KOLKATA"
        sheet.merge_cells("A2:F2"); sheet["A2"] = report["title"]
        sheet.merge_cells("A3:F3"); sheet["A3"] = labels[kind][0]
        headers = ["Sl No.", "Intra-Regional Corridor", "Import / Export", f"Total Hours of {kind} Violation", f"Percentage Time of {kind} Violation", "Intimation for Corrective Action"]
        for column, value in enumerate(headers, 1):
            sheet.cell(4, column, value)
        for index, row in enumerate(report["sections"][kind]["rows"], 1):
            values = [index, row["corridor"], row["direction"], row["violation_hours"], row["violation_percent"], row["intimation"]]
            for column, value in enumerate(values, 1):
                sheet.cell(index + 4, column, value)
        sheet.merge_cells("A11:F11"); sheet["A11"] = report["sections"][kind]["corrective_action"]
        sheet.merge_cells("A12:F13"); sheet["A12"] = labels[kind][1]
        sheet.merge_cells("A14:F14"); sheet["A14"] = f"Source: {report['source_file']} | Limits: DSO TTC/ATC Master | Interval: {report['sample_interval_minutes']:g} minute(s)"
        for row in range(1, 15):
            for column in range(1, 7):
                cell = sheet.cell(row, column)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = Border(*(Side(style="thin", color=border_color) for _ in range(4)))
        for row in (1, 2, 3):
            sheet.cell(row, 1).font = Font(bold=True, size=13 if row == 1 else 12, color=navy)
            sheet.cell(row, 1).fill = PatternFill("solid", fgColor=pale)
        for cell in sheet[4]:
            cell.font = Font(bold=True, color=navy)
            cell.fill = PatternFill("solid", fgColor=blue)
        sheet["A11"].font = Font(bold=True, color="9C0006")
        sheet["A11"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet["A12"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        sheet["A14"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.column_dimensions["A"].width = 9; sheet.column_dimensions["B"].width = 27; sheet.column_dimensions["C"].width = 15
        sheet.column_dimensions["D"].width = 24; sheet.column_dimensions["E"].width = 27; sheet.column_dimensions["F"].width = 32
        sheet.row_dimensions[1].height = 54; sheet.row_dimensions[3].height = 28; sheet.row_dimensions[4].height = 70
        sheet.row_dimensions[11].height = 34; sheet.row_dimensions[12].height = 55; sheet.row_dimensions[14].height = 28
        sheet.freeze_panes = "A5"
        sheet.page_setup.orientation = "landscape"; sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins = PageMargins(left=.2, right=.2, top=.3, bottom=.3, header=.1, footer=.1)
        logo_path = Path(__file__).resolve().parents[2] / "frontend" / "public" / "logo.png"
        if logo_path.is_file():
            try:
                logo = ExcelImage(str(logo_path)); logo.width = 68; logo.height = 48; sheet.add_image(logo, "F1")
            except Exception:
                pass
    output = io.BytesIO(); workbook.save(output); output.seek(0)
    return output


def pdf_font():
    name = "Helvetica"
    for path in (Path(r"C:\Windows\Fonts\Nirmala.ttf"), Path(r"C:\Windows\Fonts\Arial.ttf")):
        if path.is_file():
            try:
                pdfmetrics.registerFont(TTFont("SRIUnicode", str(path)))
                return "SRIUnicode"
            except Exception:
                continue
    return name


def build_pdf(report: dict):
    output = io.BytesIO()
    font = pdf_font()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=8 * mm, bottomMargin=8 * mm)
    centered = ParagraphStyle("centered", fontName=font, fontSize=9, leading=12, alignment=TA_CENTER, textColor=colors.HexColor("#17365D"))
    heading = ParagraphStyle("heading", parent=centered, fontSize=13, leading=16, spaceAfter=4, fontName=font)
    note = ParagraphStyle("note", fontName=font, fontSize=8, leading=11, alignment=TA_LEFT, textColor=colors.HexColor("#334155"))
    story = []
    logo_path = Path(__file__).resolve().parents[2] / "frontend" / "public" / "logo.png"
    for section_index, kind in enumerate(("TTC", "ATC")):
        if section_index:
            from reportlab.platypus import PageBreak
            story.append(PageBreak())
        if logo_path.is_file():
            story.append(Image(str(logo_path), width=18 * mm, height=14 * mm))
        story.append(Paragraph("GRID CONTROLLER OF INDIA LIMITED<br/>(A Government of India Enterprise)<br/>EASTERN REGIONAL LOAD DESPATCH CENTRE, KOLKATA", heading))
        story.append(Paragraph(report["title"], heading))
        story.append(Paragraph(f"Reporting of {kind} Violation", heading))
        headers = ["Sl No.", "Intra-Regional Corridor", "Import / Export", f"Total Hours of {kind} Violation", f"Percentage Time of {kind} Violation", "Intimation for Corrective Action"]
        table_data = [[Paragraph(value, centered) for value in headers]]
        for index, row in enumerate(report["sections"][kind]["rows"], 1):
            table_data.append([str(index), row["corridor"], row["direction"], f"{row['violation_hours']:.2f}", f"{row['violation_percent']:.2f}%", row["intimation"]])
        table = Table(table_data, colWidths=[14 * mm, 50 * mm, 29 * mm, 45 * mm, 48 * mm, 72 * mm], rowHeights=[18 * mm] + [10 * mm] * 6)
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#17365D")), ("FONTNAME", (0, 0), (-1, 0), font),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), .55, colors.HexColor("#8EA9C1")),
        ]))
        story.extend([table, Spacer(1, 4 * mm), Paragraph(report["sections"][kind]["corrective_action"], note), Spacer(1, 3 * mm)])
        story.append(Paragraph(f"In case of violation of {kind}, regional entities in the importing/exporting control area shall restrict drawal or injection so that {kind} is maintained within limit and SLDCs shall take appropriate action.", note))
        story.append(Spacer(1, 2 * mm))
        story.append(Paragraph(f"Source: {report['source_file']} | TTC/ATC limits: portal database | Sampling interval: {report['sample_interval_minutes']:g} minute(s)", note))
    document.build(story); output.seek(0)
    return output


def build_excel_from_template(report: dict):
    """Populate only report cells so the supplied workbook's artwork/layout stays intact."""
    template = template_file().read_bytes()
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pkg_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    source_zip = zipfile.ZipFile(io.BytesIO(template), "r")
    shared_xml = source_zip.read("xl/sharedStrings.xml").decode("utf-8")
    shared_indices = {}
    shared_count = len(re.findall(r"<si(?:\s|>)", shared_xml))

    def add_shared_string(value):
        nonlocal shared_xml, shared_count
        value = str(value)
        if value in shared_indices:
            return shared_indices[value]
        index = shared_count
        shared_indices[value] = index
        shared_count += 1
        shared_xml = shared_xml.replace("</sst>", f"<si><t>{xml_escape(value)}</t></si></sst>")
        return index

    def replace_cell(xml, ref, value, text_value):
        pattern = re.compile(rf'(<c\b(?=[^>]*\br="{re.escape(ref)}")[^>]*)(?:/>|>.*?</c>)', re.DOTALL)
        match = pattern.search(xml)
        if not match:
            raise ValueError(f"The SRI template is missing required cell {ref}.")
        opening = re.sub(r'\s+t="[^"]*"', "", match.group(1))
        if text_value:
            replacement = f'{opening} t="s"><v>{add_shared_string(value)}</v></c>'
        else:
            numeric = f"{float(value):.4f}".rstrip("0").rstrip(".") or "0"
            replacement = f"{opening}><v>{numeric}</v></c>"
        return xml[:match.start()] + replacement + xml[match.end():]

    workbook_root = ET.fromstring(source_zip.read("xl/workbook.xml"))
    rels_root = ET.fromstring(source_zip.read("xl/_rels/workbook.xml.rels"))
    rel_targets = {
        relation.attrib["Id"]: relation.attrib["Target"]
        for relation in rels_root.findall(f"{{{pkg_rel_ns}}}Relationship")
    }
    sheet_paths = []
    for sheet in workbook_root.findall(f".//{{{main_ns}}}sheet")[:2]:
        relation_id = sheet.attrib[f"{{{rel_ns}}}id"]
        sheet_paths.append(posixpath.normpath(posixpath.join("xl", rel_targets[relation_id])))
    if len(sheet_paths) < 2:
        raise ValueError("The SRI template must contain separate TTC and ATC report sheets.")

    replacements = {}
    for sheet_path, kind in zip(sheet_paths, ("TTC", "ATC")):
        xml = source_zip.read(sheet_path).decode("utf-8")
        xml = replace_cell(xml, "A2", report["title"], True)
        for excel_row, row in enumerate(report["sections"][kind]["rows"], start=5):
            xml = replace_cell(xml, f"D{excel_row}", hours_text(row["violation_hours"]), True)
            xml = replace_cell(xml, f"E{excel_row}", row["violation_percent"], False)
            xml = replace_cell(xml, f"F{excel_row}", row["intimation"] or "NA", True)
        xml = replace_cell(xml, "A11", report["sections"][kind]["corrective_action"], True)
        replacements[sheet_path] = xml.encode("utf-8")

    if shared_indices:
        added = len(shared_indices)
        for attribute in ("count", "uniqueCount"):
            match = re.search(rf'\b{attribute}="(\d+)"', shared_xml)
            if match:
                updated = int(match.group(1)) + added
                shared_xml = shared_xml[:match.start(1)] + str(updated) + shared_xml[match.end(1):]
        replacements["xl/sharedStrings.xml"] = shared_xml.encode("utf-8")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target_zip:
        for filename in source_zip.namelist():
            target_zip.writestr(filename, replacements.get(filename, source_zip.read(filename)))
    source_zip.close()
    output.seek(0)
    return output


def build_pdf_from_template(report: dict):
    """Render the same bilingual A:G portrait layout used by the Excel template."""
    output = io.BytesIO()
    font = pdf_font()
    document = SimpleDocTemplate(
        output, pagesize=A4, leftMargin=8 * mm, rightMargin=8 * mm,
        topMargin=7 * mm, bottomMargin=7 * mm,
    )
    centered = ParagraphStyle("sri-centered", fontName=font, fontSize=7, leading=9, alignment=TA_CENTER)
    header_style = ParagraphStyle("sri-header", parent=centered, fontSize=10, leading=13)
    section_style = ParagraphStyle("sri-section", parent=centered, fontSize=8, leading=10)
    note_style = ParagraphStyle("sri-note", fontName=font, fontSize=7, leading=9, alignment=TA_LEFT)
    template_book = load_workbook(template_file(), data_only=False)
    logo_data = None
    if template_book.worksheets and template_book.worksheets[0]._images:
        logo_data = template_book.worksheets[0]._images[0]._data()
    story = []
    for section_index, kind in enumerate(("TTC", "ATC")):
        if section_index:
            story.append(PageBreak())
        template_sheet = template_book.worksheets[section_index]
        heading_text = xml_escape(str(template_sheet["A1"].value or "")).replace("\n", "<br/>")
        logo = Image(io.BytesIO(logo_data), width=42 * mm, height=15 * mm) if logo_data else ""
        header = Table([["", Paragraph(heading_text, header_style), logo]], colWidths=[42 * mm, 110 * mm, 42 * mm], rowHeights=[54 * mm])
        header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
        table_data = [
            [header, "", "", "", "", "", ""],
            [Paragraph(xml_escape(report["title"]), centered), "", "", "", "", "", ""],
            [Paragraph(xml_escape(str(template_sheet["A3"].value or "")), section_style), "", "", "", "", "", ""],
            [Paragraph(xml_escape(str(template_sheet.cell(4, column).value or "")).replace("\n", "<br/>"), centered) for column in range(1, 7)] + [""],
        ]
        for index, row in enumerate(report["sections"][kind]["rows"], 1):
            table_data.append([
                str(index), row["corridor"], kind, hours_text(row["violation_hours"]),
                f"{float(row['violation_percent']):g}",
                Paragraph(xml_escape(str(row["intimation"] or "NA")), centered), "",
            ])
        table_data.extend([
            [Paragraph(xml_escape(report["sections"][kind]["corrective_action"]), note_style), "", "", "", "", "", ""],
            [Paragraph(xml_escape(str(template_sheet["A12"].value or "")), note_style), "", "", "", "", "", ""],
            [Paragraph(xml_escape(str(template_sheet["A13"].value or "")).replace("\n", "<br/>"), note_style), "", "", "", "", "", ""],
        ])
        table = Table(
            table_data,
            colWidths=[17 * mm, 49 * mm, 31 * mm, 27 * mm, 25 * mm, 15 * mm, 30 * mm],
            rowHeights=[54 * mm, 7 * mm, 8 * mm, 48 * mm] + [6 * mm] * 6 + [8 * mm, 23 * mm, 22 * mm],
        )
        spans = [("SPAN", (0, row), (6, row)) for row in (0, 1, 2, 10, 11, 12)]
        spans += [("SPAN", (5, row), (6, row)) for row in range(3, 10)]
        table.setStyle(TableStyle([
            *spans,
            ("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 1), (-1, 9), .5, colors.black),
            ("ALIGN", (0, 10), (-1, 12), "LEFT"),
            ("LEFTPADDING", (0, 10), (-1, 12), 3), ("RIGHTPADDING", (0, 10), (-1, 12), 3),
        ]))
        story.append(table)
    template_book.close()
    document.build(story)
    output.seek(0)
    return output


# Keep the original helper names as the public export builders.
build_excel = build_excel_from_template
build_pdf = build_pdf_from_template


@router.get("/export.xlsx")
def export_sri_excel(report_date: str = Query(...)):
    report = report_for_date(report_date)
    return StreamingResponse(build_excel_from_template(report), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="SRI {datetime.strptime(report_date, "%Y-%m-%d"):%d%m%Y}.xlsx"'})


@router.get("/export.pdf")
def export_sri_pdf(report_date: str = Query(...)):
    report = report_for_date(report_date)
    return StreamingResponse(build_pdf_from_template(report), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="SRI {datetime.strptime(report_date, "%Y-%m-%d"):%d%m%Y}.pdf"'})
