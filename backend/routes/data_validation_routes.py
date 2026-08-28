import io
import html
import base64
import math
import os
import re
import time
import uuid
from datetime import datetime
from pathlib import Path

import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from typing import Optional
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from crew_legacy.admin_logic.auth_utils import (
    get_authenticated_user,
    require_page_view,
    require_page_write,
)
from crew_legacy.admin_logic.notification_service import (
    replacement_mail_settings,
    send_email,
    workflow_mail_templates,
)


router = APIRouter(prefix="/api/data-validation", tags=["Data Validation"])
PAGE_KEY = "data_validation"
MAX_UPLOAD_BYTES = 30 * 1024 * 1024
REPORT_TTL_SECONDS = 2 * 60 * 60
REPORT_CACHE = {}
PSP_SOURCE_ROOT = Path(os.getenv(
    "PSP_VOLTAGE_SOURCE_ROOT",
    r"\\10.3.95.200\HTTP-Access\ScadaData\Web_based_PSP",
))

LEVEL_RULES = {
    400: {"nominal_min": 340.0, "nominal_max": 460.0, "band_min": 380.0, "band_max": 420.0},
    765: {"nominal_min": 650.25, "nominal_max": 879.75, "band_min": 728.0, "band_max": 800.0},
}


class SourceReportRequest(BaseModel):
    reportDate: str = Field(min_length=10, max_length=10)


class MailReportRequest(BaseModel):
    reportId: str = Field(min_length=8, max_length=80)
    subject: Optional[str] = Field(default=None, max_length=300)
    html: Optional[str] = Field(default=None, max_length=500000)


class _MailValues(dict):
    def __missing__(self, key):
        return "{" + key + "}"


def detect_voltage_level(name: str):
    """Detect 400/765 kV even when separators and the kV suffix vary."""
    text = str(name or "").upper().replace("KV", " KV ")
    matches = re.findall(r"(?<!\d)(765|400)(?!\d)", text)
    return int(matches[-1]) if matches else None


def display_number(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    return round(float(value), 4)


def display_datetime(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (int, float)) and 20000 <= float(value) <= 80000:
        try:
            return from_excel(float(value)).isoformat(sep=" ", timespec="seconds")
        except (TypeError, ValueError, OverflowError):
            pass
    parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)
    if not pd.isna(parsed):
        return parsed.to_pydatetime().isoformat(sep=" ", timespec="seconds")
    return str(value).strip()


def contiguous_blocks(indices):
    ordered = sorted(set(int(index) for index in indices))
    if not ordered:
        return []
    blocks = []
    start = previous = ordered[0]
    for index in ordered[1:]:
        if index != previous + 1:
            blocks.append((start, previous))
            start = index
        previous = index
    blocks.append((start, previous))
    return blocks


def flat_run_blocks(values, minimum_samples=121):
    blocks = []
    start = None
    previous = None
    for index, raw_value in enumerate(values):
        value = None if pd.isna(raw_value) else round(float(raw_value), 6)
        if value is None:
            if start is not None and index - start >= minimum_samples:
                blocks.append((start, index - 1))
            start = previous = None
            continue
        if start is None or value != previous:
            if start is not None and index - start >= minimum_samples:
                blocks.append((start, index - 1))
            start = index
        previous = value
    if start is not None and len(values) - start >= minimum_samples:
        blocks.append((start, len(values) - 1))
    return blocks


def read_voltage_dump(contents: bytes, filename: str):
    suffix = Path(filename or "").suffix.lower()
    if suffix not in {".xlsx", ".xlsm", ".xls", ".xlsb", ".csv"}:
        raise ValueError("Upload an Excel or CSV voltage dump (.xlsx, .xlsm, .xls, .xlsb or .csv).")
    source = io.BytesIO(contents)
    if suffix == ".csv":
        frame = pd.read_csv(source, header=None)
        sheet_name = "CSV"
    elif suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(source, read_only=True, data_only=True)
        candidates = [sheet for sheet in workbook.worksheets if not sheet.title.startswith("_")]
        if not candidates:
            candidates = list(workbook.worksheets)

        selected = None
        selected_last_column = 0
        selected_score = -1
        for sheet in candidates:
            first_rows = list(sheet.iter_rows(min_row=1, max_row=min(3, sheet.max_row), values_only=True))
            if len(first_rows) < 3:
                continue
            last_column = 0
            for row in first_rows[:2]:
                for index, value in enumerate(row, start=1):
                    if value is not None and str(value).strip():
                        last_column = max(last_column, index)
            numeric_samples = sum(
                isinstance(value, (int, float)) and not isinstance(value, bool)
                for value in first_rows[2][1:last_column]
            )
            score = last_column + numeric_samples * 10
            if last_column >= 2 and numeric_samples and score > selected_score:
                selected = sheet
                selected_last_column = last_column
                selected_score = score
        if selected is None:
            raise ValueError("No worksheet matching the PSP voltage layout was found.")

        values = list(selected.iter_rows(
            min_row=1,
            max_row=selected.max_row,
            min_col=1,
            max_col=selected_last_column,
            values_only=True,
        ))
        frame = pd.DataFrame(values)
        sheet_name = selected.title
        workbook.close()
    else:
        engine = "calamine" if suffix in {".xls", ".xlsb"} else "openpyxl"
        workbook = pd.ExcelFile(source, engine=engine)
        candidate_names = [name for name in workbook.sheet_names if not str(name).startswith("_")] or workbook.sheet_names
        selected_name = None
        selected_score = -1
        for name in candidate_names:
            preview = pd.read_excel(workbook, sheet_name=name, header=None, nrows=3)
            preview = preview.dropna(axis=1, how="all")
            numeric_samples = pd.to_numeric(preview.iloc[2, 1:], errors="coerce").notna().sum() if preview.shape[0] >= 3 else 0
            score = preview.shape[1] + int(numeric_samples) * 10
            if preview.shape[1] >= 2 and numeric_samples and score > selected_score:
                selected_name = name
                selected_score = score
        if selected_name is None:
            raise ValueError("No worksheet matching the PSP voltage layout was found.")
        frame = pd.read_excel(workbook, sheet_name=selected_name, header=None)
        sheet_name = selected_name
    frame = frame.dropna(axis=1, how="all").dropna(axis=0, how="all")
    if frame.shape[0] < 3 or frame.shape[1] < 2:
        raise ValueError("Expected station names in row 1, SCADA keys in row 2, and voltage samples from row 3 onward.")
    return frame.reset_index(drop=True), sheet_name


def validate_voltage_dump(contents: bytes, filename: str):
    frame, sheet_name = read_voltage_dump(contents, filename)
    sample_frame = frame.iloc[2:].reset_index(drop=True)
    date_times = [display_datetime(value) for value in sample_frame.iloc[:, 0].tolist()]
    total_samples = len(sample_frame)
    if not total_samples:
        raise ValueError("No voltage samples were found below the two header rows.")

    stations = []
    issue_blocks = []
    values_by_station = {}
    flags_by_station = {}
    issue_serial = 0

    def add_issue(station, code, title, remark, start_index=None, end_index=None, severity="error"):
        nonlocal issue_serial
        issue_serial += 1
        sample_count = 0 if start_index is None else end_index - start_index + 1
        issue = {
            "id": f"issue-{issue_serial}",
            "stationId": station["id"],
            "stationName": station["name"],
            "code": code,
            "title": title,
            "remark": remark,
            "severity": severity,
            "startIndex": start_index,
            "endIndex": end_index,
            "startTime": date_times[start_index] if start_index is not None else None,
            "endTime": date_times[end_index] if end_index is not None else None,
            "sampleCount": sample_count,
        }
        issue_blocks.append(issue)
        station["issues"].append(issue)
        if start_index is not None:
            for row_index in range(start_index, end_index + 1):
                flags_by_station[station["id"]][row_index].append({"code": code, "remark": remark})

    for column_index in range(1, frame.shape[1]):
        name_value = frame.iat[0, column_index]
        key_value = frame.iat[1, column_index]
        name = str(name_value).strip() if not pd.isna(name_value) else f"Column {column_index + 1}"
        station_id = f"station-{column_index}"
        level = detect_voltage_level(name)
        numeric = pd.to_numeric(sample_frame.iloc[:, column_index], errors="coerce")
        values = numeric.tolist()
        valid = numeric.dropna()
        rules = LEVEL_RULES.get(level)
        flags_by_station[station_id] = [[] for _ in range(total_samples)]
        values_by_station[station_id] = [display_number(value) for value in values]
        station = {
            "id": station_id,
            "name": name,
            "scadaKey": "" if pd.isna(key_value) else str(key_value).strip(),
            "voltageLevel": level,
            "totalSamples": total_samples,
            "validSamples": int(valid.count()),
            "missingSamples": int(numeric.isna().sum()),
            "minimum": display_number(valid.min()) if len(valid) else None,
            "maximum": display_number(valid.max()) if len(valid) else None,
            "average": display_number(valid.mean()) if len(valid) else None,
            "vdiPercent": None,
            "outsideNominalSamples": 0,
            "outsideBandSamples": 0,
            "flatRunSamples": 0,
            "flaggedSamples": 0,
            "issues": [],
        }
        if rules:
            station.update(rules)
        stations.append(station)

        for start, end in contiguous_blocks(numeric[numeric.isna()].index):
            add_issue(station, "MISSING_VALUE", "Missing/non-numeric voltage", f"{end - start + 1} consecutive sample(s) are blank or non-numeric.", start, end)

        flat_blocks = flat_run_blocks(values)
        station["flatRunSamples"] = sum(end - start + 1 for start, end in flat_blocks)
        for start, end in flat_blocks:
            add_issue(
                station, "FLAT_RUN", "Voltage unchanged for more than 120 samples",
                f"The same voltage ({values_by_station[station_id][start]:g} kV) continues for {end - start + 1} consecutive one-minute samples.", start, end,
            )

        if len(valid) and math.isclose(float(valid.min()), float(valid.max()), rel_tol=0, abs_tol=1e-9) and math.isclose(float(valid.mean()), float(valid.max()), rel_tol=0, abs_tol=1e-9):
            add_issue(
                station, "IDENTICAL_DAILY_STATS", "Maximum, minimum and average are identical",
                f"Daily maximum, minimum and average are all {float(valid.mean()):g} kV; verify the SCADA source.",
                0, total_samples - 1, severity="error",
            )

        if not rules:
            add_issue(station, "UNKNOWN_LEVEL", "Voltage level not detected", "Could not identify 400 kV or 765 kV from the substation name.", severity="warning")
            station["status"] = "Issue" if any(issue["severity"] == "error" for issue in station["issues"]) else "Review"
            station["issueCount"] = len(station["issues"])
            station["flaggedSamples"] = sum(bool(flags) for flags in flags_by_station[station_id])
            continue

        outside_nominal = numeric.notna() & ((numeric < rules["nominal_min"]) | (numeric > rules["nominal_max"]))
        station["outsideNominalSamples"] = int(outside_nominal.sum())
        for start, end in contiguous_blocks(numeric[outside_nominal].index):
            add_issue(
                station, "OUTSIDE_NOMINAL_RANGE", "Outside ±15% nominal range",
                f"Voltage must remain between {rules['nominal_min']:g} and {rules['nominal_max']:g} kV for a {level} kV station.", start, end,
            )

        outside_band = numeric.notna() & ((numeric < rules["band_min"]) | (numeric > rules["band_max"]))
        station["outsideBandSamples"] = int(outside_band.sum())
        station["vdiPercent"] = round((station["outsideBandSamples"] / station["validSamples"] * 100), 3) if station["validSamples"] else None
        if station["vdiPercent"] is not None and station["vdiPercent"] > 10:
            for start, end in contiguous_blocks(numeric[outside_band].index):
                add_issue(
                    station, "VDI_ABOVE_10", "Voltage Deviation Index above 10%",
                    f"VDI is {station['vdiPercent']:.2f}% using the {rules['band_min']:g}–{rules['band_max']:g} kV band; outside-band samples are marked.", start, end,
                )

        station["status"] = "Issue" if station["issues"] else "Healthy"
        station["issueCount"] = len(station["issues"])
        station["flaggedSamples"] = sum(bool(flags) for flags in flags_by_station[station_id])

    rows = []
    flagged_cells = 0
    for row_index in range(total_samples):
        flags = {station["id"]: flags_by_station[station["id"]][row_index] for station in stations if flags_by_station[station["id"]][row_index]}
        flagged_cells += len(flags)
        rows.append({
            "index": row_index,
            "dateTime": date_times[row_index],
            "values": {station["id"]: values_by_station[station["id"]][row_index] for station in stations},
            "flags": flags,
        })

    return {
        "fileName": filename,
        "sheetName": sheet_name,
        "sampleCount": total_samples,
        "stationCount": len(stations),
        "stationsWithIssues": sum(station["status"] != "Healthy" for station in stations),
        "healthyStations": sum(station["status"] == "Healthy" for station in stations),
        "issueBlockCount": len(issue_blocks),
        "flaggedCellCount": flagged_cells,
        "stations": stations,
        "issues": issue_blocks,
        "rows": rows,
    }


def purge_report_cache():
    cutoff = time.time() - REPORT_TTL_SECONDS
    for report_id in [key for key, item in REPORT_CACHE.items() if item["created"] < cutoff]:
        REPORT_CACHE.pop(report_id, None)
    while len(REPORT_CACHE) > 10:
        oldest = min(REPORT_CACHE, key=lambda key: REPORT_CACHE[key]["created"])
        REPORT_CACHE.pop(oldest, None)


def make_marked_workbook(report):
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Validation Summary"
    navy_fill = PatternFill("solid", fgColor="09245C")
    blue_fill = PatternFill("solid", fgColor="DCEBFA")
    red_fill = PatternFill("solid", fgColor="FECACA")
    amber_fill = PatternFill("solid", fgColor="FEF3C7")
    green_fill = PatternFill("solid", fgColor="DCFCE7")
    white_font = Font(color="FFFFFF", bold=True)

    summary_sheet.append(["PSP Voltage Dump Validation Report"])
    summary_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    summary_sheet["A1"].fill = navy_fill
    summary_sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    summary_sheet.append(["Source file", report["fileName"], "Sheet", report["sheetName"], "Samples", report["sampleCount"], "Stations", report["stationCount"], "Issue stations", report["stationsWithIssues"]])
    headers = ["Substation", "SCADA key", "Level", "Status", "Valid", "Missing", "Minimum", "Maximum", "Average", "VDI %", "Flagged samples", "Issue blocks"]
    summary_sheet.append(headers)
    for cell in summary_sheet[3]:
        cell.fill = blue_fill
        cell.font = Font(bold=True, color="003B78")
    for station in report["stations"]:
        summary_sheet.append([
            station["name"], station["scadaKey"], f"{station['voltageLevel']} kV" if station["voltageLevel"] else "Undetected",
            station["status"], station["validSamples"], station["missingSamples"], station["minimum"], station["maximum"],
            station["average"], station["vdiPercent"], station["flaggedSamples"], station["issueCount"],
        ])
        summary_sheet.cell(summary_sheet.max_row, 4).fill = green_fill if station["status"] == "Healthy" else red_fill if station["status"] == "Issue" else amber_fill
    summary_sheet.freeze_panes = "A4"
    summary_sheet.auto_filter.ref = summary_sheet.dimensions
    summary_sheet.column_dimensions["A"].width = 34
    for column in range(2, 13):
        summary_sheet.column_dimensions[get_column_letter(column)].width = 16

    issues_sheet = workbook.create_sheet("Issue Blocks")
    issue_headers = ["Substation", "Level", "Rule", "Start", "End", "Samples", "Remark"]
    issues_sheet.append(issue_headers)
    for cell in issues_sheet[1]:
        cell.fill = navy_fill
        cell.font = white_font
    station_map = {station["id"]: station for station in report["stations"]}
    for issue in report["issues"]:
        station = station_map[issue["stationId"]]
        issues_sheet.append([station["name"], station["voltageLevel"], issue["title"], issue["startTime"], issue["endTime"], issue["sampleCount"], issue["remark"]])
        for cell in issues_sheet[issues_sheet.max_row]:
            cell.fill = red_fill if issue["severity"] == "error" else amber_fill
    issues_sheet.freeze_panes = "A2"
    issues_sheet.auto_filter.ref = issues_sheet.dimensions
    issues_sheet.column_dimensions["A"].width = 34
    issues_sheet.column_dimensions["C"].width = 40
    issues_sheet.column_dimensions["G"].width = 80

    data_sheet = workbook.create_sheet("Marked Data")
    data_sheet.append(["DateTime", *[station["name"] for station in report["stations"]]])
    data_sheet.append(["SCADA Key (ignored)", *[station["scadaKey"] for station in report["stations"]]])
    for row_number in (1, 2):
        for cell in data_sheet[row_number]:
            cell.fill = navy_fill if row_number == 1 else blue_fill
            cell.font = white_font if row_number == 1 else Font(bold=True, color="003B78")
            cell.alignment = Alignment(horizontal="center")
    for row in report["rows"]:
        data_sheet.append([row["dateTime"], *[row["values"].get(station["id"]) for station in report["stations"]]])
        excel_row = data_sheet.max_row
        for station_column, station in enumerate(report["stations"], start=2):
            cell_flags = row["flags"].get(station["id"]) or []
            if not cell_flags:
                continue
            cell = data_sheet.cell(excel_row, station_column)
            cell.fill = red_fill
            remarks = list(dict.fromkeys(flag["remark"] for flag in cell_flags))
            cell.comment = Comment("\n".join(remarks), "COMPASS Data Validation")
    data_sheet.freeze_panes = "B3"
    data_sheet.auto_filter.ref = f"A1:{get_column_letter(data_sheet.max_column)}{data_sheet.max_row}"
    data_sheet.column_dimensions["A"].width = 22
    for column in range(2, data_sheet.max_column + 1):
        data_sheet.column_dimensions[get_column_letter(column)].width = 18
    return workbook


def cache_report(report):
    purge_report_cache()
    report_id = uuid.uuid4().hex
    report["reportId"] = report_id
    REPORT_CACHE[report_id] = {"created": time.time(), "report": report}
    return report


def get_cached_report(report_id):
    purge_report_cache()
    cached = REPORT_CACHE.get(str(report_id or ""))
    if not cached:
        raise HTTPException(404, "This validation report has expired. Generate the report again.")
    return cached["report"]


def source_file_for_date(report_date):
    try:
        parsed = datetime.strptime(str(report_date), "%Y-%m-%d")
    except ValueError as exc:
        raise HTTPException(400, "Report date must be in YYYY-MM-DD format.") from exc
    folder = PSP_SOURCE_ROOT / parsed.strftime("%Y") / parsed.strftime("%m")
    stem = f"ER_PSP_Voltage_{parsed.strftime('%d%m%Y')}"
    for extension in (".xlsx", ".xlsm", ".xlsb", ".xls", ".csv"):
        candidate = folder / f"{stem}{extension}"
        if candidate.is_file():
            return candidate, parsed
    raise HTTPException(404, f"PSP voltage dump is not available for {parsed.strftime('%d-%m-%Y')} in {folder}.")


def discrepancy_mail(report):
    template = workflow_mail_templates().get("psp_voltage_discrepancy") or {}
    mail_settings = replacement_mail_settings()
    report_date_iso = str(report.get("reportDate") or "")
    if not report_date_iso and report.get("rows"):
        report_date_iso = str(report["rows"][0].get("dateTime") or "")[:10]
    try:
        report_date = datetime.strptime(report_date_iso, "%Y-%m-%d").strftime("%d-%m-%Y")
    except ValueError:
        report_date = report_date_iso or "N/A"
    values = _MailValues({
        "report_date": report_date,
        "file_name": str(report.get("fileName") or ""),
        "station_count": str(report.get("stationCount") or 0),
        "issue_station_count": str(report.get("stationsWithIssues") or 0),
        "issue_block_count": str(report.get("issueBlockCount") or 0),
    })
    try:
        subject = str(template.get("subjectTemplate") or "PSP Voltage Discrepancy Report - {report_date}").format_map(values)
        introduction = str(template.get("bodyTemplate") or "PSP voltage discrepancies for {report_date}.").format_map(values)
    except ValueError as exc:
        raise HTTPException(400, "The PSP voltage discrepancy mail template is invalid.") from exc

    issue_rows = []
    def mail_cell(value, *, align="left", raw=False, alert=False):
        content = str(value) if raw else html.escape(str(value))
        colors = "background:#fee2e2;color:#991b1b;font-weight:bold;" if alert else ""
        return f'<td style="border:1px solid #cbd5e1;padding:7px;vertical-align:top;text-align:{align};{colors}">{content}</td>'

    for station in [item for item in report.get("stations") or [] if item.get("status") != "Healthy"]:
        observations = []
        for issue in station.get("issues") or []:
            period = "Station-level"
            if issue.get("startTime"):
                period = str(issue["startTime"])
                if issue.get("endTime") and issue["endTime"] != issue["startTime"]:
                    period += f" to {issue['endTime']}"
            sample_text = f"; {int(issue.get('sampleCount') or 0)} sample(s)" if issue.get("sampleCount") else ""
            observations.append(
                f"<li><strong>{html.escape(str(issue.get('title') or 'Issue'))}</strong> — "
                f"{html.escape(str(issue.get('remark') or ''))} "
                f"<span style=\"color:#64748b\">({html.escape(period)}{sample_text})</span></li>"
            )
        nominal_min = station.get("nominal_min")
        nominal_max = station.get("nominal_max")
        def statistic_outside(value):
            return value is not None and nominal_min is not None and (value < nominal_min or value > nominal_max)
        issue_rows.append(
            "<tr>"
            + mail_cell(station.get("name") or "")
            + mail_cell(station.get("scadaKey") or "")
            + mail_cell(f"{station.get('voltageLevel') or 'Undetected'} kV")
            + mail_cell(station.get("minimum") if station.get("minimum") is not None else "—", align="right", alert=statistic_outside(station.get("minimum")))
            + mail_cell(station.get("maximum") if station.get("maximum") is not None else "—", align="right", alert=statistic_outside(station.get("maximum")))
            + mail_cell(station.get("average") if station.get("average") is not None else "—", align="right", alert=statistic_outside(station.get("average")))
            + mail_cell(f"{station.get('vdiPercent')}%" if station.get("vdiPercent") is not None else "—", align="right", alert=(station.get("vdiPercent") or 0) > 10)
            + mail_cell(station.get("flatRunSamples") or 0, align="right", alert=bool(station.get("flatRunSamples")))
            + mail_cell(station.get("flaggedSamples") or 0, align="right")
            + mail_cell(f'<ol style="margin:0;padding-left:18px">{"".join(observations)}</ol>', raw=True)
            + "</tr>"
        )
    table = (
        "<table style=\"border-collapse:collapse;width:100%;font-family:Arial,sans-serif;font-size:12px\">"
        "<thead><tr style=\"background:#0b5f55;color:#fff\">"
        + "".join(f"<th style=\"border:1px solid #cbd5e1;padding:8px;text-align:left\">{title}</th>" for title in (
            "Substation", "SCADA Key", "Level", "Minimum", "Maximum", "Average", "VDI", "No. of Flat Samples", "Marked", "Observations"
        ))
        + "</tr></thead><tbody>"
        + ("".join(issue_rows) if issue_rows else "<tr><td colspan=\"10\" style=\"padding:12px;text-align:center\">No discrepancy found.</td></tr>")
        + "</tbody></table>"
    )
    body = (
        "<div style=\"font-family:Arial,sans-serif;color:#1e293b;font-size:13px;line-height:1.5\">"
        f"{introduction}<br><br>{table}"
        "<p style=\"color:#64748b;font-size:11px\">Generated by COMPASS · Data Validation · PSP Voltage Dump</p>"
        "</div>"
    )
    recipients = [value.strip() for value in re.split(r"[,;\n]+", str(template.get("recipients") or "")) if value.strip()]
    recipient_keys = {value.lower() for value in recipients}
    cc_recipients = []
    cc_keys = set()
    for raw_value in re.split(r"[,;\n]+", str(template.get("ccRecipients") or "")):
        value = raw_value.strip()
        key = value.lower()
        if value and key not in recipient_keys and key not in cc_keys:
            cc_keys.add(key)
            cc_recipients.append(value)
    return {
        "subject": subject,
        "html": body,
        "recipientCount": len(recipients),
        "recipients": recipients,
        "ccRecipientCount": len(cc_recipients),
        "ccRecipients": cc_recipients,
        "mailEnabled": bool(mail_settings.get("enabled")),
        "templateEnabled": bool(template.get("enabled")),
        "sender": mail_settings.get("sender") or "",
        "reportDate": report_date,
        "issueStationCount": report.get("stationsWithIssues") or 0,
        "issueBlockCount": report.get("issueBlockCount") or 0,
    }


@router.post("/psp-voltage/validate")
async def validate_psp_voltage(file: UploadFile = File(...), user=Depends(get_authenticated_user)):
    require_page_write(user, PAGE_KEY)
    contents = await file.read()
    if not contents:
        raise HTTPException(400, "The uploaded voltage dump is empty.")
    if len(contents) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "The voltage dump exceeds the 30 MB upload limit.")
    try:
        report = validate_voltage_dump(contents, file.filename or "ER_PSP_Voltage.xlsx")
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(422, f"Unable to parse the voltage dump: {exc}") from exc
    return cache_report(report)


@router.post("/psp-voltage/generate")
def generate_psp_voltage_report(payload: SourceReportRequest, user=Depends(get_authenticated_user)):
    require_page_write(user, PAGE_KEY)
    source_file, parsed_date = source_file_for_date(payload.reportDate)
    try:
        report = validate_voltage_dump(source_file.read_bytes(), source_file.name)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except OSError as exc:
        raise HTTPException(503, f"The SCADA source file could not be read: {exc}") from exc
    except Exception as exc:
        raise HTTPException(422, f"Unable to parse the voltage dump: {exc}") from exc
    report["reportDate"] = parsed_date.strftime("%Y-%m-%d")
    report["source"] = "SCADA network folder"
    return cache_report(report)


@router.post("/psp-voltage/mail-preview")
def preview_psp_voltage_mail(payload: MailReportRequest, user=Depends(get_authenticated_user)):
    require_page_view(user, PAGE_KEY)
    preview = discrepancy_mail(get_cached_report(payload.reportId))
    return {key: value for key, value in preview.items() if key not in {"recipients", "ccRecipients"}}


@router.post("/psp-voltage/send-mail")
def send_psp_voltage_mail(payload: MailReportRequest, user=Depends(get_authenticated_user)):
    require_page_write(user, PAGE_KEY)
    report = get_cached_report(payload.reportId)
    mail = discrepancy_mail(report)
    if not mail["mailEnabled"]:
        raise HTTPException(409, "Enable workflow email delivery in Mail Settings before sending.")
    if not mail["templateEnabled"]:
        raise HTTPException(409, "Enable the PSP Voltage discrepancy template in Mail Settings before sending.")
    if not mail["recipients"]:
        raise HTTPException(409, "Configure at least one PSP Voltage discrepancy recipient in Mail Settings.")
    subject = str(payload.subject or mail["subject"]).strip()
    body_html = str(payload.html or mail["html"]).strip()
    if not subject or not body_html:
        raise HTTPException(400, "Mail subject and content are required.")
    # Preserve ordinary email formatting but remove executable content from an edited preview.
    body_html = re.sub(r"<script\b[^>]*>.*?</script\s*>", "", body_html, flags=re.IGNORECASE | re.DOTALL)
    body_html = re.sub(r"\son[a-z]+\s*=\s*(['\"]).*?\1", "", body_html, flags=re.IGNORECASE | re.DOTALL)
    body_html = re.sub(r"javascript\s*:", "", body_html, flags=re.IGNORECASE)
    workbook = make_marked_workbook(report)
    workbook_output = io.BytesIO()
    workbook.save(workbook_output)
    attachment_name = f"{Path(report.get('fileName') or 'PSP_Voltage').stem}_validation.xlsx"
    result = send_email(
        mail["recipients"], subject, body_html, html=True,
        sender=mail["sender"], enabled=True, cc_list=mail["ccRecipients"],
        attachments=[{
            "name": attachment_name,
            "contentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "contentBytes": base64.b64encode(workbook_output.getvalue()).decode("ascii"),
        }],
    )
    if result.get("status") != "sent":
        raise HTTPException(503, result.get("error") or "Discrepancy mail could not be sent.")
    return {
        "status": "sent",
        "recipientCount": result.get("recipientCount", 0),
        "ccRecipientCount": result.get("ccRecipientCount", 0),
        "attachmentName": attachment_name,
        "subject": subject,
        "message": (
            f"Discrepancy report and validation workbook sent to {result.get('recipientCount', 0)} receiver(s)"
            f" with {result.get('ccRecipientCount', 0)} CC recipient(s)."
        ),
    }


@router.get("/psp-voltage/{report_id}/excel")
def download_psp_voltage_excel(report_id: str, user=Depends(get_authenticated_user)):
    require_page_view(user, PAGE_KEY)
    report = get_cached_report(report_id)
    workbook = make_marked_workbook(report)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    source_stem = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(report["fileName"]).stem).strip("_") or "PSP_Voltage"
    headers = {"Content-Disposition": f'attachment; filename="{source_stem}_validation.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)
