"""Event report presentation. Images must come from the existing HTML charts."""
import base64
import io
from html import escape

LEVELS = ('49.50', '49.70', '49.90')
CATEGORIES = ('Alert', 'Emergency', 'Extreme Emergency', 'Non-Compliance', 'Warning')
SECTIONS = ('summary', 'states', 'generators', 'actions', 'defence', 'annexure_states', 'annexure_generators')
NAVY = '17365D'


def text(value):
    return 'Data not available' if value is None else str(round(value, 3) if isinstance(value, float) else value)


def report_blocks(event, options, charts):
    """One presentation model for HTML, editable Word and PDF; no diagnostics."""
    analysis = event.get('threshold_analysis') or {}
    summary = analysis.get('summary') or {}
    performance = analysis.get('overall_performance') or {}
    included = options.get('operation_sections', SECTIONS)
    blocks = []
    def heading(title): blocks.append(('heading', title))
    def paragraph(value): blocks.append(('paragraph', value))
    def table(headers, rows, grouped=False): blocks.append(('table', (headers, rows, grouped)))
    period = event['start_time'].replace('T', ' ') + ' to ' + event['end_time'].replace('T', ' ') + ' IST'
    minimum = text(summary.get('minimum_frequency')) + ' Hz at ' + text(summary.get('minimum_timestamp'))
    from datetime import datetime
    compact_period = '\n'.join(datetime.fromisoformat(event[key]).strftime('%d-%b-%y %H:%M') for key in ('start_time','end_time'))
    compact_minimum = text(summary.get('minimum_frequency')) + ' Hz\n' + (datetime.fromisoformat(summary['minimum_timestamp']).strftime('%d-%b-%y %H:%M:%S') if summary.get('minimum_timestamp') else 'Data not available')
    if 'summary' in included:
        heading('1 Executive Summary & General Notes')
        duration = summary.get('thresholds', {}).get('49.90', {}).get('frequency_minutes')
        paragraph(options.get('executive_summary') or f"On {event['start_time'][:10]}, a low-frequency event occurred from {event['start_time'][11:]} to {event['end_time'][11:]} IST. Frequency was below 49.9 Hz for {text(duration)} minutes. Minimum frequency was {minimum}.")
        # Metadata remains authoritative even when the narrative is edited.
        paragraph(f'Reporting period: {period}. Duration below 49.9 Hz: {text(duration)} minutes. Minimum frequency: {minimum}.')
        for chart in charts:
            if chart.get('kind') == 'System Frequency': blocks.append(('image', chart))
        counts = {category: set() for category in CATEGORIES}
        for row in event.get('chronology', []):
            if row.get('record_kind') == 'physical': continue
            categories = row.get('message_categories') or [row.get('message_type')]
            if isinstance(categories, str): categories = [categories]
            for category in categories:
                key = next((c for c in CATEGORIES if c.lower().replace('-', '').replace(' ', '') == str(category).lower().replace('-', '').replace(' ', '')), None)
                if key: counts[key].add((row.get('timestamp'), row.get('message_no'), row.get('message_details')))
        table(list(CATEGORIES), [[len(counts[c]) if event.get('messages_complete') else None for c in CATEGORIES]])
        majors = sorted(performance.get('State', []), key=lambda row: row.get('thresholds', {}).get('49.90', {}).get('maximum_od_ui_mw') or 0, reverse=True)
        paragraph('Major state overdrawal: ' + ('; '.join(f"{r['entity']}: {text(r['thresholds']['49.90'].get('maximum_od_ui_mw'))} MW" for r in majors[:3]) if majors else 'Data not available'))
    for section, title, groups in [('states', '2 State Performance Details', ['State']), ('generators', '3 Central Sector Generator Performance', ['ISGS', 'IPP'])]:
        if section not in included: continue
        heading(title)
        paragraph('Reporting period: ' + period + '. Minimum frequency: ' + minimum + '.')
        headers = ['Reporting Period', 'State Name' if section == 'states' else 'Generator Name (Agency)', 'Min. Freq. (Hz) & Time']
        headers += ['OD/UI Duration (Min & %)', 'Avg. OD/UI (MW)', 'Max. OD/UI (MW)', 'Violation Messages'] * 3
        rows = []
        for group in groups:
            for row in performance.get(group, []):
                thresholds = row.get('thresholds', {})
                if section == 'generators' and not any((v.get('adverse_minutes') or 0) > 0 for v in thresholds.values()): continue
                values = [compact_period, row['entity'] + (f' ({group})' if section == 'generators' else ''), compact_minimum]
                for level in LEVELS:
                    stats = thresholds.get(level, {})
                    duration = 'Data not available' if (stats.get('unknown_deviation_minutes') or 0) > 0 else f"{text(stats.get('adverse_minutes'))} min / {text(stats.get('adverse_pct'))}%"
                    values.extend([duration, stats.get('average_od_ui_mw'), stats.get('maximum_od_ui_mw'), stats.get('message_count')])
                rows.append(values)
        table(headers, rows, True)
    if 'actions' in included:
        heading('4 Action & Chronology of Events')
        paragraph(options.get('action_summary') or 'Data not available')
        headers = ['Time (IST)', 'Frequency (Hz)', 'State / Entity', 'OD/UI (MW)', 'Message Type / Action', 'Message No.', 'Message / Details']
        messages = {}; rows = []
        for row in sorted(event.get('report_chronology', event.get('chronology', [])), key=lambda r: r['timestamp']):
            key = (row['timestamp'], row.get('message_no'), row.get('message_details'), row.get('record_kind'))
            messages.setdefault(key, {})[row.get('state') or 'Data not available'] = row
        for recipients in messages.values():
            first = next(iter(recipients.values()))
            names = ', '.join(recipients)
            def measurement(key):
                values = [row.get(key) for row in recipients.values()]
                if len(recipients) == 1 or all(value == values[0] for value in values): return values[0]
                return '; '.join(name + ': ' + text(row.get(key)) for name, row in recipients.items())
            rows.append([first['timestamp'], measurement('frequency_hz'), names, measurement('deviation_mw'), first.get('message_type'), first.get('message_no'), first.get('message_details')])
        table(headers, rows)
    if 'defence' in included:
        heading('5 Action of Defence Mechanism (ADMS & UFR)')
        defence = []
        for row in event.get('chronology', []):
            categories = row.get('message_categories') or [row.get('message_type')]
            if isinstance(categories, str): categories = [categories]
            mechanisms = [category for category in categories if str(category).upper().replace(' ', '') in {'ADMS','UFR','ADMS/UFR'}]
            if mechanisms: defence.append([row.get('timestamp'), ', '.join(mechanisms), row.get('state'), row.get('message_details')])
        if defence: table(['Time (IST)', 'Mechanism', 'Entity', 'CRMS Record'], defence)
        else: paragraph('Data not available')
    for section, title, state in [('annexure_states', 'Annexure 1 State-wise Low Frequency Analysis', True), ('annexure_generators', 'Annexure 2 Generator-wise Low Frequency Analysis', False)]:
        if section not in included: continue
        heading(title)
        selected = [c for c in charts if c.get('kind') != 'System Frequency' and bool(c.get('is_state')) == state]
        if not selected: paragraph('Data not available')
        for entity in event.get('report_entities', []):
            if (entity['group'] == 'State') != state or entity['group'] not in ('State', 'ISGS', 'IPP'): continue
            heading(entity['display_name'] + ' | Analysis window: ' + period)
            for chart in selected:
                if chart.get('entity_id') == entity['entity_id']: blocks.append(('image', chart))
            heading(entity['display_name'] + ' CRMS Messages Issued')
            rows = [[r.get('timestamp'), r.get('message_type'), r.get('message_no'), r.get('message_details')] for r in event.get('chronology', []) if r.get('entity_id') == entity['entity_id']]
            table(['Time (IST)', 'Message Type', 'Message No.', 'Details'], rows)
    return blocks, summary


def render(event, options, charts, fmt):
    blocks, summary = report_blocks(event, options, charts)
    title = 'LOW FREQUENCY OPERATION REPORT'
    def group_labels():
        return [f"Freq <{float(level):g} Hz" for level in LEVELS]
    def duration_labels():
        return [f"Total frequency duration: {text(summary.get('thresholds', {}).get(level, {}).get('frequency_minutes'))} min" for level in LEVELS]
    def widths(headers, grouped):
        if grouped: return [1.35,1.25,1.35]+[.68,.55,.55,.48]*3
        if len(headers)==7: return [1.3,.7,1.2,.7,1.1,1.1,3.65]
        if len(headers)==4: return [1.4,1.2,1.2,6.0]
        return [1]*len(headers)
    if fmt == 'html':
        out = ['<!doctype html><html><head><meta charset="utf-8"><title>'+title+'</title><style>body{font-family:Arial;margin:30px;color:#172b4d}h1,h2{color:#17365d}h1{text-align:center;font-size:28px}table{border-collapse:collapse;width:100%;font-size:11px;margin:12px 0}th,td{border:1px solid #cbd5e1;padding:6px;overflow-wrap:anywhere}th{background:#17365d;color:white}tr:nth-child(even){background:#f1f5f9}img{width:100%;max-height:480px;object-fit:contain}footer{text-align:center;color:#17365d} @media print{@page{size:A4 landscape}thead{display:table-header-group}img{break-inside:avoid}}</style></head><body><h1>'+title+'</h1>']
        for kind, value in blocks:
            if kind == 'heading': out.append('<h2>'+escape(value)+'</h2>')
            elif kind == 'paragraph': out.append('<p>'+escape(value).replace('\n', '<br>')+'</p>')
            elif kind == 'image': out.append('<p>'+escape(value['title'])+'</p><img src="data:image/png;base64,'+value['image']+'">')
            else:
                headers, rows, grouped = value
                weights = widths(headers, grouped)
                out.append('<table><colgroup>'+''.join(f'<col style="width:{100*w/sum(weights):.2f}%">' for w in weights)+'</colgroup><thead>')
                if grouped:
                    out.append('<tr>'+''.join('<th rowspan="3">'+escape(h)+'</th>' for h in headers[:3])+''.join('<th colspan="4">'+escape(label)+'</th>' for label in group_labels())+'</tr>')
                    out.append('<tr>'+''.join('<th colspan="4" style="background:#dbe5f1;color:#172b4d">'+escape(label)+'</th>' for label in duration_labels())+'</tr>')
                    out.append('<tr>'+''.join('<th style="background:#dbe5f1;color:#172b4d">'+escape(h)+'</th>' for h in headers[3:])+'</tr>')
                else: out.append('<tr>'+''.join('<th>'+escape(h)+'</th>' for h in headers)+'</tr>')
                out.append('</thead><tbody>')
                out.extend('<tr>'+''.join('<td>'+escape(text(v))+'</td>' for v in row)+'</tr>' for row in rows)
                if not rows: out.append(f'<tr><td colspan="{len(headers)}">Data not available</td></tr>')
                out.append('</tbody></table>')
        return ''.join(out)+'<footer>ERLDC</footer></body></html>'
    if fmt == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.section import WD_ORIENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from pathlib import Path
        doc = Document(Path(__file__).resolve().parents[1] / 'report_templates' / 'low_frequency_operation.docx')
        for node in list(doc._element.body):
            if node.tag != qn('w:sectPr'): doc._element.body.remove(node)
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69); section.page_height = Inches(8.27)
        section.left_margin = section.right_margin = Inches(.45)
        section.top_margin = section.bottom_margin = Inches(.5)
        doc.styles['Normal'].font.size = Pt(10)
        for name in ('Title', 'Heading 1', 'Heading 2'): doc.styles[name].font.color.rgb = RGBColor.from_string(NAVY)
        title_style = doc.styles['Title']
        title_style.font.size = Pt(21); title_style.font.bold = True
        title_style.paragraph_format.alignment = 1
        for border in title_style.element.xpath('./w:pPr/w:pBdr'): border.getparent().remove(border)
        doc.add_heading(title, 0)
        section.footer.paragraphs[0].text = 'Eastern Regional Load Despatch Centre (ERLDC)'
        section.footer.paragraphs[0].alignment = 1
        section.footer.paragraphs[0].add_run(' | ')
        page = OxmlElement('w:fldSimple'); page.set(qn('w:instr'), 'PAGE')
        section.footer.paragraphs[0]._p.append(page)
        for kind, value in blocks:
            if kind == 'heading':
                if value.startswith('Annexure'): doc.add_page_break()
                doc.add_heading(value, 1)
            elif kind == 'paragraph':
                paragraph = doc.add_paragraph(value)
                paragraph.paragraph_format.keep_with_next = value.startswith('Reporting period:')
            elif kind == 'image':
                caption = doc.add_paragraph(value['title'])
                caption.paragraph_format.keep_with_next = True
                doc.add_picture(io.BytesIO(base64.b64decode(value['image'])), width=Inches(10.5))
            else:
                headers, rows, grouped = value
                table = doc.add_table(rows=3 if grouped else 1, cols=len(headers)); table.style = 'Table Grid'
                table.autofit = False
                weights = widths(headers, grouped)
                for column, weight in zip(table.columns, weights): column.width = Inches(10.79*weight/sum(weights))
                for row in table.rows:
                    for cell, weight in zip(row.cells, weights): cell.width = Inches(10.79*weight/sum(weights))
                if grouped:
                    for at, h in enumerate(headers[:3]): table.cell(0, at).merge(table.cell(2, at)).text = h
                    for row_index, labels in enumerate((group_labels(), duration_labels())):
                        for at, label in enumerate(labels): table.cell(row_index, 3+4*at).merge(table.cell(row_index, 6+4*at)).text = label
                    for c, h in zip(table.rows[-1].cells[3:], headers[3:]): c.text = h
                else:
                    for c, h in zip(table.rows[0].cells, headers): c.text = h
                for row_index, row in enumerate(table.rows):
                    row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                    for column_index, c in enumerate(row.cells):
                        if grouped and row_index > 0 and column_index < 3: continue
                        light = grouped and row_index > 0
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DBE5F1' if light else NAVY); c._tc.get_or_add_tcPr().append(shade)
                        for p in c.paragraphs:
                            for r in p.runs: r.font.color.rgb = RGBColor(0,0,0) if light else RGBColor(255,255,255)
                for values in rows or [['Data not available'] + ['']*(len(headers)-1)]:
                    for c, v in zip(table.add_row().cells, values): c.text = text(v)
                for row in table.rows:
                    for c in row.cells:
                        for p in c.paragraphs:
                            for r in p.runs: r.font.size = Pt(7 if grouped else 9)
        output = io.BytesIO(); doc.save(output); output.seek(0); return output
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Image, PageBreak, Spacer
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from PIL import Image as PILImage
    styles = getSampleStyleSheet()
    styles['Heading1'].textColor = styles['Heading2'].textColor = colors.HexColor('#'+NAVY)
    styles['Heading1'].alignment = 1; styles['Heading1'].fontSize = 21
    styles['Heading2'].keepWithNext = True
    styles['Normal'].fontSize = 8; styles['Normal'].leading = 10
    styles.add(styles['Normal'].clone('ChartCaption', keepWithNext=True))
    story = [Paragraph(title, styles['Heading1'])]
    width = landscape(A4)[0]-48
    for kind, value in blocks:
        if kind == 'heading':
            if value.startswith('Annexure'): story.append(PageBreak())
            story.append(Paragraph(escape(value), styles['Heading2']))
        elif kind == 'paragraph': story.append(Paragraph(escape(value).replace('\n','<br/>'), styles['ChartCaption'] if value.startswith('Reporting period:') else styles['Normal']))
        elif kind == 'image':
            raw = base64.b64decode(value['image']); image = PILImage.open(io.BytesIO(raw))
            w, h = image.size; scale = min(width/w, 380/h)
            story.extend([Paragraph(escape(value['title']), styles['ChartCaption']), Image(io.BytesIO(raw), width=w*scale, height=h*scale)])
        else:
            headers, rows, grouped = value
            def cell(v): return Paragraph(escape(text(v)).replace('\n','<br/>'), styles['Normal'])
            data = [[cell(h) for h in headers]] + [[cell(v) for v in r] for r in rows or [['Data not available']+['']*(len(headers)-1)]]
            commands = [('GRID',(0,0),(-1,-1),.4,colors.lightgrey), ('VALIGN',(0,0),(-1,-1),'TOP'), ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#'+NAVY)), ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f1f5f9')])]
            if grouped:
                data[0] = [cell('')]*3 + [cell(h) for h in headers[3:]]
                data.insert(0, [cell('')]*3 + sum(([cell(label)]+[cell('')]*3 for label in duration_labels()), []))
                data.insert(0, [cell(h) for h in headers[:3]] + sum(([cell(label)]+[cell('')]*3 for label in group_labels()), []))
                commands += [('BACKGROUND',(3,1),(-1,2),colors.HexColor('#dbe5f1')), ('BACKGROUND',(0,0),(2,2),colors.HexColor('#'+NAVY))]
                commands += [('SPAN',(i,0),(i,2)) for i in range(3)]
                commands += [('SPAN',(3+4*i,r),(6+4*i,r)) for i in range(3) for r in (0,1)]
            for row_index, row in enumerate(data[:3 if grouped else 1]):
                for column_index, p in enumerate(row):
                    p.style = styles['Normal'].clone('Header')
                    p.style.textColor = colors.black if grouped and row_index > 0 and column_index >= 3 else colors.white
            weights = widths(headers, grouped)
            table = Table(data, colWidths=[width*w/sum(weights) for w in weights], repeatRows=3 if grouped else 1, splitInRow=1)
            table.setStyle(TableStyle(commands)); story.extend([table, Spacer(1,10)])
    output = io.BytesIO()
    def footer(canvas, doc):
        canvas.setFont('Helvetica',8); canvas.drawCentredString(landscape(A4)[0]/2,14,'ERLDC | '+str(doc.page))
    SimpleDocTemplate(output,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=30).build(story,onFirstPage=footer,onLaterPages=footer)
    output.seek(0); return output
