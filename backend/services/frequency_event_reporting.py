"""Saved-event supplements; existing chronology and report generators remain owners."""
import math
from collections import defaultdict, OrderedDict
from datetime import datetime, timedelta
from threading import Lock
from time import monotonic

_cache = OrderedDict()
_cache_lock = Lock()
GROUPS = ("State", "ISGS", "IPP")


def event_period(event):
    # Structured metadata is authoritative. Never guess from the display name.
    try:
        start = datetime.fromisoformat(str(event["start_time"]).replace("Z", "+00:00"))
        end = datetime.fromisoformat(str(event["end_time"]).replace("Z", "+00:00"))
        if start.tzinfo:
            from zoneinfo import ZoneInfo
            start = start.astimezone(ZoneInfo("Asia/Kolkata")).replace(tzinfo=None)
        if end.tzinfo:
            from zoneinfo import ZoneInfo
            end = end.astimezone(ZoneInfo("Asia/Kolkata")).replace(tzinfo=None)
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError("Saved event is missing valid structured start/end metadata.") from exc
    if end < start:
        raise ValueError("Saved event end precedes its start.")
    return start, end


def number(value):
    try:
        result = float(value) if value is not None and not isinstance(value, bool) else None
        return result if result is not None and math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def checked_event_period(event):
    """Structured bounds first; narrowly support old frequency instance names."""
    import re
    try:
        start, end = event_period(event)
        return start, end, "structured"
    except ValueError:
        if event.get("start_time") or event.get("end_time"):
            raise
    match = re.fullmatch(r"(?:Low|High) Freq (\d{2}-[A-Za-z]{3}-\d{2}) \((\d{2}:\d{2})-(\d{2}:\d{2})\)", event.get("name") or "")
    if not match:
        raise ValueError("Saved event has no usable period metadata.")
    day, first, last = match.groups()
    start = datetime.strptime(f"{day} {first}", "%d-%b-%y %H:%M")
    end = datetime.strptime(f"{day} {last}", "%d-%b-%y %H:%M")
    if end < start:
        end += timedelta(days=1)
    return start, end, "legacy_name"


def check_frequency_periods(db, periods):
    """One metadata-only DB read for all selections; never fetch saved series."""
    from routes.frequency_routes import EVENT_COLLECTION
    if not 1 <= len(periods) <= 1000:
        raise ValueError("Select between 1 and 1000 periods per check.")
    requested = []
    for period in periods:
        start, end = event_period(period)
        if end <= start:
            raise ValueError("Period end must be after its start.")
        requested.append((period, start, end))
    documents = db.db[EVENT_COLLECTION].find(
        {"event_id": {"$exists": True}, "data_points": {"$elemMatch": {"series.timestamps.0": {"$exists": True}}}},
        {"_id": 0, "event_id": 1, "name": 1, "start_time": 1, "end_time": 1, "dates": 1},
    )
    available = []
    for event in documents:
        try:
            start, end, source = checked_event_period(event)
        except ValueError:
            continue
        if end > start:
            available.append((event, start, end, source))
    output = []
    for period, start, end in requested:
        matches = []
        for event, first, last, source in available:
            left, right = max(start, first), min(end, last)
            if left < right:
                matches.append({"event_id": event["event_id"], "name": event.get("name"), "start_time": first.isoformat(), "end_time": last.isoformat(),
                    "coverage_start": left.isoformat(), "coverage_end": right.isoformat(), "metadata_source": source})
        matches.sort(key=lambda item: (datetime.fromisoformat(item["end_time"]) - datetime.fromisoformat(item["start_time"])).total_seconds())
        coverage = []
        for left, right in sorted((datetime.fromisoformat(item["coverage_start"]), datetime.fromisoformat(item["coverage_end"])) for item in matches):
            if coverage and left <= coverage[-1][1]:
                coverage[-1] = (coverage[-1][0], max(coverage[-1][1], right))
            else:
                coverage.append((left, right))
        missing, cursor = [], start
        for left, right in coverage:
            if cursor < left:
                missing.append((cursor, left))
            cursor = right
        if cursor < end:
            missing.append((cursor, end))
        covering = next((item for item in matches if datetime.fromisoformat(item["start_time"]) <= start and datetime.fromisoformat(item["end_time"]) >= end), None)
        encode = lambda ranges: [{"start_time": a.isoformat(), "end_time": b.isoformat()} for a, b in ranges]
        output.append({"id": period.get("id"), "start_time": start.isoformat(), "end_time": end.isoformat(),
            "status": "existing" if not missing else "partial" if coverage else "new", "matches": matches,
            "event_id": covering["event_id"] if covering else None, "event_name": covering["name"] if covering else None,
            "coverage": encode(coverage), "missing": encode(missing), "coverage_seconds": sum((b-a).total_seconds() for a,b in coverage),
            "completion_supported": False})
    return {"success": True, "periods": output}


def point_group(point, mapping):
    kind = str(point.get("type") or "").upper()
    if kind == "STATE" or point.get("is_state") or mapping.get("is_state"):
        return "State"
    if kind in {"ISGS", "IPP"}:
        return kind
    mapped = str(mapping.get("type") or "").upper()
    if mapped in {"ISGS", "IPP"}:
        return mapped
    # Match the existing generator table's fallback, retaining State Generator
    # classification rather than relabelling those units as state drawal.
    if mapped in {"STATE", "STATE_IPP"}:
        return None
    return "IPP"


def event_entities(db, event):
    mappings = list(db.map_collection.find({}, {"_id": 0}))
    by_id = defaultdict(list)
    for mapping in mappings:
        by_id[str(mapping.get("plant_id") or "")].append(mapping)
    entities = []
    for index, point in enumerate(event.get("data_points") or []):
        if point.get("is_frequency") or str(point.get("plant_id")) == "SYSTEM_FREQUENCY":
            continue
        choices = by_id.get(str(point.get("plant_id") or ""), [])
        stage = str(point.get("stage_id") or point.get("STAGE_ID") or "")
        mapping = next((item for item in choices if str(item.get("STAGE_ID") or "") == stage), choices[0] if choices else {})
        group = point_group(point, mapping)
        name = str(point.get("plant_name") or mapping.get("plant_name") or point.get("plant_id") or "Unnamed entity")
        if point.get("stage_name"):
            name += " / " + str(point["stage_name"])
        entities.append({"entity_id": f"{point.get('plant_id', '')}:{stage}:{index}", "display_name": name, "group": group, "point": point, "mapping": mapping})
    return entities


def timeline_aliases(db, event, entities=None):
    from routes.frequency_routes import _timeline_state_mappings, crms_text_list, normalize_crms_lookup
    aliases = defaultdict(list)
    entities = entities if entities is not None else event_entities(db, event)
    state_aliases = _timeline_state_mappings(db)
    for entity in entities:
        point, mapping = entity["point"], entity["mapping"]
        names = [entity["display_name"], point.get("plant_name"), mapping.get("plant_name"), point.get("plant_id"), point.get("state"), point.get("state_name"), mapping.get("mis_name"), mapping.get("wbes_name")]
        names += crms_text_list(mapping.get("crms_utility_name")) + crms_text_list(point.get("crms_utility_name"))
        if entity["group"] == "State":
            for alias, state in state_aliases.items():
                if normalize_crms_lookup(state.get("display_name")) == normalize_crms_lookup(entity["display_name"]) or str(state.get("plant_id")) == str(point.get("plant_id")):
                    names.append(alias)
        for name in set(normalize_crms_lookup(name) for name in names if name):
            aliases[name].append(entity)
    return aliases


def point_samples(point, start, end):
    series = point.get("series") or {}
    output = []
    def value(key, index):
        values = series.get(key) or []
        return number(values[index]) if index < len(values) else None
    for index, text in enumerate(series.get("timestamps") or []):
        try:
            stamp, _ = event_period({"start_time": text, "end_time": text})
        except ValueError:
            continue
        if not start <= stamp <= end:
            continue
        deviation = value("deviation", index)
        if deviation is None:
            actual, schedule = value("actual", index), value("schedule", index)
            if actual is not None and schedule is not None:
                deviation = actual - schedule
        output.append((stamp, value("frequency", index), deviation))
    return sorted(output, key=lambda sample: sample[0])


def timeline_point_values(entity, message_dt, event):
    start, end = event_period(event)
    if "_timeline_samples" not in entity:
        entity["_timeline_samples"] = point_samples(entity["point"], start, end)
    samples = entity["_timeline_samples"]
    if not samples:
        return None, None
    nearest = min(samples, key=lambda sample: abs((sample[0] - message_dt).total_seconds()))
    if abs((nearest[0] - message_dt).total_seconds()) > 900:
        return None, None
    return nearest[1], nearest[2]


def build_performance(event, entities, chronology, messages_complete=True):
    from routes.frequency_routes import compute_frequency_statistics
    start, end = event_period(event)
    result = {group: [] for group in GROUPS}
    lowest = None
    for entity in entities:
        samples = point_samples(entity["point"], start, end)
        for _, frequency, _ in samples:
            if frequency is not None:
                lowest = frequency if lowest is None else min(lowest, frequency)
        if not entity["group"]:
            continue
        blocks = defaultdict(list)
        for stamp, frequency, deviation in samples:
            block = stamp.replace(minute=stamp.minute // 15 * 15, second=0, microsecond=0)
            if stamp == end == block and end > start:
                block -= timedelta(minutes=15)
            blocks[block].append((stamp, frequency, deviation))
        for block, values in sorted(blocks.items()):
            valid = [sample for sample in values if sample[1] is not None and sample[2] is not None]
            is_state = entity["group"] == "State"
            stats = compute_frequency_statistics([sample[1] for sample in valid], [sample[0] for sample in valid], [sample[2] for sample in valid], is_state, event.get("event_type"))
            deviations = [sample[2] for sample in values if sample[2] is not None]
            offending = [(max(value, 0) if is_state else min(value, 0)) for value in deviations]
            frequencies = [sample[1] for sample in values if sample[1] is not None]
            count = len({(row.get("timestamp"), row.get("message_no")) for row in chronology if row.get("entity_id") == entity["entity_id"] and block <= datetime.fromisoformat(row["timestamp"]) and (datetime.fromisoformat(row["timestamp"]) < block + timedelta(minutes=15) or datetime.fromisoformat(row["timestamp"]) == end == block + timedelta(minutes=15))})
            result[entity["group"]].append({
                "entity_id": entity["entity_id"], "entity": entity["display_name"],
                "period_start": max(start, block).isoformat(timespec="seconds"), "period_end": min(end, block + timedelta(minutes=15)).isoformat(timespec="seconds"),
                "average_od_ui_mw": round(sum(deviations) / len(deviations), 3) if deviations else None,
                "od_ui_time_pct": round(stats["positive_pct"] if is_state else stats["negative_pct"], 3) if valid else None,
                "maximum_od_ui_mw": (max(offending) if is_state else min(offending)) if offending else None,
                "lowest_frequency": min(frequencies) if frequencies else None,
                "message_count": count if messages_complete else None,
                "sample_count": len(values), "valid_deviation_samples": len(deviations),
            })
    return result, lowest


async def saved_event_analysis(db, event_id, refresh=False):
    from routes.frequency_routes import EVENT_COLLECTION, FrequencyMessageTimelinePayload, FrequencyMessageRange, build_frequency_message_timeline
    event = db.db[EVENT_COLLECTION].find_one({"event_id": event_id}, {"_id": 0})
    if not event:
        raise ValueError(f"Saved event not found: {event_id}")
    if not event.get("data_points"):
        raise ValueError("This saved instance contains no entity datasets.")
    start, end = event_period(event)
    key = (event_id, str(event.get("updated_at")))
    with _cache_lock:
        cached = _cache.get(key)
        if not refresh and cached and monotonic() - cached[0] < 300:
            return cached[1], event
    timeline = await build_frequency_message_timeline(FrequencyMessageTimelinePayload(ranges=[FrequencyMessageRange(start_time=start.isoformat(), end_time=end.isoformat())], event_id=event_id))
    entities = event_entities(db, event)
    performance, lowest = build_performance(event, entities, timeline["rows"], timeline.get("messages_complete", True))
    from services.frequency_threshold_analysis import dataset_from_event, calculate
    try:
        threshold_analysis=calculate(dataset_from_event(db,event),[(start.isoformat(),end.isoformat())],timeline["rows"],timeline.get("messages_complete",True))
    except ValueError as exc:
        threshold_analysis=None
        timeline.setdefault("warnings",[]).append(f"Threshold analysis unavailable: {exc}")
    result = {
        "threshold_analysis": threshold_analysis,
        "event_id": event_id, "event_name": event.get("name") or event_id,
        "start_time": start.isoformat(timespec="seconds"), "end_time": end.isoformat(timespec="seconds"), "event_type": event.get("event_type") or "low",
        "lowest_frequency": lowest, "chronology": timeline["rows"], "performance": performance,
        "messages_complete": timeline.get("messages_complete", True), "warnings": timeline.get("warnings", []),
        "calculation_note": "Deviation = Actual - Schedule. State OD is positive; generator UI is negative. Block average is the existing signed Actual-minus-Schedule average; opposite deviations offset. Percentages use the existing frequency-qualified sample denominator. Periods are clipped to event bounds; partial blocks are not padded.",
    }
    if any(not entity["group"] for entity in entities):
        result["warnings"].append("State-sector generator rows retain their existing classification and are outside the State drawal / ISGS / IPP tables.")
    with _cache_lock:
        _cache[key] = (monotonic(), result)
        while len(_cache) > 32:
            _cache.popitem(last=False)
    return result, event


CHRONOLOGY_COLUMNS = [("timestamp", "Time (IST)"), ("frequency_hz", "Frequency (Hz)"), ("state", "State / Entity"), ("deviation_mw", "OD/UI (MW)"), ("message_type", "Message Type"), ("message_no", "Message No."), ("message_details", "Message / Details")]
PERFORMANCE_COLUMNS = [("entity", "Entity"), ("period", "Period (IST)"), ("average_od_ui_mw", "15-Min Average OD/UI (MW)"), ("od_ui_time_pct", "% Time OD/UI"), ("maximum_od_ui_mw", "Maximum OD/UI (MW)"), ("lowest_frequency", "Lowest Frequency (Hz)"), ("message_count", "No. of Messages")]


def threshold_columns(wide=False):
    if not wide:
        return [("entity", "Entity"), ("period", "Period (IST)"), ("threshold", "Frequency Below (Hz)"),
                ("frequency_minutes", "Freq Minutes"), ("adverse_minutes", "OD/UI Minutes"), ("adverse_pct", "OD/UI %"),
                ("average_od_ui_mw", "15-Min Avg OD/UI (MW)"), ("maximum_od_ui_mw", "Max OD/UI (MW)"),
                ("lowest_frequency", "Lowest Hz"), ("message_count", "Threshold Messages")]
    columns=[("entity", "Entity"), ("period", "Period (IST)")]
    for level in ("49.90", "49.70", "49.50"):
        columns.extend((f"{level}_{key}", f"<{level} {label}") for key,label in [("frequency_minutes","Freq Minutes"),("adverse_minutes","OD/UI Minutes"),("adverse_pct","OD/UI %"),("average_od_ui_mw","Avg OD/UI (MW)"),("maximum_od_ui_mw","Max OD/UI (MW)")])
    return columns+[("lowest_frequency","Lowest Hz"),("message_count","Messages")]


def threshold_rows(rows,wide=False):
    output=[]
    for row in rows:
        base={"entity":row["entity"],"period":row["period_start"].replace("T"," ")+" - "+row["period_end"].replace("T"," "),"lowest_frequency":row.get("lowest_frequency"),"message_count":row.get("message_count")}
        if wide:
            output.append({**base,**{f"{level}_{key}":value for level,values in row["thresholds"].items() for key,value in values.items()}})
        else:
            output.extend({**base,"threshold":level,**values} for level,values in row["thresholds"].items())
    return output


SUMMARY_COLUMNS=[("threshold","Frequency Below (Hz)"),("frequency_minutes","Frequency Minutes"),("occurrences","Occurrences"),("longest_minutes","Longest Occurrence (min)"),("lowest_frequency","Lowest Hz")]


def summary_rows(event):
    analysis=event.get("threshold_analysis") or {}
    return [{"threshold":level,**values} for level,values in analysis.get("summary",{}).get("thresholds",{}).items()]


def table_weights(columns,title):
    if len(columns)==7:
        return [1.25,.7,1.6,.8,1.4,1.2,3.75] if title=="Chronology of Messages" else [1.8,2.6,1.4,1,1.3,1.3,1.3]
    return [2 if key=="entity" else 2.5 if key=="period" else 1 for key,_ in columns]


def report_tables(event, payload):
    tables=[]
    analysis=event.get("threshold_analysis") if payload.get("include_threshold_performance",False) else None
    if payload.get("include_analysis_summary") and analysis:
        tables.append(("Overall Frequency Statistics",SUMMARY_COLUMNS,summary_rows(event)))
    if payload.get("include_chronology",False):
        tables.append(("Chronology of Messages",CHRONOLOGY_COLUMNS,event.get("chronology") or []))
    if payload.get("include_entity_performance",False):
        for group in payload.get("performance_groups",GROUPS):
            if group not in GROUPS:continue
            if analysis:
                if payload.get("include_analysis_summary"):
                    tables.append((f"{group} Overall Performance",threshold_columns(),threshold_rows(analysis["overall_performance"].get(group,[]))))
                tables.append((f"{group} Performance",threshold_columns(),threshold_rows(analysis["performance"].get(group,[]))))
            else:
                rows=[{**row,"period":row["period_start"].replace("T"," ")+" - "+row["period_end"].replace("T"," ")} for row in event.get("performance",{}).get(group,[])]
                tables.append((f"{group} Performance",PERFORMANCE_COLUMNS,rows))
    return tables


def cell_text(value):
    if value is None:
        return "-"
    if isinstance(value, float):
        return f"{value:.3f}".rstrip("0").rstrip(".")
    return str(value)


def append_docx_supplements(doc, payload):
    import docx.shared
    from docx.enum.section import WD_SECTION, WD_ORIENT
    from docx.oxml import OxmlElement
    events = payload.get("supplemental_events") or []
    if not events or not any(report_tables(event, payload) for event in events):
        return
    # Use the existing Word generator and its table style; add wide appendices.
    section = doc.add_section(WD_SECTION.NEW_PAGE) if payload.get("include_existing_sections", True) else doc.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = docx.shared.Inches(11.69), docx.shared.Inches(8.27)
    section.left_margin = section.right_margin = docx.shared.Inches(.45)
    section.top_margin = section.bottom_margin = docx.shared.Inches(.45)
    for index, event in enumerate(events):
        if index:
            doc.add_page_break()
        doc.add_heading(event["event_name"], level=1)
        doc.add_paragraph(f"Event Date: {event['start_time'][:10]} | Start Time: {event['start_time'].replace('T', ' ')} | End Time: {event['end_time'].replace('T', ' ')} | Lowest Frequency: {cell_text(event.get('lowest_frequency'))} Hz")
        if payload.get("include_analysis_summary") and event.get("threshold_analysis"):
            stats=event["threshold_analysis"]["summary"]
            doc.add_paragraph(f"Selected: {stats['selected_minutes']} min | Covered frequency: {stats['covered_frequency_minutes']} min | Average frequency: {cell_text(stats['average_frequency'])} Hz | Minimum at: {stats['minimum_timestamp']} | Low-frequency occurrences: {stats['low_frequency_events']} | Sampling interval(s): {stats.get('sampling_intervals_seconds', [stats['sampling_seconds']])} sec")
        if event.get("selection_note"):
            doc.add_paragraph(event["selection_note"])
        for warning in event.get("warnings", []):
            doc.add_paragraph(warning)
        if payload.get("include_entity_performance"):
            doc.add_paragraph(event["calculation_note"])
        for title, columns, rows in report_tables(event, payload):
            doc.add_heading(title, level=2)
            if not rows:
                doc.add_paragraph("No records available for this event period.")
                continue
            table = doc.add_table(rows=1, cols=len(columns))
            table.style = "Light Shading Accent 1"
            table.autofit = False
            widths = table_weights(columns, title)
            # Fit to the same usable landscape frame as the existing annexures.
            scale = 10.79 / sum(widths)
            for column, width in zip(table.columns, widths):
                column.width = docx.shared.Inches(width * scale)
            for cell, (_, label) in zip(table.rows[0].cells, columns):
                cell.text = label
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows:
                cells = table.add_row().cells
                for cell, (key, _) in zip(cells, columns):
                    cell.text = cell_text(row.get(key))
            for row in table.rows:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for cell, width in zip(row.cells, widths):
                    cell.width = docx.shared.Inches(width * scale)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = docx.shared.Pt(2)
                        for run in paragraph.runs:
                            run.font.size = docx.shared.Pt(8)


def append_pdf_supplements(story, payload, styles, table_style, pdf_cell, header_style):
    from reportlab.platypus import Paragraph, Table, PageBreak
    from xml.sax.saxutils import escape
    for index, event in enumerate(payload.get("supplemental_events") or []):
        if not report_tables(event, payload):
            continue
        if index or payload.get("include_existing_sections", True):
            story.append(PageBreak())
        story.append(Paragraph(escape(event["event_name"]), styles["Heading2"]))
        story.append(Paragraph(escape(f"Event Date: {event['start_time'][:10]} | Start Time: {event['start_time'].replace('T', ' ')} | End Time: {event['end_time'].replace('T', ' ')} | Lowest Frequency: {cell_text(event.get('lowest_frequency'))} Hz"), styles["Normal"]))
        for warning in event.get("warnings", []):
            story.append(Paragraph(escape(warning), styles["Normal"]))
        if payload.get("include_entity_performance"):
            story.append(Paragraph(escape(event["calculation_note"]), styles["Normal"]))
        for title, columns, rows in report_tables(event, payload):
            story.append(Paragraph(title, styles["Heading3"]))
            if not rows:
                story.append(Paragraph("No records available for this event period.", styles["Normal"]))
                continue
            values = [[pdf_cell(label, header_style) for _, label in columns]]
            values += [[pdf_cell(cell_text(row.get(key))) for key, _ in columns] for row in rows]
            weights = table_weights(columns, title)
            table = Table(values, colWidths=[weight / sum(weights) * 790 for weight in weights], repeatRows=1, splitInRow=1)
            table.setStyle(table_style)
            story.append(table)


def supplements_html(events, payload):
    from html import escape
    if not any(report_tables(event, payload) for event in events):
        return ""
    output = ['<section class="frequency-supplements" style="padding:24px;font-family:Arial,sans-serif;color:#102a43"><style>.frequency-supplements table{border-collapse:collapse;width:100%;font-size:12px;margin-bottom:24px}.frequency-supplements th,.frequency-supplements td{border:1px solid #CBD5E1;padding:7px;text-align:left;vertical-align:top;overflow-wrap:anywhere}.frequency-supplements th{background:#EAF2FF}.frequency-supplements h2{color:#03624C}.frequency-supplements .table-scroll{overflow-x:auto}</style>']
    for event in events:
        output.append(f"<h2>{escape(event['event_name'])}</h2><p>{escape(event['start_time'])} to {escape(event['end_time'])} IST | Lowest Frequency: {escape(cell_text(event.get('lowest_frequency')))} Hz</p>")
        if payload.get("include_analysis_summary") and event.get("threshold_analysis"):
            stats=event["threshold_analysis"]["summary"]
            output.append("<p>"+escape(f"Selected: {stats['selected_minutes']} min | Covered frequency: {stats['covered_frequency_minutes']} min | Average frequency: {cell_text(stats['average_frequency'])} Hz | Minimum at: {stats['minimum_timestamp']} | Low-frequency occurrences: {stats['low_frequency_events']} | Sampling interval(s): {stats.get('sampling_intervals_seconds', [stats['sampling_seconds']])} sec")+"</p>")
        if event.get("selection_note"):
            output.append("<p>"+escape(event["selection_note"])+"</p>")
        output.extend(f"<p>{escape(warning)}</p>" for warning in event.get("warnings", []))
        if payload.get("include_entity_performance"):
            output.append(f"<p>{escape(event['calculation_note'])}</p>")
        for title, columns, rows in report_tables(event, payload):
            output.append(f"<h3>{escape(title)}</h3><div class='table-scroll'><table><thead><tr>" + "".join(f"<th>{escape(label)}</th>" for _, label in columns) + "</tr></thead><tbody>")
            for row in rows:
                output.append("<tr>" + "".join(f"<td>{escape(cell_text(row.get(key)))}</td>" for key, _ in columns) + "</tr>")
            if not rows:
                output.append(f"<tr><td colspan='{len(columns)}'>No records available for this event period.</td></tr>")
            output.append("</tbody></table></div>")
    return "".join(output) + "</section>"


def supplements_excel(events, payload):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    specs = []
    threshold_mode=payload.get("include_threshold_performance",False) and any(event.get("threshold_analysis") for event in events)
    if payload.get("include_analysis_summary") and threshold_mode:
        specs.append(("Analysis Summary",[("metric","Metric"),("value","Value")],"metadata"))
        specs.append(("Overall Statistics",SUMMARY_COLUMNS,"summary"))
        if payload.get("include_entity_performance"):
            specs.extend((f"{group} Overall",threshold_columns(True),f"overall:{group}") for group in payload.get("performance_groups",GROUPS))
    if payload.get("include_chronology"):
        specs.append(("Chronology", CHRONOLOGY_COLUMNS, "chronology"))
    if payload.get("include_entity_performance"):
        specs.extend((f"{group} Performance", threshold_columns(True) if threshold_mode else PERFORMANCE_COLUMNS, group) for group in payload.get("performance_groups", GROUPS))
    for name, columns, group in specs:
        sheet = workbook.create_sheet(name)
        headers = ["Event / Instance", "Event Start (IST)", "Event End (IST)"] + [label for _, label in columns] + ["Data Quality Notes"]
        grouped=threshold_mode and group not in {"summary","metadata","chronology"}
        header_row=2 if grouped else 1
        if grouped:
            sheet.append(["Event / Entity"]+[None]*4+["Freq <49.90"]+[None]*4+["Freq <49.70"]+[None]*4+["Freq <49.50"]+[None]*4+["Other"]+[None]*2)
            for left,right in [(1,5),(6,10),(11,15),(16,20),(21,23)]:sheet.merge_cells(start_row=1,start_column=left,end_row=1,end_column=right)
        sheet.append(headers)
        for event in events:
            if group=="metadata":
                stats=event["threshold_analysis"]["summary"]
                rows=[{"metric":key.replace("_"," ").title(),"value":str(value) if isinstance(value,list) else value} for key,value in stats.items() if key!="thresholds"]
                rows.append({"metric":"Calculation", "value":event["calculation_note"]})
                if event.get("selection_note"):rows.append({"metric":"Selected Slots (IST)","value":event["selection_note"]})
            elif group=="summary":rows=summary_rows(event)
            elif group=="chronology":rows=event["chronology"]
            elif threshold_mode:
                analysis=event.get("threshold_analysis") or {}
                source=analysis.get("overall_performance" if group.startswith("overall:") else "performance",{})
                rows=threshold_rows(source.get(group.split(":")[-1],[]),True)
            else:rows=event["performance"].get(group,[])
            for row in rows:
                record = {**row}
                if not threshold_mode and group != "chronology":
                    record["period"] = row["period_start"].replace("T", " ") + " - " + row["period_end"].replace("T", " ")
                sheet.append([event["event_name"], event["start_time"], event["end_time"]] + [record.get(key) for key, _ in columns] + ["; ".join(event.get("warnings", []))])
            # Retain an event's identity even when its selected section is empty.
            if not rows:
                sheet.append([event["event_name"], event["start_time"], event["end_time"], "No records available"] + [None] * (len(columns) - 1) + ["; ".join(event.get("warnings", []))])
        for cell in list(sheet[header_row])+(list(sheet[1]) if grouped else []):
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="03624C")
        for row in sheet.iter_rows(min_row=header_row+1):
            for cell in row:
                if isinstance(cell.value, str):
                    # Message text/labels must stay text, including leading '='.
                    cell.data_type = "s"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.000"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, label in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 54 if label in {"Message / Details", "Period (IST)"} else 34 if label in {"Event / Instance", "State / Entity", "Entity"} else 24
        sheet.freeze_panes = "F3" if grouped else "D2"
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{sheet.max_row}"
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = f"1:{header_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    if not workbook.sheetnames:
        raise ValueError("Select chronology or entity performance for Excel export.")
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
