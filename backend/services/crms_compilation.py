"""Chronological state-message compilation; no changes to source records."""
from datetime import datetime, timedelta, timezone
from math import isfinite
import re

IST = timezone(timedelta(hours=5, minutes=30))
ALIASES = {
    "Bihar": ("BIHAR", "BSPTCL"), "DVC": ("DVC",),
    "Jharkhand": ("JHARKHAND", "JUSNL"), "Odisha": ("ODISHA", "GRIDCO"),
    "Sikkim": ("SIKKIM",), "West Bengal": ("WEST BENGAL", "WBSETCL", "WB"),
}
HEADERS = ["Time stamp", "Frequency (Hz)", "State / Control Area",
           "Overdrawal / Deviation (MW)", "Message type", "Message No."]


def token(value):
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def number(value):
    try:
        result = float(value)
        return result if isfinite(result) else None
    except (TypeError, ValueError):
        return None


def local_datetime(value):
    dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return dt.astimezone(IST).replace(tzinfo=None) if dt.tzinfo else dt


def normalize_ranges(ranges):
    if not isinstance(ranges, list) or not 1 <= len(ranges) <= 50:
        raise ValueError("Add between 1 and 50 date/time ranges.")
    intervals = []
    days = set()
    for item in ranges:
        try:
            start, end = local_datetime(item["start"]), local_datetime(item["end"])
        except (ValueError, TypeError, KeyError):
            raise ValueError("Each range needs valid start and end dates/times.")
        if start.second or end.second or start.microsecond or end.microsecond:
            raise ValueError("Choose times to the minute, without seconds.")
        if end < start:
            raise ValueError("End must be on or after start; select the next date for an overnight range.")
        # Include the whole final minute: a message at 19:30:45 belongs to 19:30.
        intervals.append((start, end + timedelta(minutes=1)))
        day = start.date()
        while day <= end.date():
            days.add(day.isoformat())
            if len(days) > 31:
                raise ValueError("Select at most 31 distinct dates per report.")
            day += timedelta(days=1)
    merged = []
    for start, end in sorted(intervals):
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(end, merged[-1][1]))
        else:
            merged.append((start, end))
    # Split midnight crossings so each HH:mm timestamp has an explicit date.
    sections = []
    for start, end in merged:
        while start < end:
            stop = min(end, datetime.combine(start.date() + timedelta(days=1), datetime.min.time()))
            last = stop - timedelta(minutes=1)
            sections.append({"start": start, "stop": stop,
                "label": f"{start:%d.%m.%Y} {start:%H:%M}-{last:%H:%M}hrs"})
            start = stop
    return sections, sorted(days)


def state_mappings(mappings):
    by_alias, by_name = {}, {}
    for name, aliases in ALIASES.items():
        for alias in aliases:
            by_alias[token(alias)] = name
        by_name[name] = {"plant_id": "STATE_" + name.upper().replace(" ", "_")}
    for mapping in mappings:
        if not mapping.get("is_state") or mapping.get("is_frequency"):
            continue
        name = by_alias.get(token(mapping.get("plant_name")))
        if not name:
            # Only explicitly mapped control areas, never the regional total.
            if not mapping.get("crms_utility_name"):
                continue
            name = str(mapping.get("plant_name") or mapping.get("plant_id"))
        by_name[name] = mapping
        for alias in [mapping.get("plant_name"), *re.split(r"[;,|/]+", str(mapping.get("crms_utility_name") or ""))]:
            if token(alias):
                by_alias[token(alias)] = name
    return by_alias, by_name


class Measurements:
    def __init__(self, raw_docs, events):
        self.raw = {(doc.get("date"), str(doc.get("plant_id"))): doc for doc in raw_docs}
        self.saved = {}
        # Caller supplies newest events first; overlapping older events cannot overwrite them.
        for event in events:
            for point in event.get("data_points") or []:
                series = point.get("series") or {}
                for i, stamp in enumerate(series.get("timestamps") or []):
                    try:
                        dt = local_datetime(stamp)
                    except (ValueError, TypeError):
                        continue
                    if dt.second or dt.microsecond:
                        continue
                    key = (dt.strftime("%Y-%m-%d %H:%M"), str(point.get("plant_id")))
                    values = {field: number(items[i]) for field in ("frequency", "actual", "schedule", "deviation")
                              if isinstance((items := series.get(field)), list) and i < len(items)}
                    self.saved.setdefault(key, values)
                    if values.get("frequency") is not None:
                        self.saved.setdefault((key[0], "SYSTEM_FREQUENCY"), {"actual": values["frequency"]})

    def value(self, dt, plant_id, field, sources, allow_block=False):
        doc = self.raw.get((dt.date().isoformat(), str(plant_id)), {})
        for source in sources:
            values = ((doc.get("sources") or {}).get(source) or {}).get(field)
            if not isinstance(values, list):
                continue
            if len(values) == 1440:
                index, resolution = dt.hour * 60 + dt.minute, "minute"
            elif allow_block and len(values) == 96:
                index, resolution = (dt.hour * 60 + dt.minute) // 15, "15-minute schedule"
            else:
                continue
            result = number(values[index])
            if result is not None:
                return result, f"{source}: {resolution}"
        result = self.saved.get((dt.strftime("%Y-%m-%d %H:%M"), str(plant_id)), {}).get(field)
        return result, "Saved frequency event" if result is not None else "Missing"

    def at(self, dt, mapping):
        pid = mapping["plant_id"]
        frequency, freq_source = self.value(dt, "SYSTEM_FREQUENCY", "actual", ("scada", "scada_file"))
        actual, actual_source = self.value(dt, pid, "actual", ("scada_file", "scada"))
        schedule, schedule_source = self.value(dt, pid, "schedule", ("scada_file", "wbes", "rtg"), allow_block=True)
        deviation = actual - schedule if actual is not None and schedule is not None else None
        return {"frequency": round(frequency, 3) if frequency is not None else None,
                "deviation": round(deviation) if deviation is not None else None,
                "frequency_source": freq_source, "deviation_source": f"{actual_source} actual - {schedule_source}"}


def compile_messages(sections, messages, mappings, measurements, warnings=None):
    aliases, states = state_mappings(mappings)
    output = [{"label": section["label"], "start": section["start"].isoformat(),
               "end": (section["stop"] - timedelta(minutes=1)).isoformat(), "rows": []} for section in sections]
    seen, included_messages, excluded = set(), set(), set()
    for message in messages:
        dt = local_datetime(message["timestamp"])
        index = next((i for i, section in enumerate(sections) if section["start"] <= dt < section["stop"]), None)
        if index is None:
            continue
        recipients = message.get("issued_to") or []
        if isinstance(recipients, str):
            recipients = re.split(r"[;,]+", recipients)
        category = message.get("category") or []
        if isinstance(category, str):
            category = [category]
        category = ", ".join(dict.fromkeys(str(c).strip() for c in category if str(c).strip())) or "Not specified"
        for recipient in recipients:
            state = aliases.get(token(recipient))
            if not state:
                excluded.add(str(recipient))
                continue
            identity = (message.get("message_no") or "", message["timestamp"], state, category)
            if identity in seen:
                continue
            seen.add(identity)
            included_messages.add(identity[:2])
            output[index]["rows"].append({"timestamp": dt.isoformat(), "time": dt.strftime("%H:%M"),
                "state": state, "message_type": category, "message_no": message.get("message_no") or "Not specified",
                **measurements.at(dt, states[state])})
    for section in output:
        section["rows"].sort(key=lambda row: (row["timestamp"], row["state"], row["message_no"]))
    rows = [row for section in output for row in section["rows"]]
    missing = sum(row["frequency"] is None or row["deviation"] is None for row in rows)
    notes = list(warnings or [])
    if missing:
        notes.append(f"{missing} row(s) have missing measurements. Fetch/upload the date's SCADA data in Frequency Event Analysis and generate again. Missing values are shown as N/A.")
    if any("15-minute" in row["deviation_source"] for row in rows):
        notes.append("Some deviations use the applicable 15-minute schedule with minute actuals; row source details identify these values.")
    return {"sections": output, "row_count": len(rows), "message_count": len(included_messages),
            "missing_count": missing, "warnings": notes, "excluded_recipients": sorted(excluded),
            "generated_at": datetime.now(IST).isoformat(), "timezone": "Asia/Kolkata"}


def export_report(report, format):
    import io
    from xml.sax.saxutils import escape
    stream = io.BytesIO()
    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "CRMS Messages"
        for section in report["sections"]:
            ws.append([section["label"]])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=6)
            ws.cell(ws.max_row, 1).font = Font(bold=True, color="FFFFFF")
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor="12386B")
            ws.append(HEADERS)
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True)
            for row in section["rows"]:
                ws.append([row["time"], row["frequency"] if row["frequency"] is not None else "N/A",
                           row["state"], row["deviation"] if row["deviation"] is not None else "N/A",
                           row["message_type"], row["message_no"]])
                # Source strings are text, never Excel formulas.
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                ws.cell(ws.max_row, 2).number_format = "0.000"
                ws.cell(ws.max_row, 4).number_format = "#,##0"
            if not section["rows"]:
                ws.append(["No state/control-area messages in this range."])
            ws.append([])
        for i, width in enumerate((15, 18, 25, 29, 26, 35), 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A3"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        notes = wb.create_sheet("Sources and notes")
        notes.append(["Times are IST. Deviation = actual - schedule; positive means overdrawal."])
        notes.append(["Frequency: exact message minute. N/A means source measurements are unavailable."])
        for warning in report.get("warnings", []):
            notes.append([warning])
        if report.get("excluded_recipients"):
            notes.append(["Recipients outside mapped state/control-area scope: " + ", ".join(report["excluded_recipients"])])
        notes.append(["Timestamp", "State", "Frequency source", "Deviation source"])
        for section in report["sections"]:
            for row in section["rows"]:
                notes.append([row["timestamp"], row["state"], row.get("frequency_source", ""), row.get("deviation_source", "")])
        for cells in notes:
            for cell in cells:
                if isinstance(cell.value, str): cell.data_type = "s"
        wb.save(stream)
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        styles = getSampleStyleSheet()
        small = styles["BodyText"]
        small.fontSize, small.leading = 8, 11
        story = [Paragraph("CRMS Message Compilation", styles["Title"]),
                 Paragraph("Times: IST. Deviation = actual - schedule (positive: overdrawal). N/A: measurement unavailable.", small), Spacer(1, 10)]
        for warning in report.get("warnings", []):
            story.extend([Paragraph(escape(warning), small), Spacer(1, 6)])
        for section in report["sections"]:
            story.append(Paragraph(escape(section["label"]), styles["Heading2"]))
            table_rows = [HEADERS]
            for row in section["rows"]:
                table_rows.append([row["time"], f'{row["frequency"]:.3f}' if row["frequency"] is not None else "N/A",
                    row["state"], f'{row["deviation"]:,.0f}' if row["deviation"] is not None else "N/A", row["message_type"], row["message_no"]])
            if not section["rows"]:
                table_rows.append(["No messages", "", "", "", "", ""])
            table = Table([[Paragraph(escape(str(value)), small) for value in row] for row in table_rows],
                          colWidths=[64, 77, 113, 116, 119, 190], repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F0FA")),
                ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#CCD6E4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            story.extend([table, Spacer(1, 12)])
        SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=25, bottomMargin=25).build(story)
    stream.seek(0)
    return stream
