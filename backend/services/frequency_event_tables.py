"""Generic, reusable Dynamic Frequency Event Analysis Table engine.

Generates dynamic frequency-event analysis tables (Table-2, Table-2A, Table-2B, Table-2C)
for continuous frequency dips below selected thresholds (<49.50 Hz, <49.70 Hz, <49.90 Hz)
with constituent Max OD (positive deviation only) and Avg Dev (all deviation points).

Follows ERLDC/GRID-INDIA operational reporting standards.
"""

from io import BytesIO
from datetime import datetime
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from services.od_frequency_report import STATE_COLUMNS, load_od_frequency_dataframe

DEFAULT_CONSTITUENTS = ["Bihar", "Jharkhand", "Odisha", "West Bengal", "DVC"]
DEFAULT_THRESHOLDS = [49.5, 49.7, 49.9]


def _format_date_range(start_dt, end_dt):
    if not start_dt or not end_dt:
        return ""
    try:
        s = pd.to_datetime(start_dt)
        e = pd.to_datetime(end_dt)
        return f"{s.strftime('%d.%m.%Y')} to {e.strftime('%d.%m.%Y')}"
    except Exception:
        return f"{str(start_dt)[:10]} to {str(end_dt)[:10]}"


def identify_frequency_events(timestamps, frequencies, threshold, max_gap_seconds=90):
    """
    Identify continuous frequency dip events where frequency < threshold.
    
    A continuous event consists of consecutive 1-minute data points satisfying
    the threshold. If frequency becomes >= threshold or there is a time gap > 90s,
    the event ends and a new one starts on the next dip.
    """
    events = []
    current = []
    
    for i in range(len(timestamps)):
        t = pd.to_datetime(timestamps[i]) if pd.notna(timestamps[i]) else None
        f = frequencies[i]
        
        # Check if frequency is strictly below threshold and valid
        is_below = pd.notna(f) and np.isfinite(f) and (45.0 <= f < threshold) and (t is not None)
        
        if is_below:
            # Check for gap between consecutive records
            if current:
                prev_t = current[-1]['time']
                gap = (t - prev_t).total_seconds()
                if gap > max_gap_seconds:
                    events.append(current)
                    current = []
            current.append({'index': i, 'time': t, 'freq': float(f)})
        else:
            if current:
                events.append(current)
                current = []
                
    if current:
        events.append(current)
        
    return events


def extract_event_rows(df, events, threshold, constituents):
    """
    Build row records for an identified list of frequency dip events.
    
    Columns:
    - Sr. No
    - Date (e.g. 01-Sep)
    - Window (Hrs) (e.g. 19:15–19:19)
    - Min <{threshold} (duration count)
    - Lowest (Hz) (3 decimal places)
    - At (Hrs) (time lowest frequency occurred)
    - Constituent Max OD (positive deviation only or '–') & Avg Dev (arithmetic mean)
    """
    rows = []
    thresh_str = f"{threshold:.2f}"
    
    for sr_no, ev in enumerate(events, 1):
        start_time = ev[0]['time']
        end_time = ev[-1]['time']
        duration = len(ev)
        
        # Find lowest frequency and timestamp of lowest occurrence
        lowest_freq = min(p['freq'] for p in ev)
        lowest_time = next(p['time'] for p in ev if p['freq'] == lowest_freq)
        
        indices = [p['index'] for p in ev]
        
        date_str = start_time.strftime('%d-%b')
        if start_time.date() != end_time.date():
            date_str = f"{start_time.strftime('%d-%b')} / {end_time.strftime('%d-%b')}"
            
        window_str = f"{start_time.strftime('%H:%M')}–{end_time.strftime('%H:%M')}"
        at_str = lowest_time.strftime('%H:%M')
        
        constituent_metrics = {}
        
        for state in constituents:
            sched_col = None
            act_col = None
            
            # Find schedule and actual columns in dataframe
            if state in STATE_COLUMNS:
                s_name, a_name = STATE_COLUMNS[state]
                if s_name in df.columns:
                    sched_col = s_name
                if a_name in df.columns:
                    act_col = a_name
            
            # Fallback search if column casing or name differs
            if not sched_col or not act_col:
                for col in df.columns:
                    c_up = col.upper()
                    st_up = state.upper()
                    if st_up in c_up:
                        if 'SCHED' in c_up and not sched_col:
                            sched_col = col
                        elif 'ACT' in c_up and not act_col:
                            act_col = col
                            
            if sched_col and act_col:
                sched_vals = pd.to_numeric(df[sched_col].iloc[indices], errors='coerce')
                act_vals = pd.to_numeric(df[act_col].iloc[indices], errors='coerce')
                dev_vals = act_vals - sched_vals
                
                pos_dev = dev_vals[dev_vals > 0]
                if len(pos_dev):
                    max_od_val = float(pos_dev.max())
                    max_od_str = f"{round(max_od_val):.0f}"
                else:
                    max_od_val = None
                    max_od_str = "–"
                    
                valid_dev = dev_vals.dropna()
                if len(valid_dev):
                    avg_dev_val = float(valid_dev.mean())
                    avg_dev_str = f"{round(avg_dev_val):.0f}"
                else:
                    avg_dev_val = None
                    avg_dev_str = "–"
            else:
                max_od_val = None
                max_od_str = "–"
                avg_dev_val = None
                avg_dev_str = "–"
                
            constituent_metrics[state] = {
                "max_od": int(round(max_od_val)) if max_od_val is not None else "–",
                "max_od_val": max_od_val,
                "max_od_str": max_od_str,
                "avg_dev": int(round(avg_dev_val)) if avg_dev_val is not None else "–",
                "avg_dev_val": avg_dev_val,
                "avg_dev_str": avg_dev_str
            }
            
        rows.append({
            "sr_no": sr_no,
            "date": date_str,
            "window": window_str,
            "window_hrs": window_str,
            "duration_min": duration,
            "lowest_hz": round(lowest_freq, 3),
            "lowest_freq": round(lowest_freq, 3),
            "lowest_hz_str": f"{lowest_freq:.3f}",
            "lowest_at": at_str,
            "lowest_at_time": at_str,
            "start_iso": start_time.isoformat(),
            "end_iso": end_time.isoformat(),
            "constituents": constituent_metrics,
            "constituents_data": constituent_metrics,
        })
        
    return rows


def generate_frequency_event_tables(
    source,
    start_date=None,
    end_date=None,
    thresholds=None,
    constituents=None
):
    """
    Main engine for generating dynamic frequency event analysis tables.
    
    Args:
        source: Excel path, BytesIO, bytes, or DataFrame.
        start_date: Filter start date (YYYY-MM-DD or string).
        end_date: Filter end date (YYYY-MM-DD or string).
        thresholds: List of float thresholds (e.g. [49.5, 49.7, 49.9]).
        constituents: List of constituent names (e.g. ["Bihar", "Jharkhand", ...]).
        
    Returns:
        {
          "success": True,
          "period": "01.09.2026 to 08.10.2026",
          "constituents": [...],
          "thresholds": [...],
          "tables": [
             {
               "table_id": "Table-2A",
               "threshold": 49.5,
               "threshold_display": "49.50",
               "title": "Table-2A: Frequency dips below 49.50 Hz ...",
               "rows": [...]
             },
             ...
          ]
        }
    """
    if thresholds is None:
        thresholds = [49.5]
    if isinstance(thresholds, (int, float)):
        thresholds = [float(thresholds)]
    thresholds = sorted([float(t) for t in thresholds])
    
    if constituents is None:
        constituents = list(DEFAULT_CONSTITUENTS)
        
    # Load DataFrame
    if isinstance(source, pd.DataFrame):
        df = source.copy()
    else:
        df = load_od_frequency_dataframe(source).copy()
        
    # Locate timestamp and frequency columns
    ts_col = None
    for col in df.columns:
        if any(keyword in col.upper() for keyword in ["DATE & TIME", "TIMESTAMP", "DATE", "TIME", "UNNAMED: 0"]):
            ts_col = col
            break
    if not ts_col:
        ts_col = df.columns[0]
        
    if "FREQUENCY" not in df.columns:
        raise ValueError("The workbook does not contain a 'FREQUENCY' column.")
        
    df['DATETIME_PARSED'] = pd.to_datetime(df[ts_col], errors='coerce')
    df = df.dropna(subset=['DATETIME_PARSED']).sort_values('DATETIME_PARSED').reset_index(drop=True)
    
    # Filter by date range if provided
    if start_date:
        start_ts = pd.to_datetime(start_date)
        df = df[df['DATETIME_PARSED'] >= start_ts]
    if end_date:
        end_ts = pd.to_datetime(end_date)
        # Include full end day if end_ts has no time component
        if end_ts.hour == 0 and end_ts.minute == 0 and end_ts.second == 0:
            end_ts = end_ts + pd.Timedelta(days=1)
        df = df[df['DATETIME_PARSED'] <= end_ts]
        
    if df.empty:
        raise ValueError("No 1-minute frequency data available for the selected period.")
        
    actual_start = df['DATETIME_PARSED'].iloc[0]
    actual_end = df['DATETIME_PARSED'].iloc[-1]
    period_str = _format_date_range(actual_start, actual_end)
    
    timestamps = df['DATETIME_PARSED'].to_numpy()
    frequencies = pd.to_numeric(df['FREQUENCY'], errors='coerce').to_numpy(dtype=float)
    
    tables = []
    multi_table = len(thresholds) > 1
    
    for idx, threshold in enumerate(thresholds):
        letter = chr(ord('A') + idx) if multi_table else ""
        table_id = f"Table-2{letter}" if multi_table else "Table-2"
        thresh_fmt = f"{threshold:.2f}"
        thresh_col_name = f"Min <{threshold:g}"
        
        title = f"{table_id}: Frequency dips below {thresh_fmt} Hz during {period_str} with maximum overdraw and average deviation (MW) of the Eastern Region constituents during the event"
        
        # Identify continuous dip events
        raw_events = identify_frequency_events(timestamps, frequencies, threshold)
        rows = extract_event_rows(df, raw_events, threshold, constituents)
        
        tables.append({
            "table_id": table_id,
            "threshold": threshold,
            "threshold_display": thresh_fmt,
            "threshold_col_name": thresh_col_name,
            "title": title,
            "period": period_str,
            "constituents": constituents,
            "rows": rows,
            "total_events": len(rows),
            "total_minutes": sum(r["duration_min"] for r in rows),
        })
        
    return {
        "success": True,
        "period": period_str,
        "constituents": constituents,
        "thresholds": thresholds,
        "tables": tables
    }


def generate_event_tables_excel(tables_data):
    """
    Generate a publication-grade Excel workbook for the frequency event tables.
    Uses grouped constituent headers, ERLDC corporate navy styling, and proper numeric alignment.
    """
    wb = Workbook()
    wb.remove(wb.active)  # Remove default blank sheet
    
    title_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="17365D", end_color="17365D", fill_type="solid")
    
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="17365D", end_color="17365D", fill_type="solid")
    sub_header_fill = PatternFill(start_color="204B7D", end_color="204B7D", fill_type="solid")
    
    data_font = Font(name="Segoe UI", size=10)
    bold_data_font = Font(name="Segoe UI", size=10, bold=True)
    
    thin_border_side = Side(style='thin', color='CBD5E1')
    table_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    for table_info in tables_data.get("tables", []):
        sheet_name = f"{table_info['table_id']} (<{table_info['threshold_display']} Hz)"[:31]
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True
        
        constituents = table_info["constituents"]
        total_cols = 6 + len(constituents) * 2
        
        # Row 1: Title Banner
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        title_cell = ws.cell(row=1, column=1, value=table_info["title"])
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 40
        
        # Row 2 & 3: Two-level headers
        # Level 1 Base Columns
        base_cols = [
            ("Sr. No", 1),
            ("Date", 2),
            ("Window (Hrs)", 3),
            (table_info["threshold_col_name"], 4),
            ("Lowest (Hz)", 5),
            ("At (Hrs)", 6)
        ]
        
        for col_name, col_idx in base_cols:
            ws.merge_cells(start_row=2, start_column=col_idx, end_row=3, end_column=col_idx)
            cell = ws.cell(row=2, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = table_border
            # Bottom merged cell border
            ws.cell(row=3, column=col_idx).border = table_border
            
        # Constituent Columns
        col_cursor = 7
        for state in constituents:
            # Level 1 merged across 2 columns
            ws.merge_cells(start_row=2, start_column=col_cursor, end_row=2, end_column=col_cursor + 1)
            state_cell = ws.cell(row=2, column=col_cursor, value=state)
            state_cell.font = header_font
            state_cell.fill = header_fill
            state_cell.alignment = Alignment(horizontal="center", vertical="center")
            state_cell.border = table_border
            ws.cell(row=2, column=col_cursor + 1).border = table_border
            
            # Level 2 sub-headers
            max_cell = ws.cell(row=3, column=col_cursor, value="Max OD")
            max_cell.font = header_font
            max_cell.fill = sub_header_fill
            max_cell.alignment = Alignment(horizontal="center", vertical="center")
            max_cell.border = table_border
            
            avg_cell = ws.cell(row=3, column=col_cursor + 1, value="Avg Dev")
            avg_cell.font = header_font
            avg_cell.fill = sub_header_fill
            avg_cell.alignment = Alignment(horizontal="center", vertical="center")
            avg_cell.border = table_border
            
            col_cursor += 2
            
        ws.row_dimensions[2].height = 24
        ws.row_dimensions[3].height = 22
        
        # Populate Data Rows
        row_cursor = 4
        for row in table_info.get("rows", []):
            fill = zebra_fill if (row_cursor % 2 == 0) else white_fill
            
            # Base data
            c_sr = ws.cell(row=row_cursor, column=1, value=row["sr_no"])
            c_date = ws.cell(row=row_cursor, column=2, value=row["date"])
            c_win = ws.cell(row=row_cursor, column=3, value=row["window"])
            c_dur = ws.cell(row=row_cursor, column=4, value=row["duration_min"])
            c_low = ws.cell(row=row_cursor, column=5, value=row["lowest_hz"])
            c_at = ws.cell(row=row_cursor, column=6, value=row["lowest_at"])
            
            c_sr.alignment = Alignment(horizontal="center", vertical="center")
            c_date.alignment = Alignment(horizontal="center", vertical="center")
            c_win.alignment = Alignment(horizontal="center", vertical="center")
            c_dur.alignment = Alignment(horizontal="center", vertical="center")
            c_low.alignment = Alignment(horizontal="center", vertical="center")
            c_at.alignment = Alignment(horizontal="center", vertical="center")
            
            c_low.number_format = "0.000"
            
            for c in (c_sr, c_date, c_win, c_dur, c_low, c_at):
                c.font = data_font
                c.fill = fill
                c.border = table_border
                
            # Constituent data
            st_col = 7
            for state in constituents:
                state_data = row["constituents"].get(state, {})
                max_od = state_data.get("max_od")
                avg_dev = state_data.get("avg_dev")
                
                c_max = ws.cell(row=row_cursor, column=st_col, value=max_od if max_od is not None else "–")
                c_avg = ws.cell(row=row_cursor, column=st_col + 1, value=avg_dev if avg_dev is not None else "–")
                
                c_max.alignment = Alignment(horizontal="center", vertical="center")
                c_avg.alignment = Alignment(horizontal="center", vertical="center")
                
                if isinstance(max_od, (int, float)):
                    c_max.number_format = "#,##0"
                if isinstance(avg_dev, (int, float)):
                    c_avg.number_format = "#,##0"
                    
                c_max.font = data_font
                c_avg.font = data_font
                c_max.fill = fill
                c_avg.fill = fill
                c_max.border = table_border
                c_avg.border = table_border
                
                st_col += 2
                
            ws.row_dimensions[row_cursor].height = 20
            row_cursor += 1
            
        if not table_info.get("rows"):
            ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=total_cols)
            no_data = ws.cell(row=4, column=1, value=f"No continuous frequency dips below {table_info['threshold_display']} Hz recorded during the selected period.")
            no_data.alignment = Alignment(horizontal="center", vertical="center")
            no_data.font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
            ws.row_dimensions[4].height = 25
            
        # Adjust Column Widths
        for col_idx in range(1, total_cols + 1):
            letter = get_column_letter(col_idx)
            if col_idx == 1:
                ws.column_dimensions[letter].width = 8
            elif col_idx in (2, 3):
                ws.column_dimensions[letter].width = 14
            elif col_idx in (4, 6):
                ws.column_dimensions[letter].width = 12
            elif col_idx == 5:
                ws.column_dimensions[letter].width = 14
            else:
                ws.column_dimensions[letter].width = 13
                
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
