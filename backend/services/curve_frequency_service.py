"""Frequency-only extraction through the existing PSP Curve reader."""
import math
from collections import OrderedDict
from threading import RLock
from time import monotonic
import json
from datetime import date, datetime, time, timedelta

from openpyxl.utils.datetime import from_excel
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def curve_timestamp(value, day, sample_index):
    """Interpret Curve column C, deriving empty cells from its 30-second grid."""
    expected = datetime.combine(day, time()) + timedelta(seconds=sample_index * 30)
    if value is None or value == "":
        return expected, True
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        value = from_excel(value)
    if isinstance(value, str):
        text = value.strip()
        try:
            value = datetime.fromisoformat(text)
        except ValueError:
            try:
                value = time.fromisoformat(text)
            except ValueError as exc:
                raise ValueError(f"Invalid timestamp in 30SEC!C{sample_index + 8}.") from exc
    if isinstance(value, datetime):
        stamp = value.replace(tzinfo=None)
    elif isinstance(value, time):
        stamp = datetime.combine(day, value.replace(tzinfo=None))
    elif isinstance(value, date):
        stamp = datetime.combine(value, time())
    else:
        raise ValueError(f"Invalid timestamp in 30SEC!C{sample_index + 8}.")
    # Excel fractional-day arithmetic can introduce microsecond roundoff.
    stamp = stamp.replace(microsecond=0)
    if stamp != expected:
        raise ValueError(f"30SEC!C{sample_index + 8} does not match the Curve date's 30-second timeline.")
    return stamp, False


def extract_curve_frequency_rows(rows, date_str):
    """Consume ONLY C8:D2887 pairs, keeping missing readings as null gaps."""
    day = date.fromisoformat(date_str)
    rows = list(rows)
    if len(rows) != 2880:
        raise ValueError("30SEC must contain all 2880 rows in C8:D2887.")
    points, missing, derived = [], 0, 0
    for index, (time_value, frequency_value) in enumerate(rows):
        stamp, was_derived = curve_timestamp(time_value, day, index)
        derived += int(was_derived)
        try:
            frequency = float(frequency_value) if not isinstance(frequency_value, bool) else None
        except (TypeError, ValueError):
            frequency = None
        if frequency is None or not math.isfinite(frequency) or frequency < 45:
            frequency = None
            missing += 1
        points.append({"timestamp": stamp.isoformat(timespec="seconds"), "frequency": frequency})
    if missing == len(points):
        raise ValueError("30SEC!D8:D2887 contains no valid frequency readings.")
    return points, {"missing_readings": missing, "derived_timestamps": derived}


_series_cache=OrderedDict()
_series_lock=RLock()


def load_curve_frequency_range(start_date, end_date, refresh=False):
    # Lazy import avoids a router/service import cycle. File naming, location,
    # timeout, workbook readers and fallback are owned by the PSP integration.
    from routes.psp_routes import get_psp_config_with_curve_defaults, read_curve_file_series

    try:
        start, end = date.fromisoformat(start_date), date.fromisoformat(end_date)
    except ValueError as exc:
        raise ValueError("Dates must be YYYY-MM-DD.") from exc
    if end < start:
        raise ValueError("End date must be on or after start date.")
    if (end - start).days >= 31:
        raise ValueError("Select up to 31 days per overview.")
    config = get_psp_config_with_curve_defaults()
    key=(start_date,end_date,json.dumps(config,sort_keys=True,default=str))
    with _series_lock:
        cached=_series_cache.get(key)
        if cached and not refresh and monotonic()-cached[0]<3600:return cached[1]
    points, sources, diagnostics = [], [], []
    current = start
    while current <= end:
        day = current.isoformat()
        series, meta = read_curve_file_series(day, [], config, frequency_only=True)
        if meta.get("available"):
            points.extend(series.get("frequency") or [])
            sources.append({"date": day, "file": meta.get("file"), "missing_readings": meta.get("missing_readings", 0), "derived_timestamps": meta.get("derived_timestamps", 0)})
        else:
            diagnostics.append({"date": day, "message": meta.get("message") or "Curve frequency is unavailable."})
        current += timedelta(days=1)
    stats = calculate_curve_frequency_stats(points, sample_seconds=30)
    response = {
        "success": bool(points), "start_date": start_date, "end_date": end_date,
        "timezone": "Asia/Kolkata", "sample_seconds": 30,
        "sheet": "30SEC", "frequency_range": "D8:D2887", "timestamp_range": "C8:C2887",
        "points": points, "sources": sources, "diagnostics": diagnostics,
        "statistics": stats,
    }
    if points:
        with _series_lock:
            _series_cache[key]=(monotonic(),response)
            _series_cache.move_to_end(key)
            while len(_series_cache)>2:_series_cache.popitem(last=False)
    return response


def calculate_curve_frequency_stats(points, sample_seconds=30):
    """Calculate frequency statistics and band distributions for loaded points."""
    valid_points = [
        p for p in points
        if p.get("frequency") is not None
        and isinstance(p["frequency"], (int, float))
        and math.isfinite(p["frequency"])
        and p["frequency"] >= 45
    ]
    total_samples = len(points)
    valid_samples = len(valid_points)
    missing_samples = total_samples - valid_samples
    step_minutes = sample_seconds / 60.0

    if not valid_samples:
        return {
            "total_samples": total_samples,
            "valid_samples": 0,
            "missing_samples": missing_samples,
            "covered_minutes": 0.0,
            "covered_hours": 0.0,
            "average_frequency": None,
            "max_frequency": None,
            "max_timestamp": None,
            "min_frequency": None,
            "min_timestamp": None,
            "rows": [],
        }

    min_point = min(valid_points, key=lambda p: p["frequency"])
    max_point = max(valid_points, key=lambda p: p["frequency"])
    total_freq = sum(p["frequency"] for p in valid_points)
    avg_freq = total_freq / valid_samples
    covered_minutes = valid_samples * step_minutes

    band_defs = [
        {"id": "normal", "label": "49.90 – 50.05 Hz", "condition": "49.90 ≤ F ≤ 50.05", "category": "Normal IEGC Operating Band", "test": lambda f: 49.90 <= f <= 50.05},
        {"id": "lt49_9", "label": "Frequency < 49.90 Hz", "condition": "F < 49.90", "category": "Low Frequency", "test": lambda f: f < 49.90},
        {"id": "lt49_7", "label": "Frequency < 49.70 Hz", "condition": "F < 49.70", "category": "Low Frequency (Emergency)", "test": lambda f: f < 49.70},
        {"id": "lt49_5", "label": "Frequency < 49.50 Hz", "condition": "F < 49.50", "category": "Low Frequency (Extreme)", "test": lambda f: f < 49.50},
        {"id": "gt50_05", "label": "Frequency > 50.05 Hz", "condition": "F > 50.05", "category": "High Frequency", "test": lambda f: f > 50.05},
        {"id": "gt50_10", "label": "Frequency > 50.10 Hz", "condition": "F > 50.10", "category": "High Frequency", "test": lambda f: f > 50.10},
        {"id": "gt50_20", "label": "Frequency > 50.20 Hz", "condition": "F > 50.20", "category": "High Frequency (Alert)", "test": lambda f: f > 50.20},
        {"id": "gt50_30", "label": "Frequency > 50.30 Hz", "condition": "F > 50.30", "category": "High Frequency (Warning)", "test": lambda f: f > 50.30},
        {"id": "gt50_50", "label": "Frequency > 50.50 Hz", "condition": "F > 50.50", "category": "High Frequency (Extreme)", "test": lambda f: f > 50.50},
    ]

    rows = []
    for b in band_defs:
        count = sum(1 for p in valid_points if b["test"](p["frequency"]))
        mins = count * step_minutes
        pct = (count / valid_samples * 100.0) if valid_samples else 0.0
        rows.append({
            "id": b["id"],
            "label": b["label"],
            "condition": b["condition"],
            "category": b["category"],
            "count": count,
            "minutes": round(mins, 2),
            "percent": round(pct, 2),
        })

    return {
        "total_samples": total_samples,
        "valid_samples": valid_samples,
        "missing_samples": missing_samples,
        "covered_minutes": round(covered_minutes, 2),
        "covered_hours": round(covered_minutes / 60.0, 2),
        "average_frequency": round(avg_freq, 3),
        "max_frequency": round(max_point["frequency"], 3),
        "max_timestamp": max_point["timestamp"],
        "min_frequency": round(min_point["frequency"], 3),
        "min_timestamp": min_point["timestamp"],
        "rows": rows,
    }


def format_excel_timestamp(stamp_str):
    if not stamp_str:
        return "—"
    clean = str(stamp_str).replace("T", " ")
    try:
        dt = datetime.fromisoformat(str(stamp_str).split("+")[0])
        return dt.strftime("%d-%b-%Y %H:%M:%S IST")
    except Exception:
        return f"{clean} IST"


def export_curve_frequency_excel_bytes(start_date, end_date, points=None, sample_seconds=30):
    if points is None:
        data = load_curve_frequency_range(start_date, end_date)
        points = data.get("points") or []
        sample_seconds = data.get("sample_seconds") or sample_seconds

    stats = calculate_curve_frequency_stats(points, sample_seconds=sample_seconds)

    wb = openpyxl.Workbook()

    NAVY_FILL = PatternFill("solid", fgColor="0F2F4F")
    HEADER_FONT = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    SUB_FONT = Font(name="Calibri", size=10, italic=True, color="475569")
    LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="0F2F4F")
    VALUE_FONT = Font(name="Calibri", size=10, color="1E293B")
    SECTION_FILL = PatternFill("solid", fgColor="047857")
    SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    TH_FILL = PatternFill("solid", fgColor="E2E8F0")
    TH_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")
    ROW_ALT_FILL = PatternFill("solid", fgColor="F8FAFC")

    NORMAL_FILL = PatternFill("solid", fgColor="ECFDF5")
    LOW_FILL = PatternFill("solid", fgColor="FEF2F2")
    HIGH_FILL = PatternFill("solid", fgColor="FFFBEB")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # ── Sheet 1: Frequency Statistics ─────────────────────────
    ws1 = wb.active
    ws1.title = "Frequency Statistics"
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:E1")
    ws1["A1"] = "EASTERN REGIONAL GRID — FREQUENCY OVERVIEW & STATISTICS"
    ws1["A1"].font = HEADER_FONT
    ws1["A1"].fill = NAVY_FILL
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 34

    ws1.merge_cells("A2:E2")
    ws1["A2"] = f"Period: {start_date} to {end_date} (IST) | Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} IST | Sample Interval: {sample_seconds}s"
    ws1["A2"].font = SUB_FONT
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    max_str = f"{stats['max_frequency']:.3f} Hz at {format_excel_timestamp(stats['max_timestamp'])}" if stats['max_frequency'] is not None else "—"
    min_str = f"{stats['min_frequency']:.3f} Hz at {format_excel_timestamp(stats['min_timestamp'])}" if stats['min_frequency'] is not None else "—"
    coverage_pct = (stats['valid_samples'] / stats['total_samples'] * 100.0) if stats['total_samples'] else 0.0

    meta_rows = [
        ("Analysis Date Range", f"{start_date} to {end_date}"),
        ("Total 30-Second Samples", f"{stats['total_samples']:,}"),
        ("Valid Frequency Readings", f"{stats['valid_samples']:,} ({coverage_pct:.2f}% coverage)"),
        ("Missing / Gap Readings", f"{stats['missing_samples']:,}"),
        ("Total Covered Duration", f"{stats['covered_minutes']:,.1f} Minutes ({stats['covered_hours']:.2f} Hours)"),
        ("Average Frequency", f"{stats['average_frequency']:.3f} Hz" if stats['average_frequency'] is not None else "—"),
        ("Maximum Frequency (& Date/Time)", max_str),
        ("Minimum Frequency (& Date/Time)", min_str),
    ]

    curr_row = 4
    for label, val in meta_rows:
        ws1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=2)
        c_lbl = ws1.cell(row=curr_row, column=1, value=label)
        c_lbl.font = LABEL_FONT
        c_lbl.fill = ROW_ALT_FILL
        c_lbl.alignment = Alignment(horizontal="left", vertical="center")
        c_lbl.border = thin_border

        ws1.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=5)
        c_val = ws1.cell(row=curr_row, column=3, value=val)
        c_val.font = VALUE_FONT
        c_val.alignment = Alignment(horizontal="left", vertical="center")
        c_val.border = thin_border

        if "Maximum Frequency" in label:
            c_val.font = Font(name="Calibri", size=10, bold=True, color="C2410C")
        elif "Minimum Frequency" in label:
            c_val.font = Font(name="Calibri", size=10, bold=True, color="DC2626")

        ws1.row_dimensions[curr_row].height = 22
        curr_row += 1

    curr_row += 1
    # Section Header
    ws1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=5)
    sec = ws1.cell(row=curr_row, column=1, value="FREQUENCY OPERATING BANDS & DURATION BREAKDOWN")
    sec.fill = SECTION_FILL
    sec.font = SECTION_FONT
    sec.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[curr_row].height = 26
    curr_row += 1

    headers = [
        "Frequency Band / Condition",
        "Category",
        "Duration (Minutes)",
        "% of Time",
        "Sample Count (30s)",
    ]
    for ci, h in enumerate(headers, 1):
        c = ws1.cell(row=curr_row, column=ci, value=h)
        c.fill = TH_FILL
        c.font = TH_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws1.row_dimensions[curr_row].height = 26
    curr_row += 1

    for r in stats["rows"]:
        c1 = ws1.cell(row=curr_row, column=1, value=r["label"])
        c2 = ws1.cell(row=curr_row, column=2, value=r["category"])
        c3 = ws1.cell(row=curr_row, column=3, value=r["minutes"])
        c4 = ws1.cell(row=curr_row, column=4, value=r["percent"])
        c5 = ws1.cell(row=curr_row, column=5, value=r["count"])

        fill = NORMAL_FILL if "Normal" in r["category"] else (LOW_FILL if "Low" in r["category"] else HIGH_FILL)
        for c in (c1, c2, c3, c4, c5):
            c.fill = fill
            c.border = thin_border

        c1.alignment = Alignment(horizontal="left", vertical="center")
        c1.font = Font(name="Calibri", size=10, bold=True)
        c2.alignment = Alignment(horizontal="left", vertical="center")
        c2.font = Font(name="Calibri", size=10)
        c3.alignment = Alignment(horizontal="right", vertical="center")
        c3.font = Font(name="Calibri", size=10, bold=True)
        c3.number_format = "#,##0.0"
        c4.alignment = Alignment(horizontal="right", vertical="center")
        c4.font = Font(name="Calibri", size=10, bold=True)
        c4.number_format = "0.00"
        c5.alignment = Alignment(horizontal="right", vertical="center")
        c5.font = Font(name="Calibri", size=10)
        c5.number_format = "#,##0"

        ws1.row_dimensions[curr_row].height = 22
        curr_row += 1

    col_widths_1 = {1: 30, 2: 32, 3: 20, 4: 16, 5: 20}
    for col_idx, width in col_widths_1.items():
        ws1.column_dimensions[get_column_letter(col_idx)].width = width

    # ── Sheet 2: 30SEC Frequency Data ─────────────────────────
    ws2 = wb.create_sheet(title="30SEC Frequency Data")
    ws2.views.sheetView[0].showGridLines = True

    ts_headers = [
        "S.No",
        "Timestamp (IST)",
        "Frequency (Hz)",
        "Deviation (Hz)",
        "Operating Band",
    ]
    for ci, h in enumerate(ts_headers, 1):
        c = ws2.cell(row=1, column=ci, value=h)
        c.fill = NAVY_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws2.row_dimensions[1].height = 28

    def band_label(f):
        if f is None or not math.isfinite(f):
            return "Missing / Invalid"
        if f < 49.50:
            return "< 49.50 Hz (Extreme Low)"
        if f < 49.70:
            return "< 49.70 Hz (Emergency Low)"
        if f < 49.90:
            return "< 49.90 Hz (Low)"
        if 49.90 <= f <= 50.05:
            return "49.90–50.05 Hz (Normal IEGC Band)"
        if f > 50.50:
            return "> 50.50 Hz (Extreme High)"
        if f > 50.30:
            return "> 50.30 Hz (Warning High)"
        if f > 50.20:
            return "> 50.20 Hz (Alert High)"
        if f > 50.10:
            return "> 50.10 Hz (High)"
        return "> 50.05 Hz (High)"

    for idx, p in enumerate(points, 1):
        row_num = idx + 1
        freq = p.get("frequency")
        dev = (freq - 50.0) if (freq is not None and isinstance(freq, (int, float)) and math.isfinite(freq)) else None
        b_text = band_label(freq)

        c1 = ws2.cell(row=row_num, column=1, value=idx)
        c2 = ws2.cell(row=row_num, column=2, value=str(p.get("timestamp") or "").replace("T", " "))
        c3 = ws2.cell(row=row_num, column=3, value=round(freq, 3) if freq is not None else "")
        c4 = ws2.cell(row=row_num, column=4, value=round(dev, 3) if dev is not None else "")
        c5 = ws2.cell(row=row_num, column=5, value=b_text)

        c1.alignment = Alignment(horizontal="center", vertical="center")
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c3.alignment = Alignment(horizontal="right", vertical="center")
        c4.alignment = Alignment(horizontal="right", vertical="center")
        c5.alignment = Alignment(horizontal="left", vertical="center")

        if freq is not None:
            c3.number_format = "0.000"
            c4.number_format = "+0.000;-0.000;0.000"

        if freq is None:
            c5.font = Font(name="Calibri", size=9, color="94A3B8")
        elif freq < 49.90:
            c5.font = Font(name="Calibri", size=9, bold=True, color="B91C1C")
        elif freq > 50.05:
            c5.font = Font(name="Calibri", size=9, bold=True, color="B45309")
        else:
            c5.font = Font(name="Calibri", size=9, color="047857")

        for c in (c1, c2, c3, c4, c5):
            c.border = thin_border

    ws2.freeze_panes = "B2"
    ws2.auto_filter.ref = ws2.dimensions
    col_widths_2 = {1: 10, 2: 24, 3: 18, 4: 18, 5: 34}
    for col_idx, width in col_widths_2.items():
        ws2.column_dimensions[get_column_letter(col_idx)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

