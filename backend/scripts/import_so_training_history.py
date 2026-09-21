"""Idempotently import the SO historical-training workbook into Crew Training.

The workbook is treated strictly as source data.  Historical records without an
exact date retain their financial year and stated duration; dates are never
invented.  Run without --apply for a dry run.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from bson import json_util
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from crew_legacy.api.training_assignment import shift_group_context  # noqa: E402
from crew_legacy.database.database_mongo import (  # noqa: E402
    employee_collection,
    training_master_collection,
    training_nomination_history_collection,
)


SHEET_LAYOUTS = {
    "Training details 2024-25": {"financial_year": "2024-25", "first_row": 3, "name_col": 2, "designation_col": 3, "employee_id_col": None, "training_cols": range(5, 18, 2)},
    "Training details 2025-26": {"financial_year": "2025-26", "first_row": 3, "name_col": 2, "designation_col": 3, "employee_id_col": 4, "training_cols": range(6, 19, 2)},
}

# Explicit aliases are used only where the workbook gives initials/short names.
# Missing workbook employee numbers are intentionally not redirected to a
# similarly named person.
NAME_TO_EMPLOYEE_ID = {
    "d biswas": "20020",
    "a p singh": "50042",
    "a k basak": "50040",
    "akash modi": "50047",
    "asit kumar das": "00205",
    "d mondal": "60034",
    "g patel": "00172",
    "g verma": "10028",
    "rakesh kr singh": "50084",
    "sandeep k maurya": "00137",
    "sanjeev k jha": "00304",
    "sharad kr yadav": "00307",
    "s s k suman": "50048",
    "s v agarwal": "00221",
}


def clean_text(value) -> str:
    return " ".join(str(value or "").replace("\u00a0", " ").split()).strip()


def name_key(value) -> str:
    text = clean_text(value).casefold().replace("kr.", "kumar").replace("kr ", "kumar ")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def training_key(value) -> str:
    return "".join(character for character in clean_text(value).casefold() if character.isalnum())


def employee_id_value(value) -> str:
    if value in (None, ""):
        return ""
    try:
        return str(int(float(str(value).strip()))).zfill(5)
    except (TypeError, ValueError):
        return clean_text(value).zfill(5)


def stable_key(*parts: str) -> str:
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()


def read_assignments(workbook_path: Path) -> list[dict]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    assignments: list[dict] = []
    try:
        for worksheet in workbook.worksheets:
            layout = SHEET_LAYOUTS.get(worksheet.title)
            if not layout:
                continue
            for source_row, row in enumerate(
                worksheet.iter_rows(min_row=layout["first_row"], values_only=True),
                layout["first_row"],
            ):
                employee_name = clean_text(row[layout["name_col"] - 1] if len(row) >= layout["name_col"] else "")
                if not employee_name or employee_name.casefold() == "control room":
                    continue
                designation = clean_text(row[layout["designation_col"] - 1] if len(row) >= layout["designation_col"] else "")
                provided_id = ""
                if layout["employee_id_col"]:
                    index = layout["employee_id_col"] - 1
                    provided_id = employee_id_value(row[index] if len(row) > index else "")
                for training_col in layout["training_cols"]:
                    title_index = training_col - 1
                    days_index = training_col
                    title = clean_text(row[title_index] if len(row) > title_index else "")
                    if not title:
                        continue
                    raw_days = row[days_index] if len(row) > days_index else None
                    try:
                        days = max(1, int(float(raw_days)))
                    except (TypeError, ValueError):
                        days = 1
                    assignments.append({
                        "financialYear": layout["financial_year"],
                        "employeeName": employee_name,
                        "providedEmployeeId": provided_id,
                        "designation": designation,
                        "trainingName": title,
                        "trainingDays": days,
                        "sourceSheet": worksheet.title,
                        "sourceRow": source_row,
                    })
    finally:
        workbook.close()
    return assignments


def employee_directory() -> tuple[dict[str, dict], dict[str, list[dict]]]:
    by_id: dict[str, dict] = {}
    by_name: dict[str, list[dict]] = defaultdict(list)
    for employee in employee_collection.find({}):
        employee_id = employee_id_value(employee.get("userId") or employee.get("employeeId"))
        if employee_id:
            by_id[employee_id] = employee
        by_name[name_key(employee.get("name"))].append(employee)
    return by_id, by_name


def resolve_employee(record: dict, by_id: dict[str, dict], by_name: dict[str, list[dict]]) -> tuple[dict | None, str]:
    provided_id = record["providedEmployeeId"]
    if provided_id:
        return (by_id.get(provided_id), "employee number" if provided_id in by_id else "employee number missing")
    alias_id = NAME_TO_EMPLOYEE_ID.get(name_key(record["employeeName"]))
    if alias_id:
        return (by_id.get(alias_id), "verified alias" if alias_id in by_id else "alias employee missing")
    exact = by_name.get(name_key(record["employeeName"])) or []
    if len(exact) == 1:
        return exact[0], "exact name"
    return None, "employee not found" if not exact else "ambiguous employee name"


def write_review(path: Path, rows: list[dict], summary: dict) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Import review"
    headers = ["Financial Year", "Employee Name (source)", "Employee No. (source)", "Resolved Employee", "Resolved Employee No.", "Training", "Days", "Status", "Match Method", "Source"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
    green = PatternFill("solid", fgColor="DCFCE7")
    red = PatternFill("solid", fgColor="FECACA")
    for item in rows:
        sheet.append([
            item["financialYear"], item["employeeName"], item["providedEmployeeId"],
            item.get("resolvedEmployeeName") or "", item.get("resolvedEmployeeId") or "",
            item["trainingName"], item["trainingDays"], item["status"], item["matchMethod"],
            f'{item["sourceSheet"]}!{item["sourceRow"]}',
        ])
        fill = green if item["status"] == "IMPORTED" else red
        for cell in sheet[sheet.max_row]:
            cell.fill = fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [16, 27, 18, 27, 20, 65, 9, 18, 22, 28]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    info = workbook.create_sheet("Summary")
    info.append(["Metric", "Value"])
    for key, value in summary.items():
        info.append([key, value])
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True, type=Path)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--review", type=Path)
    args = parser.parse_args()

    assignments = read_assignments(args.file)
    by_id, by_name = employee_directory()
    reviewed: list[dict] = []
    programmes: dict[tuple[str, str], dict] = {}

    for record in assignments:
        employee, method = resolve_employee(record, by_id, by_name)
        resolved_id = employee_id_value((employee or {}).get("userId") or (employee or {}).get("employeeId"))
        reviewed_record = {
            **record,
            "resolvedEmployeeId": resolved_id,
            "resolvedEmployeeName": clean_text((employee or {}).get("name")),
            "status": "IMPORTED" if employee else "EMPLOYEE MISSING",
            "matchMethod": method,
        }
        reviewed.append(reviewed_record)
        programme = programmes.setdefault(
            (record["financialYear"], training_key(record["trainingName"])),
            {
                "financialYear": record["financialYear"],
                "trainingName": record["trainingName"],
                "durations": set(),
                "matched": set(),
                "unmatched": {},
            },
        )
        programme["durations"].add(record["trainingDays"])
        if employee:
            programme["matched"].add(resolved_id)
        else:
            missing_key = f'{record["employeeName"]}|{record["providedEmployeeId"]}'
            programme["unmatched"][missing_key] = {
                "name": record["employeeName"],
                "providedEmployeeId": record["providedEmployeeId"],
                "designation": record["designation"],
                "trainingDays": record["trainingDays"],
                "sourceSheet": record["sourceSheet"],
                "sourceRow": record["sourceRow"],
            }

    now = datetime.utcnow()
    if args.apply:
        backup_dir = BACKEND_ROOT / "backups"
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup_path = backup_dir / f"training_history_before_so_import_{now:%Y%m%d_%H%M%S}.json"
        backup = {
            "training_master_collection": list(training_master_collection.find({"financialYear": {"$in": ["2024-25", "2025-26"]}})),
            "training_nomination_history_collection": list(training_nomination_history_collection.find({"$or": [{"financialYear": {"$in": ["2024-25", "2025-26"]}}, {"historicalImport.source": args.file.name}]})),
        }
        backup_path.write_text(json_util.dumps(backup, indent=2), encoding="utf-8")

        for (_, normalized_training), programme in programmes.items():
            import_key = stable_key(programme["financialYear"], normalized_training)
            durations = sorted(programme["durations"])
            training_master_collection.update_one(
                {"historicalImportKey": import_key},
                {
                    "$set": {
                        "financialYear": programme["financialYear"],
                        "trainingName": programme["trainingName"],
                        "location": "",
                        "startDate": None,
                        "endDate": None,
                        "durationDays": max(durations),
                        "durationOptions": durations,
                        "dateStatus": "Financial year only - exact dates not supplied",
                        "trainingType": "Historical",
                        "status": "Completed",
                        "historicalImport": True,
                        "historicalImportKey": import_key,
                        "historicalSource": args.file.name,
                        "matchedEmployeeCount": len(programme["matched"]),
                        "unmatchedEmployees": list(programme["unmatched"].values()),
                        "updatedOn": now,
                    },
                    "$setOnInsert": {"createdOn": now},
                },
                upsert=True,
            )

        for item in reviewed:
            if item["status"] != "IMPORTED":
                continue
            employee = by_id[item["resolvedEmployeeId"]]
            context = shift_group_context(item["resolvedEmployeeId"])
            import_key = stable_key(
                item["financialYear"], item["resolvedEmployeeId"],
                training_key(item["trainingName"]), str(item["sourceRow"]), item["sourceSheet"],
            )
            training_nomination_history_collection.update_one(
                {"historicalImportKey": import_key},
                {
                    "$set": {
                        "financialYear": item["financialYear"],
                        "trainingName": item["trainingName"],
                        "trainingLocation": "",
                        "trainingDate": None,
                        "startDate": None,
                        "endDate": None,
                        "trainingDays": item["trainingDays"],
                        "employeeId": item["resolvedEmployeeId"],
                        "employeeName": clean_text(employee.get("name")),
                        "employeeDesignation": employee.get("designation") or item["designation"],
                        "employeeType": "Shift" if context else "Non-shift",
                        "groupName": context.get("groupName") if context else None,
                        "isShiftEmployee": bool(context),
                        "isGroupSIC": bool((context or {}).get("isGroupSIC")),
                        "replacementRequired": False,
                        "workflowKind": "Training",
                        "requestType": "Historical Import",
                        "status": "Approved",
                        "approvalChain": [],
                        "currentApprovalIndex": 0,
                        "currentApproverId": None,
                        "historicalImport": True,
                        "historicalImportKey": import_key,
                        "historicalSource": {
                            "file": args.file.name,
                            "sheet": item["sourceSheet"],
                            "row": item["sourceRow"],
                            "sourceEmployeeName": item["employeeName"],
                            "matchMethod": item["matchMethod"],
                        },
                        "updatedOn": now,
                        "approvedOn": now,
                    },
                    "$setOnInsert": {"createdOn": now},
                },
                upsert=True,
            )
    else:
        backup_path = None

    missing_assignments = [item for item in reviewed if item["status"] != "IMPORTED"]
    summary = {
        "Mode": "APPLIED" if args.apply else "DRY RUN",
        "Source assignments": len(reviewed),
        "Verified employee assignments": len(reviewed) - len(missing_assignments),
        "Missing employee assignments": len(missing_assignments),
        "Training Master programmes": len(programmes),
        "2024-25 assignments": sum(1 for item in reviewed if item["financialYear"] == "2024-25"),
        "2025-26 assignments": sum(1 for item in reviewed if item["financialYear"] == "2025-26"),
        "Backup": str(backup_path or "Not created during dry run"),
    }
    review_path = args.review or args.file.with_name(f"{args.file.stem}_import_review.xlsx")
    write_review(review_path, reviewed, summary)
    print(json.dumps({**summary, "Review workbook": str(review_path)}, ensure_ascii=False, indent=2))
    if missing_assignments:
        print("Missing employees:")
        for name, provided_id in sorted({(item["employeeName"], item["providedEmployeeId"]) for item in missing_assignments}):
            print(f"- {name} ({provided_id or 'no employee number'})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
